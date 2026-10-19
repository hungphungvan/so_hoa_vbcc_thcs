# /// script
# requires-python = ">=3.10"
# dependencies = [
#     "openpyxl",
# ]
# ///
import os
import sys
import subprocess
import datetime
import math
import re
import io
import copy
import unicodedata
from collections import Counter
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.utils.datetime import from_excel
from openpyxl.utils import get_column_letter

# Đảm bảo mã hóa UTF-8 cho stdout trên mọi nền tảng
if sys.stdout and hasattr(sys.stdout, 'reconfigure'):
    try:
        sys.stdout.reconfigure(encoding='utf-8')
    except Exception:
        pass

# ==============================================================================
# HẰNG SỐ & ĐƯỜNG DẪN MẪU (Hỗ trợ cả môi trường PyInstaller .exe)
# ==============================================================================
if getattr(sys, 'frozen', False) and hasattr(sys, '_MEIPASS'):
    # Môi trường chạy file .exe đóng gói bởi PyInstaller
    SCRIPT_DIR = Path(sys._MEIPASS)
    if not (SCRIPT_DIR / "template.xlsx").exists():
        SCRIPT_DIR = Path(sys.executable).resolve().parent
else:
    SCRIPT_DIR = Path(__file__).resolve().parent

TEMPLATE_PATH = SCRIPT_DIR / "template.xlsx"

# Danh sách 28 cột chuẩn hóa theo file template.xlsx
TEMPLATE_COL_NAMES = [
    "STT",                                                    # 0 (A)
    "Số hiệu bằng",                                           # 1 (B)
    "Số vào sổ gốc cấp văn bằng",                             # 2 (C)
    "Tên văn bằng",                                           # 3 (D)
    "Họ, chữ đệm và tên",                                     # 4 (E)
    "Mã người học",                                           # 5 (F)
    "Số định danh cá nhân",                                   # 6 (G)
    "Hộ chiếu",                                               # 7 (H)
    "Ngày, tháng, năm sinh",                                  # 8 (I)
    "Giới tính",                                              # 9 (J)
    "Dân tộc",                                                # 10 (K)
    "Quốc tịch",                                              # 11 (L)
    "Nơi sinh",                                               # 12 (M)
    "Tên trường đã học",                                      # 13 (N)
    "Niên khóa",                                              # 14 (O)
    "Điểm các môn thi",                                       # 15 (P)
    "Hội đồng thi",                                           # 16 (Q)
    "Địa danh nơi cơ quan cấp bằng đặt trụ sở",               # 17 (R)
    "Ngày tháng năm cấp bằng",                                # 18 (S)
    "Tên cơ quan cấp bằng",                                   # 19 (T)
    "Chức danh người ký bằng",                                # 20 (U)
    "Họ, chữ đệm, tên người ký bằng",                         # 21 (V)
    "Tình trạng văn bằng",                                    # 22 (W)
    "Ghi chú",                                                # 23 (X)
    "Trạng thái số hóa",                                      # 24 (Y)
    "Thông tin chưa bảo đảm đầy đủ, chính xác và nguyên nhân",# 25 (Z)
    "Mã đơn vị",                                              # 26 (AA)
    "Tên đơn vị"                                              # 27 (AB)
]


# ==============================================================================
# CÁC HÀM TIỆN ÍCH HỖ TRỢ
# ==============================================================================
def get_folder_path(prompt="Chọn thư mục chứa các file Excel:"):
    """Mở hộp thoại native trên macOS hoặc Windows để chọn thư mục, hỗ trợ kéo thả thư mục vào Terminal/CMD."""
    if sys.platform == "darwin":
        try:
            script = f'return POSIX path of (choose folder with prompt "{prompt}")'
            result = subprocess.check_output(['osascript', '-e', script])
            path_str = result.decode('utf-8').strip()
            if path_str:
                return path_str
        except Exception:
            pass
    elif sys.platform == "win32":
        try:
            ps_script = (
                "[System.Reflection.Assembly]::LoadWithPartialName('System.windows.forms') | Out-Null; "
                f"$f = New-Object System.Windows.Forms.FolderBrowserDialog; $f.Description = '{prompt}'; "
                "if($f.ShowDialog() -eq 'OK'){ $f.SelectedPath }"
            )
            result = subprocess.check_output(["powershell", "-NoProfile", "-Command", ps_script])
            path_str = result.decode('utf-8').strip()
            if path_str:
                return path_str
        except Exception:
            pass

    # Nếu hộp thoại bị hủy hoặc lỗi, cho phép nhập hoặc kéo thả thư mục vào Terminal / CMD
    print("👉 Hoặc bạn có thể KÉO THẢ thư mục trực tiếp vào đây (hoặc dán đường dẫn):")
    manual_path = input("   Đường dẫn thư mục (bấm Enter để hủy): ").strip()
    if manual_path:
        clean_path = manual_path.strip("'\"").replace('\\ ', ' ')
        if Path(clean_path).is_dir():
            return clean_path
        else:
            print(f"⚠️ Không tìm thấy thư mục: {clean_path}")
    return None

get_folder_path_mac = get_folder_path  # Giữ alias tương thích ngược

def get_file_path(prompt="Chọn file Excel cần xử lý:"):
    """Mở hộp thoại native trên macOS hoặc Windows để chọn file Excel, hỗ trợ kéo thả file vào Terminal/CMD."""
    if sys.platform == "darwin":
        try:
            script = f'return POSIX path of (choose file with prompt "{prompt}" of type {{"xlsx", "org.openxmlformats.spreadsheetml.sheet"}})'
            result = subprocess.check_output(['osascript', '-e', script])
            path_str = result.decode('utf-8').strip()
            if path_str:
                return path_str
        except Exception:
            pass
    elif sys.platform == "win32":
        try:
            ps_script = (
                "[System.Reflection.Assembly]::LoadWithPartialName('System.windows.forms') | Out-Null; "
                f"$f = New-Object System.Windows.Forms.OpenFileDialog; $f.Title = '{prompt}'; $f.Filter = 'Excel Files (*.xlsx)|*.xlsx'; "
                "if($f.ShowDialog() -eq 'OK'){ $f.FileName }"
            )
            result = subprocess.check_output(["powershell", "-NoProfile", "-Command", ps_script])
            path_str = result.decode('utf-8').strip()
            if path_str:
                return path_str
        except Exception:
            pass

    # Nếu hộp thoại bị hủy hoặc lỗi, cho phép nhập hoặc kéo thả file vào Terminal / CMD
    print("👉 Hoặc bạn có thể KÉO THẢ file Excel trực tiếp vào đây (hoặc dán đường dẫn):")
    manual_path = input("   Đường dẫn file Excel (bấm Enter để hủy): ").strip()
    if manual_path:
        clean_path = manual_path.strip("'\"").replace('\\ ', ' ')
        if Path(clean_path).is_file():
            return clean_path
        else:
            print(f"⚠️ Không tìm thấy file: {clean_path}")
    return None

get_file_path_mac = get_file_path  # Giữ alias tương thích ngược

def clean_whitespace(value):
    """Làm sạch khoảng trắng thừa trong chuỗi, xử lý nan/none."""
    if value is None:
        return ""
    if isinstance(value, float):
        if math.isnan(value):
            return ""
    if isinstance(value, str):
        cleaned = re.sub(r'\s+', ' ', value).strip()
        if cleaned.lower() in ("none", "nan", "null", "<na>"):
            return ""
        return cleaned
    return value

def sanitize_filename(name):
    """
    Chuẩn hóa chuỗi để dùng làm tên file hợp lệ trên cả Windows và macOS.
    Loại bỏ các ký tự cấm: \\ / : * ? " < > |
    """
    if not name:
        return "Chua_Xac_Dinh"
    s = str(name).strip()
    # Thay thế các ký tự cấm bằng dấu gạch dưới
    s = re.sub(r'[\\/*?:"<>|]', '_', s)
    # Rút gọn khoảng trắng và gạch dưới lặp
    s = re.sub(r'\s+', ' ', s).strip()
    s = re.sub(r'_+', '_', s).strip('._ ')
    return s if s else "Chua_Xac_Dinh"

def get_col_letter(col_idx):
    """Lấy ký hiệu chữ của cột từ index 0 (0 -> A, 1 -> B, ..., 27 -> AB, 28 -> AC...)."""
    return get_column_letter(col_idx + 1)

def is_cell_empty(val):
    """Kiểm tra ô có bị trống dữ liệu không."""
    if val is None:
        return True
    if isinstance(val, float):
        if math.isnan(val):
            return True
    s = str(val).strip()
    return s == "" or s.lower() in ["none", "nan", "null", "<na>"]

def is_cell_yellow(cell):
    """Kiểm tra xem ô Excel có đang được tô màu vàng hay không."""
    if cell.fill and cell.fill.start_color and cell.fill.start_color.rgb:
        rgb_str = str(cell.fill.start_color.rgb).upper()
        if "FFFF00" in rgb_str:
            return True
    return False

def is_valid_stt(val):
    """Kiểm tra ô có phải là Số thứ tự hợp lệ của học sinh (số nguyên dương: 1, 2, 3...)."""
    if val is None or isinstance(val, bool):
        return False
    if isinstance(val, (int, float)):
        if math.isnan(val):
            return False
        return val > 0 and float(val).is_integer()
    s = str(val).strip()
    s_clean = re.sub(r'\.0+$', '', s)
    return s_clean.isdigit() and int(s_clean) > 0

SUMMARY_KEYWORDS = [
    "danh sách này có", "danh sách có", "danh sách gồm",
    "tổng số", "tổng cộng", "số học sinh",
    "công nhận tốt nghiệp", "người lập", "người ký", "hiệu trưởng", "thủ trưởng"
]

def is_summary_or_footer_row(row_cells):
    """Kiểm tra xem dòng có phải là dòng tổng kết, ghi chú cuối danh sách hoặc chân trang ký tên hay không."""
    for cell in row_cells[:15]:
        val = str(cell.value or '').strip().lower()
        if any(kw in val for kw in SUMMARY_KEYWORDS):
            return True
    return False

def get_cached_row_fonts(ws, reference_row_idx, max_col=28):
    """Trích xuất và lưu sẵn (cache) định dạng Font từ dòng mẫu để tái sử dụng."""
    cached = []
    for col in range(1, max_col + 1):
        cell = ws.cell(row=reference_row_idx, column=col)
        if cell.has_style and cell.font:
            cached.append(Font(name=cell.font.name, size=cell.font.size))
        else:
            cached.append(None)
    return cached

def apply_cached_fonts(ws, target_row_idx, cached_fonts):
    """Gán font từ bộ nhớ cache vào dòng mục tiêu."""
    for col_idx, font in enumerate(cached_fonts, start=1):
        if font:
            ws.cell(row=target_row_idx, column=col_idx).font = font

def get_file_format_info(ws):
    """
    Nhận diện định dạng sheet:
    - Quét tối đa 30 dòng đầu tìm dòng tiêu đề chứa 'STT' hoặc 'Họ và tên' (không phân biệt hoa/thường).
    - Tìm cột 'Tên đơn vị' để phân loại 'AE' (31 cột) hoặc 'AB' (28 cột).
    - Trả về (format_type, headers, header_row_idx)
    """
    for r_idx, row in enumerate(ws.iter_rows(min_row=1, max_row=30, values_only=True), start=1):
        if not row:
            continue
        row_str = [clean_whitespace(str(cell)) if cell is not None else "" for cell in row]
        row_lower = [s.lower() for s in row_str]

        # Kiểm tra xem dòng này có phải dòng tiêu đề không
        has_stt = any(s in ["stt", "số tt", "số thứ tự", "stt.", "tt", "tt."] for s in row_lower[:5])
        has_name = any(("họ" in s and "tên" in s) for s in row_lower)

        if has_stt or has_name:
            headers = row_str
            col_idx = -1
            for idx, h in enumerate(row_lower):
                if "tên đơn vị" in h or h == "đơn vị":
                    col_idx = idx
                    break

            if col_idx >= 29:
                return "AE", headers, r_idx
            elif 25 <= col_idx <= 28:
                return "AB", headers, r_idx
            elif col_idx == -1 and has_name:
                # Dự phòng: nếu tiêu đề không có chữ 'Tên đơn vị', suy đoán qua số cột có dữ liệu
                non_empty_cols = sum(1 for c in headers if c)
                if non_empty_cols >= 30:
                    return "AE", headers, r_idx
                elif non_empty_cols >= 26:
                    return "AB", headers, r_idx
                return "INVALID", headers, r_idx
            else:
                return "INVALID", headers, r_idx

    return "INVALID", [], None

# ==============================================================================
# DANH MỤC ÁNH XẠ MÃ SỐ SANG VĂN BẢN (4 SHEET CHUẨN)
# ==============================================================================
GENDER_MAP = {
    1: "Nam",
    2: "Khác",
    3: "Nữ"
}

ETHNIC_MAP = {
    1: "Mạ",
    2: "Kháng",
    3: "Hà Nhì",
    4: "Hrê",
    5: "Lào",
    6: "Mường",
    7: "Chu-ru",
    8: "Mnông",
    9: "Ra-glai",
    10: "Cống",
    11: "Bru-Vân Kiều",
    12: "Si La",
    13: "Xinh-mun",
    14: "La Hủ",
    15: "Khác",
    16: "Dao",
    17: "Xtiêng",
    18: "Lự",
    19: "Pu Péo",
    20: "Bố Y",
    21: "Cơ-tu",
    22: "Phù Lá",
    23: "Chăm",
    24: "Gia-rai",
    25: "Ba-na",
    26: "Ngái",
    27: "Co",
    28: "Sán Chay",
    29: "Kinh",
    30: "Ơ-đu",
    31: "Chứt",
    32: "Rơ-măm",
    33: "Thái",
    34: "La Ha",
    35: "La-chí",
    36: "Giáy",
    37: "Ê-đê",
    38: "Sán Dìu",
    39: "Gié-Triêng",
    40: "Mảng",
    41: "Nùng",
    42: "LôLô",
    43: "Brâu",
    44: "Chơ-ro",
    45: "Thổ",
    46: "Khơmú",
    47: "Ta-ôi",
    48: "H'Mông",
    49: "Pà Thẻn",
    50: "Tày",
    51: "Cờ lao"
}

NATIONALITY_MAP = {
    1: "Lit-va",
    2: "Kê-nia",
    3: "Hy Lạp",
    4: "Dăm-bia",
    5: "Vê-nê-xu-ê-la",
    6: "CH Trung Phi",
    7: "Cô-lôm-bi-a",
    8: "Nam-mi-bia",
    9: "Thụy Sỹ",
    10: "Et-to-ni",
    11: "Giooc-đa-ni",
    12: "Cooxxta Rica",
    13: "Na Uy",
    14: "I-ra-li-a",
    15: "Xu-đăng",
    16: "Dominica",
    17: "Pa-Ki-xtan",
    18: "Ka-dắc-xtan",
    19: "Li Băng",
    20: "Áp-ga-ni-xtan",
    21: "Turkmênistan",
    22: "Phần Lan",
    23: "Quatar",
    24: "Vương Quốc Anh",
    25: "Sri-lan-ka",
    26: "Ê-cu-a-đo",
    27: "Goa-tê-ma-la",
    28: "Công-Kinshasa",
    29: "Ả rập Xê út",
    30: "Pa-na-ma",
    31: "Hà Lan",
    32: "CHDCND Lào",
    33: "Ma-la0uy",
    34: "Ma-na-cô",
    35: "Guy-a-na",
    36: "Pa-pua Niu Ghi - nê",
    37: "Gam-bi-a",
    38: "Ha-i-ti",
    39: "Libi",
    40: "Thái Lan",
    41: "Môn-đô-va",
    42: "Căm-Pu-Chia",
    43: "Úc",
    44: "Đài Loan",
    45: "Tây Ban Nha",
    46: "Nam phi",
    47: "Bồ Đào Nha",
    48: "No-giê",
    49: "Gha-na",
    50: "I-rắc",
    51: "Ni Di Lân",
    52: "Ca-na-đa",
    53: "Bu-run-đi",
    54: "CHND Trung Hoa",
    55: "Tat-gi-ki-xtan",
    56: "Dim-ba-buê",
    57: "Mi-cờ-rô-nê-xia-a",
    58: "Ru-ma-no",
    59: "Nga",
    60: "Phi-li-pin",
    61: "Croát-chia-a",
    62: "Be-lo-rut-xia",
    63: "Sat",
    64: "Ba-ranh",
    65: "U-krai-na",
    66: "En Xan-va-đo",
    67: "Buốc-ki-na-pha-xô",
    68: "Ma Rốc",
    69: "Liên Bang Nga",
    70: "Cư-rơ-gư-dơ-xtan",
    71: "I-ran",
    72: "Pu-ec-to Ri-cô",
    73: "Nê-Pan",
    74: "My-an-ma",
    75: "Quần đảo Sô-lô-môn",
    76: "Cô-oét",
    77: "Ôn-đu-rát",
    78: "Hàn Quốc",
    79: "U-dơ-bê-ki-xtan",
    80: "Tôn-ga",
    81: "Sec-bi-a",
    82: "Ma-đa-gát-xka",
    83: "Việt Nam",
    84: "Ma-lai-xi-a",
    85: "Man-đi-vơ",
    86: "Hung-ga-ri",
    87: "Tan-da-nia",
    88: "Thụy Điển",
    89: "Việt Nam",
    90: "Ma Cao",
    91: "Mỹ",
    92: "Ma-li",
    93: "Các Tiểu vương quốc Ả Rập Thống Nhất",
    94: "Tu-va-lu",
    95: "Li-bê-ri-a",
    96: "B-ru-nây",
    97: "I-xra-en",
    98: "U-ru-goay",
    99: "Phi-Ghi",
    100: "Cuba",
    101: "Ai-xlen",
    102: "Latvia",
    103: "Hồng Công",
    104: "Xê nê gan",
    105: "Man-ta",
    106: "Pa-ra-goay",
    107: "Mông-tê-nê-groo",
    108: "Thổ Nhĩ Kỳ",
    109: "Mông Cổ",
    110: "Mê hi cô",
    111: "Đông -ti-mo",
    112: "Quần đảo Ma-rơ-san",
    113: "Slô-vac-ki-a",
    114: "Ô man",
    115: "Gru-dia",
    116: "An giê ri",
    117: "Ki-ri-ba-ti",
    118: "Ấn Độ",
    119: "Ni-ca-ra-goa",
    120: "Pê - ru",
    121: "Ruanđa",
    122: "Gia-mai-ca",
    123: "Bun-ga-ri",
    124: "Luc-xem-bua",
    125: "Ba Lan",
    126: "Nam Xu-đăng",
    127: "Chi lê",
    128: "Xi-ga-po",
    129: "U-gan-da",
    130: "Nhật Bản",
    131: "Si-ri-a",
    132: "A-rơ-len",
    133: "Ghi nê",
    134: "Tuy-ni-di",
    135: "Áo",
    136: "Đan Mạch",
    137: "Sô-ma-li",
    138: "Mau-ri-ti-út",
    139: "Cộng hòa Séc",
    140: "Đức",
    141: "To-go",
    142: "Pháp",
    143: "Va-nu-a-tu",
    144: "Palestine",
    145: "A-déc-bai-gian",
    146: "Ăng -gô-la",
    147: "Can-mơ-run",
    148: "Lit-ten-xten",
    149: "Tây Sa-moa",
    150: "Ma-xê-đô-ni-a",
    151: "Ni-giê-ri-a",
    152: "Y-ê-men",
    153: "In-đô-nê-xi-a",
    154: "Mô-dăm-bich",
    155: "Ac-hen-ti-na",
    156: "Xi-ê-a Lê- ôn",
    157: "Slô-ven-ni-a",
    158: "Triều Tiên",
    159: "Ai cập"
}

BIRTHPLACE_MAP = {
    1: "Tỉnh Bắc Ninh",
    2: "Tỉnh Bà Rịa - Vũng Tàu",
    3: "Thành phố Hải Phòng",
    4: "Tỉnh Lâm Đồng",
    5: "Tỉnh Phú Yên",
    6: "Tỉnh Thái Nguyên",
    7: "Tỉnh Đắk Lắk",
    8: "Tỉnh Kon Tum",
    9: "Tỉnh Nghệ Tĩnh",
    10: "Tỉnh Long An",
    11: "Tỉnh Vĩnh Phúc",
    12: "Tỉnh Yên Bái",
    13: "Tỉnh Thừa Thiên Huế",
    14: "Tỉnh Quảng Nam",
    15: "Tỉnh Hà Tây",
    16: "Tỉnh Bình Trị Thiên",
    17: "Tỉnh Trà Vinh",
    18: "Tỉnh Hà Nam Ninh",
    19: "Tỉnh Ninh Bình",
    20: "Tỉnh Lạng Sơn",
    21: "Tỉnh Quảng Ninh",
    22: "Tỉnh Quảng Bình",
    23: "Tỉnh Gia Lai",
    24: "Tỉnh Đồng Tháp",
    25: "Tỉnh Tiền Giang",
    26: "Thành phố Hà Nội",
    27: "Tỉnh Phú Khánh",
    28: "Tỉnh Bình Dương",
    29: "Tỉnh Bắc Giang",
    30: "Tỉnh Bắc Kạn",
    31: "Tỉnh Phú Thọ",
    32: "Tỉnh Hà Giang",
    33: "Tỉnh Minh Hải",
    34: "Tỉnh Hà Bắc",
    35: "Tỉnh Quảng Ngãi",
    36: "Tỉnh Đồng Nai",
    37: "Tỉnh Hoà Bình",
    38: "Tỉnh Thái Bình",
    39: "Tỉnh Điện Biên",
    40: "Tỉnh Bình Định",
    41: "Tỉnh Lào Cai",
    42: "Tỉnh Hải Hưng",
    43: "Tỉnh Đắk Nông",
    44: "Tỉnh Hải Dương",
    45: "Tỉnh Hưng Yên",
    46: "Tỉnh Kiên Giang",
    47: "Thành phố Cần Thơ",
    48: "Thành phố Hồ Chí Minh",
    49: "Tỉnh Bạc Liêu",
    50: "Tỉnh Lai Châu",
    51: "Tỉnh Cà Mau",
    52: "Tỉnh Hà Sơn Bình",
    53: "Tỉnh Quảng Trị",
    54: "Tỉnh An Giang",
    55: "Thành phố Đà Nẵng",
    56: "Tỉnh Cao Bằng",
    57: "Tỉnh Tuyên Quang",
    58: "Tỉnh Nghệ An",
    59: "Tỉnh Bình Thuận",
    60: "Tỉnh Sơn La",
    61: "Tỉnh Bến Tre",
    62: "Tỉnh Khánh Hòa",
    63: "Tỉnh Vĩnh Phú",
    64: "Tỉnh Ninh Thuận",
    65: "Tỉnh Hà Nam",
    66: "Tỉnh Bình Phước",
    67: "Tỉnh Tây Ninh",
    68: "Tỉnh Nam Định",
    69: "Tỉnh Sông Bé",
    70: "Tỉnh Hà Tĩnh",
    71: "Tỉnh Thanh Hóa",
    72: "Tỉnh Bắc Thái",
    73: "Tỉnh Vĩnh Long",
    74: "Tỉnh Sóc Trăng",
    75: "Tỉnh Hậu Giang"
}

def map_code_to_text(col_idx, value):
    """
    Chuyển đổi mã số sang văn bản theo 4 danh mục chuẩn:
    - Cột J (Index 9 - Giới tính): 1 -> Nam, 2 -> Khác, 3 -> Nữ
    - Cột K (Index 10 - Dân tộc): 1..51 -> Tên dân tộc (29 -> Kinh)
    - Cột L (Index 11 - Quốc tịch): 1..159 -> Tên quốc tịch (83, 89 -> Việt Nam)
    - Cột M (Index 12 - Nơi sinh): 1..75 -> Tên tỉnh/thành phố
    Nếu giá trị đã là tên chữ hoặc không có trong danh mục, giữ nguyên giá trị hợp lệ.
    """
    if value is None:
        return ""
    s_val = str(value).strip()
    if not s_val or s_val.lower() in ("none", "nan", "null", "<na>"):
        return ""

    # Loại bỏ phần thập phân .0 nếu là số nguyên lưu dạng float (ví dụ: 1.0 -> 1, 29.0 -> 29)
    s_clean = re.sub(r'\.0+$', '', s_val)

    # 1. Cột J (Index 9): Giới tính
    if col_idx == 9:
        if s_clean.isdigit():
            c_num = int(s_clean)
            if c_num in GENDER_MAP:
                return GENDER_MAP[c_num]
        s_lower = s_clean.lower()
        if s_lower == "nam":
            return "Nam"
        elif s_lower in ["nữ", "nu"]:
            return "Nữ"
        elif s_lower in ["khác", "khac"]:
            return "Khác"
        return s_val

    # 2. Cột K (Index 10): Dân tộc
    elif col_idx == 10:
        if s_clean.isdigit():
            c_num = int(s_clean)
            if c_num in ETHNIC_MAP:
                return ETHNIC_MAP[c_num]
        if s_clean.lower() == "kinh":
            return "Kinh"
        return s_val

    # 3. Cột L (Index 11): Quốc tịch
    elif col_idx == 11:
        if s_clean.isdigit():
            c_num = int(s_clean)
            if c_num in NATIONALITY_MAP:
                return NATIONALITY_MAP[c_num]
        if s_clean.lower() in ["việt nam", "viet nam", "vn"]:
            return "Việt Nam"
        return s_val

    # 4. Cột M (Index 12): Nơi sinh
    elif col_idx == 12:
        if s_clean.isdigit():
            c_num = int(s_clean)
            if c_num in BIRTHPLACE_MAP:
                return BIRTHPLACE_MAP[c_num]
        return s_val

    return value


# ==============================================================================
# CHUẨN HÓA & ĐỐI SOÁT NIÊN KHÓA (CỘT O) VÀ NGÀY CẤP BẰNG (CỘT S)
# ==============================================================================
def parse_nien_khoa(val_nien_khoa):
    """
    Kiểm tra và bóc tách Niên khóa (Cột O):
    - Chuẩn hóa về dạng YYYY-YYYY.
    - Trả về: (is_valid_diff_4, y_vao, y_ra, formatted_str)
      Nếu hiệu Y_ra - Y_vao == 4: is_valid_diff_4 = True.
    """
    if val_nien_khoa is None:
        return False, None, None, ""
    s_val = str(val_nien_khoa).strip()
    if not s_val or s_val.lower() in ("none", "nan", "null", "<na>"):
        return False, None, None, ""

    s_clean = clean_whitespace(s_val)
    s_clean = re.sub(r'\.0+$', '', s_clean)

    m = re.search(r'(\d{4})\s*[-/–—._~]\s*(\d{2,4})', s_clean)
    if m:
        y1_str, y2_str = m.groups()
        y1 = int(y1_str)
        if len(y2_str) == 2:
            y2_val = int(y2_str)
            century = (y1 // 100) * 100
            if y2_val < (y1 % 100):
                y2 = century + 100 + y2_val
            else:
                y2 = century + y2_val
        else:
            y2 = int(y2_str)

        if y2 - y1 == 4:
            return True, y1, y2, f"{y1}-{y2}"
        else:
            return False, y1, y2, f"{y1}-{y2}"

    return False, None, None, s_clean


def normalize_nien_khoa(val_nien_khoa, val_ngay_cap_bang=None):
    """
    Chuẩn hóa sơ bộ Niên khóa (Cột O):
    Nếu bóc tách được 2 năm có Y_ra - Y_vao = 4 thì trả về định dạng YYYY-YYYY.
    Nếu không, giữ nguyên chuỗi gốc đã làm sạch.
    """
    is_valid, y1, y2, fmt = parse_nien_khoa(val_nien_khoa)
    if is_valid:
        return fmt
    return clean_whitespace(str(val_nien_khoa)) if val_nien_khoa is not None else ""


# ==============================================================================
# CHUẨN HÓA NGÀY THÁNG
# ==============================================================================
def format_date_to_text(value):
    """
    Chuyển đổi mọi loại dữ liệu ngày tháng thành văn bản thuần túy dạng dd/MM/yyyy.
    Hỗ trợ datetime, date, số serial Excel, số liền 7-8 chữ số (Custom format DDMMYYYY),
    lỗi lặp dấu (//, ///, .., --), chuỗi ISO (YYYY-MM-DD), chuỗi phân cách dấu chấm/gạch ngang, năm sinh.
    """
    if value is None:
        return ""

    if isinstance(value, float):
        if math.isnan(value):
            return ""

    # 1. Hỗ trợ cả datetime.datetime và datetime.date
    if isinstance(value, (datetime.datetime, datetime.date)):
        return value.strftime("%d/%m/%Y")

    # 2. Xử lý số serial ngày tháng của Excel (10000 đến 80000, tương ứng năm 1927 đến 2119)
    if isinstance(value, (int, float)) and 10000 <= value <= 80000:
        try:
            dt = from_excel(value)
            return dt.strftime("%d/%m/%Y")
        except:
            pass

    val_str = str(value).strip()
    if not val_str or val_str.lower() in ("none", "nan", "null"):
        return ""

    # Loại bỏ phần thập phân .0 nếu là số nguyên lưu dạng float (ví dụ: 19061998.0 -> 19061998, 1990.0 -> 1990)
    val_clean = re.sub(r'\.0+$', '', val_str)

    # 3. Xử lý dạng số liền 7-8 chữ số (ví dụ: 19061998 hoặc 1061998 do gõ số dùng Custom format 00/00/0000)
    if val_clean.isdigit() and len(val_clean) in (7, 8):
        padded = val_clean.zfill(8)
        # Thử kiểu DDMMYYYY (phổ biến nhất tại Việt Nam)
        d, m, y = int(padded[:2]), int(padded[2:4]), int(padded[4:])
        if 1900 <= y <= 2100:
            try:
                dt = datetime.date(y, m, d)
                return f"{dt.day:02d}/{dt.month:02d}/{dt.year}"
            except ValueError:
                pass
        # Thử kiểu YYYYMMDD
        y, m, d = int(padded[:4]), int(padded[4:6]), int(padded[6:])
        if 1900 <= y <= 2100:
            try:
                dt = datetime.date(y, m, d)
                return f"{dt.day:02d}/{dt.month:02d}/{dt.year}"
            except ValueError:
                pass

    # 4. Gom các dấu phân cách lặp lại (//, ///, .., --) và khoảng trắng thừa, loại bỏ / ở đầu/cuối
    val_norm = re.sub(r'[/.\-\s]+', '/', val_clean).strip('/')
    if not val_norm or val_norm.lower() in ("none", "nan", "null"):
        return ""

    # 5. Chuỗi dạng ISO YYYY/MM/DD -> đảo về DD/MM/YYYY
    m_iso = re.match(r'^(\d{4})/(\d{1,2})/(\d{1,2})$', val_norm)
    if m_iso:
        y, m, d = m_iso.groups()
        return f"{int(d):02d}/{int(m):02d}/{y}"

    # 6. Chuỗi dạng DD/MM/YYYY hoặc D/M/YYYY -> chuẩn hóa thành DD/MM/YYYY
    m_dmy = re.match(r'^(\d{1,2})/(\d{1,2})/(\d{4})$', val_norm)
    if m_dmy:
        d, m, y = m_dmy.groups()
        return f"{int(d):02d}/{int(m):02d}/{y}"

    # 6b. Chuỗi dạng D/M/YY (năm 2 chữ số)
    m_dmy2 = re.match(r'^(\d{1,2})/(\d{1,2})/(\d{2})$', val_norm)
    if m_dmy2:
        d, m, y2 = m_dmy2.groups()
        y_int = int(y2)
        full_y = 2000 + y_int if y_int <= 35 else 1900 + y_int
        return f"{int(d):02d}/{int(m):02d}/{full_y}"

    # 7. Nếu chỉ có năm (ví dụ "1990" hoặc trường hợp nhập "//1990")
    m_year = re.match(r'^(\d{4})$', val_norm)
    if m_year:
        return m_year.group(1)

    # 8. Các trường hợp chuỗi còn lại
    return val_norm


def is_valid_dmy(val):
    """Kiểm tra chuỗi có đúng chuẩn dd/MM/yyyy hợp lệ thực tế không."""
    if not val or not isinstance(val, str):
        return False
    m = re.match(r'^(\d{2})/(\d{2})/(\d{4})$', val.strip())
    if not m:
        return False
    d, m, y = map(int, m.groups())
    if not (1900 <= y <= 2100):
        return False
    try:
        datetime.date(y, m, d)
        return True
    except ValueError:
        return False


def format_cell(text, width, align='<'):
    """Căn lề và cắt ngắn văn bản hiển thị trong bảng terminal."""
    s = unicodedata.normalize('NFC', str(text)) if text is not None else ""
    if len(s) > width:
        s = s[:width - 1] + "…"
    if align == '^':
        return s.center(width)
    elif align == '>':
        return s.rjust(width)
    return s.ljust(width)


def generate_date_suggestion(field_name, raw_val, norm_val, student_cohort, all_sheet_students):
    """
    Tạo gợi ý thông minh cho ngày sinh hoặc ngày cấp bằng dựa vào dữ liệu học sinh cùng sheet.
    Trả về: (suggestion_date, reason) hoặc (None, None)
    """
    if field_name == "Ngày cấp bằng":
        # Tìm ngày cấp bằng hợp lệ phổ biến nhất trong sheet
        valid_grads = [
            s["row_values"][18] for s in all_sheet_students
            if is_valid_dmy(s["row_values"][18])
        ]
        if valid_grads:
            most_common, count = Counter(valid_grads).most_common(1)[0]
            return most_common, f"Lấy từ ngày cấp bằng chuẩn của {count} bạn cùng lớp/sheet"
        return None, None

    elif field_name == "Ngày sinh":
        # 1. Xác định năm sinh dự kiến từ niên khóa hoặc ngày cấp bằng
        y_ra = None
        if student_cohort:
            m_nk = re.search(r'(\d{4})\s*-\s*(\d{4})', str(student_cohort))
            if m_nk:
                y_ra = int(m_nk.group(2))
        if not y_ra:
            # Tìm niên khóa phổ biến nhất trong sheet
            cohorts = [
                s["row_values"][14] for s in all_sheet_students
                if s["row_values"][14] and re.match(r'^\d{4}-\d{4}$', str(s["row_values"][14]))
            ]
            if cohorts:
                most_common_cohort = Counter(cohorts).most_common(1)[0][0]
                y_ra = int(most_common_cohort.split('-')[1])
        if not y_ra:
            # Tìm ngày cấp bằng phổ biến nhất trong sheet
            grads = [
                s["row_values"][18] for s in all_sheet_students
                if is_valid_dmy(s["row_values"][18])
            ]
            if grads:
                most_common_grad = Counter(grads).most_common(1)[0][0]
                y_ra = int(most_common_grad.split('/')[2])

        y_expected = (y_ra - 15) if y_ra else None

        raw_str = str(raw_val).strip() if raw_val is not None else ""
        norm_str = str(norm_val).strip() if norm_val is not None else ""

        # Tìm mẫu dd/mm hoặc dd-mm trong raw_str hoặc norm_str
        m_dm = re.search(r'\b(\d{1,2})[/-](\d{1,2})\b', raw_str)
        if not m_dm:
            m_dm = re.search(r'\b(\d{1,2})[/-](\d{1,2})\b', norm_str)

        if m_dm and y_expected:
            d_val, m_val = int(m_dm.group(1)), int(m_dm.group(2))
            if 1 <= d_val <= 31 and 1 <= m_val <= 12:
                candidate = f"{d_val:02d}/{m_val:02d}/{y_expected}"
                if is_valid_dmy(candidate):
                    return candidate, f"Ghép ngày/tháng ({d_val:02d}/{m_val:02d}) với năm sinh dự kiến {y_expected} (theo niên khóa kết thúc năm {y_ra})"

        return None, None


def prompt_fix_date_with_context(
    file_name, sheet_name, row_idx, student_name, stt_val,
    field_name, raw_val, current_norm_val, student_cohort,
    adjacent_records, smart_suggestion, suggestion_reason,
    date_cache
):
    """
    Hiển thị thông báo, bảng đối chiếu +-2 học sinh liền kề (kèm Niên khóa) và gợi ý thông minh
    cho phép người dùng nhập tay hoặc bấm 'y' để chấp nhận gợi ý.
    Trả về: (new_val, skip_all)
    """
    # Chỉ tự động áp dụng cache cho Ngày cấp bằng (vì học sinh cùng lớp tốt nghiệp cùng ngày)
    if field_name == "Ngày cấp bằng":
        raw_key = (field_name, clean_whitespace(str(raw_val)))
        if raw_key in date_cache:
            cached_val = date_cache[raw_key]
            print(f"   ⚡ Tự động áp dụng Ngày cấp bằng đã nhớ cho '{raw_val}': {cached_val}")
            return cached_val, False

    raw_display = str(raw_val).strip() if raw_val is not None else "(Trống)"
    curr_display = str(current_norm_val).strip() if current_norm_val else "(Chưa xác định)"

    print("\n" + "=" * 95)
    print("⚠️  PHÁT HIỆN NGÀY THÁNG CHƯA ĐÚNG ĐỊNH DẠNG dd/MM/yyyy")
    print(f"   • Tệp nguồn : {file_name} (Sheet: '{sheet_name}', Dòng Excel: {row_idx})")
    print(f"   • Học sinh  : {student_name} (STT: {stt_val})")
    print(f"   • Trường sai: [ {field_name} ]")
    print(f"   • Giá trị gốc trong file Excel : '{raw_display}'")
    print(f"   • Kết quả tự động chuẩn hóa   : '{curr_display}'")
    print("-" * 95)

    if adjacent_records:
        print("📋 BẢNG ĐỐI CHIẾU HỌC SINH LIỀN KỀ:")
        print("┌───┬──────┬────────────────────────┬─────────────┬──────────────┬────────────────┬─────────────────────┐")
        print(f"│   │{format_cell('STT', 6, '^')}│{format_cell(' Họ và tên', 24)}│{format_cell('Niên khóa', 13, '^')}│{format_cell('Ngày sinh', 14, '^')}│{format_cell('Ngày cấp bằng', 16, '^')}│{format_cell(' Ghi chú', 21)}│")
        print("├───┼──────┼────────────────────────┼─────────────┼──────────────┼────────────────┼─────────────────────┤")

        for rec in adjacent_records:
            is_curr = rec.get("is_current", False)
            pointer = " 👉" if is_curr else "   "
            stt_str = str(rec.get("stt", ""))
            name_str = " " + str(rec.get("name", ""))
            cohort_str = str(rec.get("cohort", "")) if rec.get("cohort") else "(Trống)"
            dob_str = str(rec.get("dob", "")) if rec.get("dob") else "(Trống)"
            grad_str = str(rec.get("grad", "")) if rec.get("grad") else "(Trống)"

            note_str = ""
            if is_curr:
                note_str = " <-- LỖI NGÀY SINH" if field_name == "Ngày sinh" else " <-- LỖI NGÀY CẤP"

            c_stt = format_cell(stt_str, 6, '^')
            c_name = format_cell(name_str, 24, '<')
            c_cohort = format_cell(cohort_str, 13, '^')
            c_dob = format_cell(dob_str, 14, '^')
            c_grad = format_cell(grad_str, 16, '^')
            c_note = format_cell(note_str, 21, '<')
            print(f"│{pointer}│{c_stt}│{c_name}│{c_cohort}│{c_dob}│{c_grad}│{c_note}│")

        print("└───┴──────┴────────────────────────┴─────────────┴──────────────┴────────────────┴─────────────────────┘")

    if smart_suggestion:
        reason_info = f" ({suggestion_reason})" if suggestion_reason else ""
        print(f"\n💡 GỢI Ý THÔNG MINH: {smart_suggestion}{reason_info}")
        print(f"👉 [Bấm 'y' hoặc 'Y' để chấp nhận gợi ý: {smart_suggestion}]")
        prompt_text = (
            "   Hoặc nhập ngày chuẩn dd/MM/yyyy (hoặc gõ ddmmyyyy)\n"
            "   [Bấm Enter để giữ nguyên | Gõ 'all' để bỏ qua tất cả các lỗi sau]: "
        )
    else:
        print()
        prompt_text = (
            "👉 Nhập ngày chuẩn dd/MM/yyyy (hoặc gõ ddmmyyyy)\n"
            "   [Bấm Enter để giữ nguyên | Gõ 'all' để bỏ qua tất cả các lỗi sau]: "
        )

    while True:
        try:
            user_input = input(prompt_text).strip()
        except (KeyboardInterrupt, EOFError):
            print("\n   -> Bỏ qua nhập tay.")
            return current_norm_val, True

        if not user_input:
            # Bấm Enter để giữ nguyên
            if field_name == "Ngày cấp bằng":
                raw_key = (field_name, clean_whitespace(str(raw_val)))
                date_cache[raw_key] = current_norm_val
            return current_norm_val, False

        if user_input.lower() in ["all", "skip", "bo qua", "bỏ qua"]:
            print("   -> Đã chọn bỏ qua tất cả cảnh báo nhập tay còn lại.")
            return current_norm_val, True

        if smart_suggestion and user_input.lower() in ["y", "yes", "c", "co", "có", "ok"]:
            print(f"   ✅ Đã chấp nhận gợi ý cho {field_name}: {smart_suggestion}")
            if field_name == "Ngày cấp bằng":
                raw_key = (field_name, clean_whitespace(str(raw_val)))
                date_cache[raw_key] = smart_suggestion
            return smart_suggestion, False

        # Thử chuẩn hóa chuỗi người dùng nhập (ví dụ gõ 5/6/2005 hoặc 05062005)
        formatted = format_date_to_text(user_input)
        if is_valid_dmy(formatted):
            print(f"   ✅ Đã cập nhật {field_name}: {formatted}")
            if field_name == "Ngày cấp bằng":
                raw_key = (field_name, clean_whitespace(str(raw_val)))
                date_cache[raw_key] = formatted
            return formatted, False
        else:
            print(f"   ⚠️ Giá trị '{user_input}' chưa đúng dạng dd/MM/yyyy hợp lệ. Vui lòng nhập lại (hoặc bấm Enter để bỏ qua):")


def prompt_fix_date(file_name, sheet_name, row_idx, student_name, stt_val, field_name, raw_val, current_norm_val, date_cache):
    """Hàm bọc tương thích ngược."""
    return prompt_fix_date_with_context(
        file_name, sheet_name, row_idx, student_name, stt_val,
        field_name, raw_val, current_norm_val, "", [], None, None, date_cache
    )


def prompt_fix_nien_khoa_with_context(
    file_name, sheet_name, row_idx, student_name, stt_val,
    raw_val, current_norm_val, adjacent_records, nien_khoa_cache
):
    """
    Hiển thị thông báo, bảng đối chiếu +-2 học sinh liền kề (Ngày sinh, Niên khóa, Ngày cấp bằng)
    cho phép người dùng nhập tay niên khóa. Tự động ghi nhớ để áp dụng cho các dòng tương tự.
    Trả về: (new_val, skip_all)
    """
    cache_key = clean_whitespace(str(raw_val)) if raw_val is not None else ""
    if cache_key in nien_khoa_cache:
        cached_val = nien_khoa_cache[cache_key]
        print(f"   ⚡ Tự động áp dụng Niên khóa đã nhớ cho '{raw_val if raw_val is not None else '(Trống)'}': {cached_val}")
        return cached_val, False

    raw_display = str(raw_val).strip() if raw_val is not None else "(Trống)"
    curr_display = str(current_norm_val).strip() if current_norm_val else "(Trống)"

    print("\n" + "=" * 95)
    print("⚠️  PHÁT HIỆN NIÊN KHÓA CHƯA ĐÚNG QUY CHUẨN (Y_ra - Y_vao khác 4)")
    print(f"   • Tệp nguồn : {file_name} (Sheet: '{sheet_name}', Dòng Excel: {row_idx})")
    print(f"   • Học sinh  : {student_name} (STT: {stt_val})")
    print(f"   • Trường sai: [ Niên khóa ]")
    print(f"   • Giá trị gốc trong file Excel : '{raw_display}'")
    print(f"   • Kết quả hiện tại             : '{curr_display}'")
    print("-" * 95)

    if adjacent_records:
        print("📋 BẢNG ĐỐI CHIẾU HỌC SINH LIỀN KỀ:")
        print("┌───┬──────┬────────────────────────┬──────────────┬─────────────┬────────────────┬─────────────────────────┐")
        print(f"│   │{format_cell('STT', 6, '^')}│{format_cell(' Họ và tên', 24)}│{format_cell('Ngày sinh', 14, '^')}│{format_cell('Niên khóa', 13, '^')}│{format_cell('Ngày cấp bằng', 16, '^')}│{format_cell(' Ghi chú', 25)}│")
        print("├───┼──────┼────────────────────────┼──────────────┼─────────────┼────────────────┼─────────────────────────┤")

        for rec in adjacent_records:
            is_curr = rec.get("is_current", False)
            pointer = " 👉" if is_curr else "   "
            stt_str = str(rec.get("stt", ""))
            name_str = " " + str(rec.get("name", ""))
            dob_str = str(rec.get("dob", "")) if rec.get("dob") else "(Trống)"
            cohort_str = str(rec.get("cohort", "")) if rec.get("cohort") else "(Trống)"
            grad_str = str(rec.get("grad", "")) if rec.get("grad") else "(Trống)"

            note_str = " <-- LỖI NIÊN KHÓA" if is_curr else ""

            c_stt = format_cell(stt_str, 6, '^')
            c_name = format_cell(name_str, 24, '<')
            c_dob = format_cell(dob_str, 14, '^')
            c_cohort = format_cell(cohort_str, 13, '^')
            c_grad = format_cell(grad_str, 16, '^')
            c_note = format_cell(note_str, 25, '<')
            print(f"│{pointer}│{c_stt}│{c_name}│{c_dob}│{c_cohort}│{c_grad}│{c_note}│")

        print("└───┴──────┴────────────────────────┴──────────────┴─────────────┴────────────────┴─────────────────────────┘")

    prompt_text = (
        "\n👉 Nhập niên khóa chuẩn YYYY-YYYY (ví dụ: 2016-2020)\n"
        "   [Bấm Enter để giữ nguyên | Gõ 'all' để bỏ qua tất cả các lỗi sau]: "
    )

    while True:
        try:
            user_input = input(prompt_text).strip()
        except (KeyboardInterrupt, EOFError):
            print("\n   -> Bỏ qua nhập tay.")
            return current_norm_val, True

        if not user_input:
            # Bấm Enter để giữ nguyên
            nien_khoa_cache[cache_key] = current_norm_val
            return current_norm_val, False

        if user_input.lower() in ["all", "skip", "bo qua", "bỏ qua"]:
            print("   -> Đã chọn bỏ qua tất cả cảnh báo nhập tay còn lại.")
            return current_norm_val, True

        is_valid, y1, y2, fmt = parse_nien_khoa(user_input)
        if is_valid:
            print(f"   ✅ Đã cập nhật Niên khóa: {fmt} (đã ghi nhớ cho các dòng tương tự)")
            nien_khoa_cache[cache_key] = fmt
            return fmt, False
        else:
            if y1 is not None and y2 is not None:
                print(f"   ⚠️ Giá trị '{user_input}' có hiệu số {y2} - {y1} = {y2 - y1} (khác 4). Vui lòng nhập lại (hoặc bấm Enter để giữ nguyên):")
            else:
                print(f"   ⚠️ Giá trị '{user_input}' chưa đúng dạng YYYY-YYYY (ví dụ: 2016-2020). Vui lòng nhập lại (hoặc bấm Enter để giữ nguyên):")


# ==============================================================================
# CÔNG ĐOẠN 1: CHUẨN HÓA TIÊU ĐỀ & DỮ LIỆU TỪNG FILE LẺ
# ==============================================================================
def stage1_normalize_individual_files():
    """
    Quét từng file nguồn (thứ tự ABC), chuẩn hóa tiêu đề và dữ liệu về dạng 28 cột.
    Chuẩn hóa ngày sinh và ngày cấp bằng (hiển thị giá trị gốc và dấu nhắc nhập tay nếu ngày chưa đúng chuẩn dd/MM/yyyy).
    Tự động tô màu vàng vào các ô bị thiếu dữ liệu.
    Xuất file lẻ cho từng file nguồn vào thư mục 'Da_Chuan_Hoa'.
    """
    print("\n" + "=" * 80)
    print(" CÔNG ĐOẠN 1: CHUẨN HÓA HÀNG TIÊU ĐỀ & DỮ LIỆU TỪNG FILE NGUỒN ".center(80, "="))
    print("=" * 80)

    print("Đang mở cửa sổ chọn thư mục nguồn...")
    folder_path = get_folder_path("Chọn thư mục chứa các file Excel nguồn cần chuẩn hóa:")
    if not folder_path:
        print("Bạn đã hủy chọn thư mục.")
        return

    folder = Path(folder_path)
    print(f"Thư mục nguồn đã chọn: {folder.absolute()}")

    print("\n🏢 THÔNG TIN ĐƠN VỊ CẤP BẰNG (Áp dụng cho toàn bộ đợt này):")
    unit_code_input = input("👉 Nhập Mã đơn vị (Cột AA - Bấm Enter để giữ nguyên file nguồn): ").strip().strip("'\"")
    unit_name_input = input("👉 Nhập Tên đơn vị (Cột AB - Bấm Enter để giữ nguyên file nguồn): ").strip().strip("'\"")

    if not TEMPLATE_PATH.exists():
        print(f"❌ Lỗi: Không tìm thấy file mẫu '{TEMPLATE_PATH.name}' tại: {TEMPLATE_PATH.parent}")
        return

    try:
        with open(TEMPLATE_PATH, "rb") as f:
            template_bytes = f.read()
        wb_tpl = openpyxl.load_workbook(io.BytesIO(template_bytes))
        ws_tpl = wb_tpl.active
        cached_fonts = get_cached_row_fonts(ws_tpl, ws_tpl.max_row, max_col=28)
        wb_tpl.close()
    except Exception as e:
        print(f"❌ Lỗi khi đọc file mẫu template.xlsx: {e}")
        return

    # Thư mục xuất file đã chuẩn hóa
    output_dir = folder / "Da_Chuan_Hoa"
    output_dir.mkdir(parents=True, exist_ok=True)

    yellow_fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")

    # Lấy danh sách file Excel nguồn, sắp xếp theo thứ tự ABC
    excel_files = sorted(
        [
            f for f in folder.glob("*.xlsx")
            if not f.name.startswith("~$")
            and not f.name.startswith("._")
            and not f.name.endswith("_TongHop.xlsx")
            and not f.name.endswith("_Xong.xlsx")
            and f.is_file()
        ],
        key=lambda x: x.name.lower()
    )

    if not excel_files:
        print(f"⚠️ Không tìm thấy file Excel (*.xlsx) nào trong thư mục: {folder.name}")
        return

    print(f"\n📁 Tìm thấy tổng cộng {len(excel_files)} file Excel cần chuẩn hóa (thứ tự quét: ABC):")
    for idx, f in enumerate(excel_files, start=1):
        print(f"   {idx:>2}. {f.name}")

    processed_files_count = 0
    total_rows_extracted = 0
    skipped_sheets_summary = []
    skip_all_date_prompt = False
    skip_all_nien_khoa_prompt = False
    date_cache = {}

    for file_idx, file_path in enumerate(excel_files, start=1):
        print("\n" + "=" * 80)
        print(f"[{file_idx}/{len(excel_files)}] ĐANG XỬ LÝ FILE: {file_path.name}")
        print("=" * 80)

        wb_in = None
        wb_out = None
        nien_khoa_cache = {}
        try:
            wb_in = openpyxl.load_workbook(file_path, data_only=True)
            file_rows = []

            for ws_in in wb_in.worksheets:
                if ws_in.sheet_state != 'visible':
                    continue

                format_type, headers, header_row_idx = get_file_format_info(ws_in)

                if format_type == "INVALID":
                    print(f"  -> Bỏ qua Sheet '{ws_in.title}' (Không khớp cấu trúc 28 hoặc 31 cột)")
                    headers_str = ", ".join([h for h in headers if h]) if headers else "Không có dòng tiêu đề"
                    skipped_sheets_summary.append({
                        "file": file_path.name,
                        "sheet": ws_in.title,
                        "columns": headers_str
                    })
                    continue

                remove_qrs = (format_type == "AE")
                print(f"  -> Quét Sheet '{ws_in.title}': {'Dạng AE 31 cột (loại bỏ 3 cột thừa Q, R, S)' if remove_qrs else 'Dạng AB chuẩn 28 cột'}")

                start_extracting = False
                empty_consecutive_rows = 0
                sheet_students = []

                for row in ws_in.iter_rows():
                    row_idx = row[0].row if hasattr(row[0], 'row') else 0

                    # Bỏ qua tuyệt đối các dòng trước và chính dòng tiêu đề
                    if header_row_idx is not None and row_idx <= header_row_idx:
                        continue

                    cell_a_val = row[0].value
                    cell_a_str = str(cell_a_val).strip() if cell_a_val is not None else ""

                    name_cell_val = row[4].value if len(row) > 4 else None
                    name_cell_str = clean_whitespace(str(name_cell_val)) if name_cell_val is not None else ""
                    has_letters_in_name = bool(re.search(r'[a-zA-ZÀ-ỹ]', name_cell_str))

                    is_row_empty = not any(c.value is not None and str(c.value).strip() != "" for c in row[:15])

                    if not start_extracting:
                        is_stt_one = bool(re.match(r'^0*1(\.0*)?$', cell_a_str))
                        if (is_stt_one and has_letters_in_name) or (has_letters_in_name and name_cell_str.lower() not in ["họ và tên", "họ tên", "họ, chữ đệm và tên"]):
                            start_extracting = True

                    if start_extracting:
                        if is_row_empty:
                            empty_consecutive_rows += 1
                            if empty_consecutive_rows >= 20:
                                break
                            continue
                        else:
                            empty_consecutive_rows = 0

                        # Bỏ qua nếu là dòng tổng kết, ghi chú cuối bảng (ví dụ: 'Danh sách này có 100 học sinh')
                        if is_summary_or_footer_row(row):
                            continue

                        # Chỉ lấy các dòng dữ liệu với ĐỒNG THỜI STT là số hợp lệ và Họ tên không trống
                        if not is_valid_stt(cell_a_val) or is_cell_empty(name_cell_val) or not has_letters_in_name:
                            continue

                        row_values = [cell.value for cell in row]

                        required_len = 31 if remove_qrs else 28
                        while len(row_values) < required_len:
                            row_values.append(None)

                        clean_stt = int(cell_a_val) if isinstance(cell_a_val, (int, float)) and float(cell_a_val).is_integer() else str(cell_a_val or '')

                        # Lấy giá trị ngày gốc & niên khóa gốc để đối chiếu khi cần nhập tay
                        raw_dob = row[8].value if len(row) > 8 else None
                        raw_nien_khoa = row[14].value if len(row) > 14 else None
                        if remove_qrs:
                            raw_grad = row[21].value if len(row) > 21 else None
                            new_row_values = (row_values[:16] + row_values[19:])[:28]
                        else:
                            raw_grad = row[18].value if len(row) > 18 else None
                            new_row_values = row_values[:28]

                        while len(new_row_values) < 28:
                            new_row_values.append(None)

                        new_row_values = [clean_whitespace(val) for val in new_row_values]

                        # Chuẩn hóa sơ bộ định dạng ngày tháng và niên khóa
                        new_row_values[8] = format_date_to_text(new_row_values[8])
                        new_row_values[18] = format_date_to_text(new_row_values[18])
                        new_row_values[14] = normalize_nien_khoa(new_row_values[14])

                        sheet_students.append({
                            "row_idx": row_idx,
                            "clean_stt": clean_stt,
                            "name": name_cell_str,
                            "raw_dob": raw_dob,
                            "raw_grad": raw_grad,
                            "raw_nien_khoa": raw_nien_khoa,
                            "row_values": new_row_values
                        })

                # Công đoạn 1.2: Rà soát & sửa ngày tháng, niên khóa tương tác với bảng đối chiếu +-2 học sinh liền kề
                for i, student in enumerate(sheet_students):
                    row_idx = student["row_idx"]
                    name_cell_str = student["name"]
                    clean_stt = student["clean_stt"]
                    raw_dob = student["raw_dob"]
                    raw_grad = student["raw_grad"]
                    raw_nien_khoa = student["raw_nien_khoa"]
                    row_values = student["row_values"]

                    # Hàm lấy danh sách +-2 học sinh liền kề (kèm Niên khóa, Ngày sinh, Ngày cấp bằng)
                    def get_adjacent_records(curr_idx):
                        adj = []
                        start_k = max(0, curr_idx - 2)
                        end_k = min(len(sheet_students), curr_idx + 3)
                        for k in range(start_k, end_k):
                            s_item = sheet_students[k]
                            adj.append({
                                "stt": s_item["clean_stt"],
                                "name": s_item["name"],
                                "cohort": s_item["row_values"][14] or "",
                                "dob": s_item["row_values"][8] or "",
                                "grad": s_item["row_values"][18] or "",
                                "is_current": (k == curr_idx)
                            })
                        return adj

                    # 1. Rà soát Ngày cấp bằng (Cột S - Index 18) TRƯỚC
                    if not skip_all_date_prompt and not is_valid_dmy(row_values[18]) and not is_cell_empty(raw_grad):
                        adj_recs = get_adjacent_records(i)
                        sugg_date, sugg_reason = generate_date_suggestion(
                            "Ngày cấp bằng", raw_grad, row_values[18], row_values[14], sheet_students
                        )
                        row_values[18], skip_all_date_prompt = prompt_fix_date_with_context(
                            file_path.name, ws_in.title, row_idx,
                            name_cell_str, clean_stt,
                            "Ngày cấp bằng", raw_grad, row_values[18], row_values[14],
                            adj_recs, sugg_date, sugg_reason, date_cache
                        )

                    # 2. Rà soát Ngày sinh (Cột I - Index 8)
                    if not skip_all_date_prompt and not is_valid_dmy(row_values[8]) and not is_cell_empty(raw_dob):
                        adj_recs = get_adjacent_records(i)
                        sugg_date, sugg_reason = generate_date_suggestion(
                            "Ngày sinh", raw_dob, row_values[8], row_values[14], sheet_students
                        )
                        row_values[8], skip_all_date_prompt = prompt_fix_date_with_context(
                            file_path.name, ws_in.title, row_idx,
                            name_cell_str, clean_stt,
                            "Ngày sinh", raw_dob, row_values[8], row_values[14],
                            adj_recs, sugg_date, sugg_reason, date_cache
                        )

                    # 3. Rà soát & chuẩn hóa Niên khóa (Cột O - Index 14)
                    is_valid_nk, y1, y2, fmt_nk = parse_nien_khoa(row_values[14])
                    if is_valid_nk:
                        row_values[14] = fmt_nk
                    else:
                        if not skip_all_nien_khoa_prompt:
                            adj_recs = get_adjacent_records(i)
                            row_values[14], skip_all_nien_khoa_prompt = prompt_fix_nien_khoa_with_context(
                                file_path.name, ws_in.title, row_idx,
                                name_cell_str, clean_stt,
                                raw_nien_khoa, row_values[14],
                                adj_recs, nien_khoa_cache
                            )

                    # Áp dụng các trường cố định và mã danh mục
                    # Cột D (Index 3): Tên văn bằng cố định
                    row_values[3] = "Bằng tốt nghiệp trung học cơ sở"
                    # Cột Y (Index 24): Trạng thái số hóa cố định
                    row_values[24] = "Đã số hóa đầy đủ, chính xác"
                    # Cột J (Index 9): Giới tính (1->Nam, 2->Khác, 3->Nữ)
                    row_values[9] = map_code_to_text(9, row_values[9])
                    # Cột K (Index 10): Dân tộc (1..51, 29->Kinh)
                    row_values[10] = map_code_to_text(10, row_values[10])
                    # Cột L (Index 11): Quốc tịch (1..159, 83/89->Việt Nam)
                    row_values[11] = map_code_to_text(11, row_values[11])
                    # Cột M (Index 12): Nơi sinh (1..75 tỉnh/thành)
                    row_values[12] = map_code_to_text(12, row_values[12])
                    # Cột AA (Index 26): Mã đơn vị
                    if unit_code_input:
                        row_values[26] = unit_code_input
                    # Cột AB (Index 27): Tên đơn vị
                    if unit_name_input:
                        row_values[27] = unit_name_input

                    file_rows.append(row_values)

            M = len(file_rows)
            if M == 0:
                print(f"⚠️ File '{file_path.name}' không có dòng dữ liệu học sinh nào. Bỏ qua.")
                continue

            print(f"  -> Trích xuất được: {M} dòng dữ liệu học sinh.")

            # ==================================================================
            # XUẤT FILE EXCEL LẺ ĐÃ CHUẨN HÓA VÀO THƯ MỤC 'Da_Chuan_Hoa'
            # ==================================================================
            wb_out = openpyxl.load_workbook(io.BytesIO(template_bytes))
            ws_out = wb_out.active
            ws_out.title = "Data"

            yellow_cell_count = 0
            for r in range(M):
                row_vals = file_rows[r]
                row_vals[0] = r + 1  # Đánh lại STT 1..M cho file lẻ
                ws_out.append(row_vals)
                curr_row = ws_out.max_row
                apply_cached_fonts(ws_out, curr_row, cached_fonts)

                for c in range(1, 28):
                    if is_cell_empty(row_vals[c]):
                        ws_out.cell(row=curr_row, column=c + 1).fill = yellow_fill
                        yellow_cell_count += 1

            out_file_path = output_dir / file_path.name
            try:
                wb_out.save(out_file_path)
            except PermissionError:
                print(f"❌ Lỗi: Không thể ghi file '{out_file_path.name}'. Vui lòng đóng file nếu đang mở trong Excel và thử lại.")
                continue

            processed_files_count += 1
            total_rows_extracted += M
            print(f"  -> Xuất file: Da_Chuan_Hoa/{out_file_path.name} ({M} học sinh | {yellow_cell_count} ô thiếu tô vàng)")

        except Exception as e:
            print(f"❌ Lỗi khi xử lý file {file_path.name}: {e}")
        finally:
            if wb_in is not None:
                try:
                    wb_in.close()
                except Exception:
                    pass
            if wb_out is not None:
                try:
                    wb_out.close()
                except Exception:
                    pass

    # Báo cáo tổng kết Công đoạn 1
    print("\n" + "=" * 80)
    print(" HOÀN THÀNH CÔNG ĐOẠN 1 ".center(80, "="))
    print(f"✅ Đã chuẩn hóa thành công: {processed_files_count}/{len(excel_files)} file Excel nguồn.")
    print(f"✅ Tổng số học sinh: {total_rows_extracted} dòng dữ liệu.")
    print(f"📁 Toàn bộ các file chuẩn hóa lẻ đã được lưu tại thư mục:")
    print(f"   {output_dir.absolute()}")
    print("=" * 80)

    if skipped_sheets_summary:
        print("\n" + "=" * 100)
        print(" DANH SÁCH CÁC SHEET BỊ BỎ QUA (SAI CẤU TRÚC)".center(100))
        print("=" * 100)
        print(f"| {'Tên File':<30} | {'Tên Sheet':<25} | {'Các cột hiện có'}")
        print("-" * 100)
        for info in skipped_sheets_summary:
            f_name = info['file'][:27] + "..." if len(info['file']) > 30 else info['file']
            s_name = info['sheet'][:22] + "..." if len(info['sheet']) > 25 else info['sheet']
            print(f"| {f_name:<30} | {s_name:<25} | {info['columns']}")
        print("=" * 100)


# ==============================================================================
# CÔNG ĐOẠN 2: QUÉT CÁC FILE ĐÃ CHUẨN HÓA & TẠO FILE EXCEL TỔNG HỢP
# ==============================================================================
def stage2_merge_normalized_files():
    """
    Quét các file đã chuẩn hóa (thứ tự ABC), gộp toàn bộ thành 1 file Excel tổng hợp duy nhất.
    Đánh lại số thứ tự (STT) liên tục 1..N.
    Giữ nguyên định dạng và màu tô vàng của các ô bị thiếu.
    KHÔNG cần chuẩn hóa lại ngày sinh và ngày cấp bằng.
    """
    print("\n" + "=" * 80)
    print(" CÔNG ĐOẠN 2: QUÉT CÁC FILE ĐÃ CHUẨN HÓA & TẠO FILE TỔNG HỢP ".center(80, "="))
    print("=" * 80)

    print("Đang mở cửa sổ chọn thư mục chứa file đã chuẩn hóa...")
    folder_path = get_folder_path("Chọn thư mục chứa file đã chuẩn hóa (thư mục Da_Chuan_Hoa hoặc thư mục nguồn):")
    if not folder_path:
        print("Bạn đã hủy chọn thư mục.")
        return

    folder = Path(folder_path)

    # Tự động nhận diện nếu người dùng chọn thư mục cha có chứa thư mục con 'Da_Chuan_Hoa'
    if folder.name != "Da_Chuan_Hoa" and (folder / "Da_Chuan_Hoa").is_dir():
        sub_files = [f for f in (folder / "Da_Chuan_Hoa").glob("*.xlsx") if not f.name.startswith("~$") and not f.name.startswith("._")]
        if sub_files:
            folder = folder / "Da_Chuan_Hoa"
            print(f"-> Tự động nhận diện và sử dụng thư mục con: {folder.name}/")

    print(f"Thư mục làm việc: {folder.absolute()}")

    if not TEMPLATE_PATH.exists():
        print(f"❌ Lỗi: Không tìm thấy file mẫu '{TEMPLATE_PATH.name}' tại: {TEMPLATE_PATH.parent}")
        return

    try:
        with open(TEMPLATE_PATH, "rb") as f:
            template_bytes = f.read()
        wb_out = openpyxl.load_workbook(io.BytesIO(template_bytes))
        ws_out = wb_out.active
        ws_out.title = "Data"
        cached_fonts = get_cached_row_fonts(ws_out, ws_out.max_row, max_col=28)
    except Exception as e:
        print(f"❌ Lỗi khi đọc file mẫu template.xlsx: {e}")
        return

    yellow_fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")

    # Lấy danh sách file chuẩn hóa, sắp xếp ABC
    excel_files = sorted(
        [
            f for f in folder.glob("*.xlsx")
            if not f.name.startswith("~$")
            and not f.name.startswith("._")
            and not f.name.endswith("_TongHop.xlsx")
            and not f.name.endswith("_Xong.xlsx")
            and f.is_file()
        ],
        key=lambda x: x.name.lower()
    )

    if not excel_files:
        print(f"⚠️ Không tìm thấy file Excel (*.xlsx) nào trong thư mục: {folder.name}")
        return

    print(f"\n📁 Tìm thấy {len(excel_files)} file đã chuẩn hóa cần gộp (thứ tự quét: ABC):")

    count_files = 0
    total_student_rows = 0

    for file_idx, file_path in enumerate(excel_files, start=1):
        print(f"  [{file_idx}/{len(excel_files)}] Đang đọc: {file_path.name} ...", end=" ")
        wb_in = None
        try:
            wb_in = openpyxl.load_workbook(file_path, data_only=True)
            ws_in = wb_in['Data'] if 'Data' in wb_in.sheetnames else wb_in.active

            # Tìm dòng tiêu đề
            header_row_idx = 2
            for r in range(1, min(10, ws_in.max_row + 1)):
                row_vals = [clean_whitespace(str(ws_in.cell(row=r, column=c).value or '')) for c in range(1, 10)]
                row_lower = [v.lower() for v in row_vals]
                if any(s in ["stt", "số tt", "số thứ tự"] for s in row_lower) or any("họ" in s and "tên" in s for s in row_lower):
                    header_row_idx = r
                    break

            file_student_count = 0
            for r in range(header_row_idx + 1, ws_in.max_row + 1):
                cells = [ws_in.cell(row=r, column=c) for c in range(1, 29)]
                # Bỏ qua dòng trống hoàn toàn hoặc thiếu STT / Họ tên / dòng tổng kết
                if not any(c.value is not None and str(c.value).strip() != "" for c in cells[:15]):
                    continue
                if is_summary_or_footer_row(cells) or not is_valid_stt(cells[0].value) or is_cell_empty(cells[4].value):
                    continue

                row_vals = [clean_whitespace(c.value) for c in cells]

                # Nhận diện trạng thái tô vàng từ ô trong file chuẩn hóa
                row_yellow = []
                for c in cells:
                    is_yel = is_cell_yellow(c) or is_cell_empty(c.value)
                    row_yellow.append(is_yel)

                total_student_rows += 1
                file_student_count += 1
                row_vals[0] = total_student_rows  # Đánh lại STT 1..N liên tục toàn bộ file tổng

                ws_out.append(row_vals)
                curr_row = ws_out.max_row
                apply_cached_fonts(ws_out, curr_row, cached_fonts)

                for col_idx, is_yel in enumerate(row_yellow, start=1):
                    if is_yel and col_idx > 1:
                        ws_out.cell(row=curr_row, column=col_idx).fill = yellow_fill

            count_files += 1
            print(f"-> Gộp {file_student_count:>4} học sinh.")

        except Exception as e:
            print(f"❌ Lỗi khi đọc file {file_path.name}: {e}")
        finally:
            if wb_in is not None:
                try:
                    wb_in.close()
                except Exception:
                    pass

    # Đường dẫn lưu file tổng hợp
    save_dir = folder.parent if folder.name == "Da_Chuan_Hoa" else folder
    out_path = save_dir / f"{save_dir.name}_TongHop.xlsx"

    try:
        wb_out.save(out_path)
        print("\n" + "=" * 80)
        print(" HOÀN THÀNH CÔNG ĐOẠN 2 ".center(80, "="))
        print(f"✅ Đã gộp thành công: {count_files}/{len(excel_files)} file Excel đã chuẩn hóa.")
        print(f"✅ Tổng số học sinh: {total_student_rows} dòng dữ liệu.")
        print(f"📁 File TỔNG HỢP đã lưu tại:")
        print(f"   {out_path.absolute()}")
        print("=" * 80)
    except PermissionError:
        print(f"\n❌ Lỗi: Không thể lưu file '{out_path.name}' do file đang được mở trong Excel.")
        print("👉 Vui lòng ĐÓNG file trong Excel rồi chạy lại để lưu file tổng hợp.")
    except Exception as e:
        print(f"❌ Lỗi khi lưu file tổng hợp: {e}")
    finally:
        try:
            wb_out.close()
        except Exception:
            pass


# ==============================================================================
# CÔNG ĐOẠN 3: LỌC & GHÉP DỮ LIỆU CHƯA NHẬP TỪ FILE NHIỀU SHEET
# ==============================================================================
def stage3_filter_unimported_records():
    """
    Xử lý một file Excel chứa nhiều sheet có cùng cấu trúc:
    - Hàng tiêu đề có cấu trúc chuẩn với cột Tên đơn vị là AB, cột Kết quả nhập là AC.
    - Trích ra các dòng dữ liệu (giữ nguyên STT ban đầu và dữ liệu gốc).
    - Với mỗi dòng có Kết quả nhập khác "Đã nhập": tô màu vàng toàn bộ dòng đó.
    - Giữ lại kèm thêm 5 dòng liền trước và 5 dòng liền sau nó để làm tham chiếu sửa dữ liệu.
    - Ghép toàn bộ vào một sheet duy nhất và xuất ra file Excel có tên "KetQuaNhap.xlsx" trong cùng thư mục với file nguồn.
    """
    print("\n" + "=" * 80)
    print(" CÔNG ĐOẠN 3: LỌC & GHÉP DỮ LIỆU CHƯA NHẬP TỪ FILE NHIỀU SHEET ".center(80, "="))
    print("=" * 80)

    print("Đang mở cửa sổ chọn file Excel nguồn...")
    file_path_str = get_file_path("Chọn file Excel chứa các sheet kết quả nhập:")
    if not file_path_str:
        print("Bạn đã hủy chọn file.")
        return

    file_path = Path(file_path_str)
    if not file_path.is_file():
        print(f"❌ Lỗi: File không tồn tại: {file_path}")
        return

    print(f"File nguồn đã chọn: {file_path.absolute()}")

    if not TEMPLATE_PATH.exists():
        print(f"❌ Lỗi: Không tìm thấy file mẫu '{TEMPLATE_PATH.name}' tại: {TEMPLATE_PATH.parent}")
        return

    try:
        with open(TEMPLATE_PATH, "rb") as f:
            template_bytes = f.read()
        wb_tpl = openpyxl.load_workbook(io.BytesIO(template_bytes))
        ws_tpl = wb_tpl.active
        cached_fonts = get_cached_row_fonts(ws_tpl, ws_tpl.max_row, max_col=28)
        # Thêm font cho cột AC (cột thứ 29)
        cached_fonts.append(Font(name=cached_fonts[0].name, size=cached_fonts[0].size) if cached_fonts and cached_fonts[0] else None)
        wb_tpl.close()
    except Exception as e:
        print(f"❌ Lỗi khi đọc file mẫu template.xlsx: {e}")
        return

    yellow_fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")

    print("\nĐang mở và đọc dữ liệu từ file Excel (vui lòng chờ)...")
    try:
        wb_in = openpyxl.load_workbook(file_path, data_only=True)
    except Exception as e:
        print(f"❌ Lỗi khi mở file Excel '{file_path.name}': {e}")
        return

    visible_sheets = [ws for ws in wb_in.worksheets if ws.sheet_state == 'visible']
    if not visible_sheets:
        print("⚠️ File Excel không có sheet nào hiển thị dữ liệu.")
        return

    print(f"📁 Tìm thấy tổng cộng {len(visible_sheets)} sheet dữ liệu cần quét:")

    all_kept_rows = []
    total_scanned_rows = 0
    total_imported_rows = 0
    total_unimported_rows = 0
    processed_sheets_count = 0
    skipped_sheets = []

    for sheet_idx, ws_in in enumerate(visible_sheets, start=1):
        print(f"\n[{sheet_idx}/{len(visible_sheets)}] Quét Sheet: '{ws_in.title}' ...", end=" ")

        # 1. Tìm dòng tiêu đề và xác định vị trí cột Kết quả nhập (AC - index 28)
        header_row_idx = None
        result_col_idx = 28  # Mặc định AC (cột thứ 29)

        for r in range(1, min(30, ws_in.max_row + 1)):
            row_vals = [clean_whitespace(str(ws_in.cell(row=r, column=c).value or '')) for c in range(1, min(35, ws_in.max_column + 1))]
            row_lower = [v.lower() for v in row_vals]

            has_stt = any(s in ["stt", "số tt", "số thứ tự", "stt.", "tt", "tt."] for s in row_lower[:5])
            has_name = any(("họ" in s and "tên" in s) for s in row_lower)
            has_res = any("kết quả" in s for s in row_lower)

            if has_stt or has_name or has_res:
                header_row_idx = r
                for idx, h in enumerate(row_lower):
                    if "kết quả" in h or "ket qua" in h:
                        result_col_idx = idx
                        break
                break

        start_extracting = False
        empty_consecutive_rows = 0
        sheet_all_students = []

        for row in ws_in.iter_rows():
            row_idx = row[0].row if hasattr(row[0], 'row') else 0

            # Bỏ qua tuyệt đối các dòng trước và chính dòng tiêu đề
            if header_row_idx is not None and row_idx <= header_row_idx:
                continue

            cell_a_val = row[0].value
            cell_a_str = str(cell_a_val).strip() if cell_a_val is not None else ""

            name_cell_val = row[4].value if len(row) > 4 else None
            name_cell_str = clean_whitespace(str(name_cell_val)) if name_cell_val is not None else ""
            has_letters_in_name = bool(re.search(r'[a-zA-ZÀ-ỹ]', name_cell_str))

            is_row_empty = not any(c.value is not None and str(c.value).strip() != "" for c in row[:15])

            if not start_extracting:
                is_stt_one = bool(re.match(r'^0*1(\.0*)?$', cell_a_str))
                if (is_stt_one and has_letters_in_name) or (has_letters_in_name and name_cell_str.lower() not in ["họ và tên", "họ tên", "họ, chữ đệm và tên"]):
                    start_extracting = True

            if start_extracting:
                if is_row_empty:
                    empty_consecutive_rows += 1
                    if empty_consecutive_rows >= 20:
                        break
                    continue
                else:
                    empty_consecutive_rows = 0

                # Bỏ qua nếu là dòng tổng kết, ghi chú cuối bảng (ví dụ: 'Danh sách này có 100 học sinh')
                if is_summary_or_footer_row(row):
                    continue

                # Chỉ lấy các dòng dữ liệu với ĐỒNG THỜI STT là số hợp lệ và Họ tên không trống
                if not is_valid_stt(cell_a_val) or is_cell_empty(name_cell_val) or not has_letters_in_name:
                    continue

                # Lấy giá trị các ô trong dòng
                row_values = [cell.value for cell in row]
                required_len = max(29, result_col_idx + 1)
                while len(row_values) < required_len:
                    row_values.append(None)

                # Kiểm tra giá trị cột Kết quả nhập (AC)
                result_raw = row_values[result_col_idx]
                result_str = clean_whitespace(str(result_raw)) if result_raw is not None else ""
                is_unimported = (result_str.lower() != "đã nhập")

                # Lấy 29 cột và làm sạch khoảng trắng thừa (giữ nguyên dữ liệu gốc ngày sinh, giới tính, dân tộc...)
                new_row_values = row_values[:29]
                while len(new_row_values) < 29:
                    new_row_values.append(None)

                new_row_values = [clean_whitespace(val) for val in new_row_values]

                # Cột D (Index 3): Tên văn bằng (nếu rỗng thì điền mặc định)
                if not new_row_values[3]:
                    new_row_values[3] = "Bằng tốt nghiệp trung học cơ sở"
                # Cột Y (Index 24): Trạng thái số hóa (nếu rỗng thì điền mặc định)
                if not new_row_values[24]:
                    new_row_values[24] = "Đã số hóa đầy đủ, chính xác"

                # Giữ nguyên giá trị cột Kết quả nhập tại index 28 (AC)
                new_row_values[28] = result_str

                sheet_all_students.append((new_row_values, is_unimported))

        total_sheet_students = len(sheet_all_students)
        if total_sheet_students == 0:
            print("-> ⚠️ Không tìm thấy dòng dữ liệu học sinh nào.")
            skipped_sheets.append(ws_in.title)
            continue

        # Tìm các index của dòng chưa nhập (khác "Đã nhập")
        target_indices = [i for i, (_, is_unimp) in enumerate(sheet_all_students) if is_unimp]
        sheet_unimp_count = len(target_indices)
        sheet_imp_count = total_sheet_students - sheet_unimp_count

        total_scanned_rows += total_sheet_students
        total_imported_rows += sheet_imp_count
        total_unimported_rows += sheet_unimp_count

        if sheet_unimp_count == 0:
            print(f"-> Quét {total_sheet_students:>4} học sinh | 100% đã nhập ({sheet_imp_count}). Bỏ qua.")
            continue

        # Giữ lại dòng chưa nhập kèm thêm 5 dòng liền trước và 5 dòng liền sau
        keep_indices = set()
        for t_idx in target_indices:
            start_k = max(0, t_idx - 5)
            end_k = min(total_sheet_students, t_idx + 5 + 1)
            for k in range(start_k, end_k):
                keep_indices.add(k)

        sorted_keep_indices = sorted(keep_indices)

        for k in sorted_keep_indices:
            r_vals, is_unimp = sheet_all_students[k]
            all_kept_rows.append({
                "values": r_vals,
                "is_highlight": is_unimp  # True: dòng chưa nhập (cần tô vàng), False: dòng tham chiếu
            })

        processed_sheets_count += 1
        print(f"-> Quét {total_sheet_students:>4} dòng | Chưa nhập: {sheet_unimp_count:>3} | Xuất kèm ±5 dòng tham chiếu: {len(sorted_keep_indices):>3}")

    try:
        wb_in.close()
    except Exception:
        pass

    M = len(all_kept_rows)
    print("\n" + "-" * 80)
    print(f"TỔNG KẾT: Đã quét {total_scanned_rows} dòng | Đã nhập: {total_imported_rows} | Chưa nhập: {total_unimported_rows} | Tổng xuất ra: {M} dòng.")
    print("-" * 80)

    if M == 0:
        print("\n" + "=" * 80)
        print("🎉 THÔNG BÁO: Tất cả học sinh trong các sheet đều ở trạng thái 'Đã nhập'!".center(80))
        print("Không có bản ghi nào bị lỗi hoặc chưa nhập cần xử lý.")
        print("=" * 80)
        return

    # Khởi tạo file xuất ra
    wb_out = None
    try:
        wb_out = openpyxl.load_workbook(io.BytesIO(template_bytes))
        ws_out = wb_out.active
        ws_out.title = "KetQuaNhap"

        # Thiết lập cột thứ 29 (AC) - Kết quả nhập
        ws_out.cell(row=2, column=29, value="Kết quả nhập")
        cell_ab2 = ws_out.cell(row=2, column=28)
        cell_ac2 = ws_out.cell(row=2, column=29)
        if cell_ab2.has_style:
            if cell_ab2.font:
                cell_ac2.font = copy.copy(cell_ab2.font)
            if cell_ab2.fill:
                cell_ac2.fill = copy.copy(cell_ab2.fill)
            if cell_ab2.border:
                cell_ac2.border = copy.copy(cell_ab2.border)
            if cell_ab2.alignment:
                cell_ac2.alignment = copy.copy(cell_ab2.alignment)

        # Mở rộng ô gộp dòng 1 từ A1:AB1 thành A1:AC1 (nếu có)
        try:
            ws_out.unmerge_cells("A1:AB1")
            ws_out.merge_cells("A1:AC1")
        except Exception:
            pass

        ws_out.column_dimensions['AC'].width = 25

        yellow_rows_count = 0
        ref_rows_count = 0
        yellow_cell_count = 0

        for item in all_kept_rows:
            row_vals = item["values"]
            is_highlight = item["is_highlight"]

            # Giữ nguyên STT gốc từ file nguồn (chuyển float dạng 1.0 thành int nếu cần)
            if isinstance(row_vals[0], float) and row_vals[0].is_integer():
                row_vals[0] = int(row_vals[0])
            elif isinstance(row_vals[0], str) and row_vals[0].isdigit():
                row_vals[0] = int(row_vals[0])

            ws_out.append(row_vals)
            curr_row = ws_out.max_row
            apply_cached_fonts(ws_out, curr_row, cached_fonts)

            if is_highlight:
                # Dòng chưa nhập: TÔ VÀNG TOÀN BỘ CỘT CỦA DÒNG (Từ cột 1 đến 29)
                for c in range(1, 30):
                    ws_out.cell(row=curr_row, column=c).fill = yellow_fill
                yellow_rows_count += 1
            else:
                # Dòng tham chiếu: Giữ nguyên màu nền bình thường (chỉ tô vàng ô thiếu thông tin nếu có)
                ref_rows_count += 1
                for c in range(1, 28):
                    if is_cell_empty(row_vals[c]):
                        ws_out.cell(row=curr_row, column=c + 1).fill = yellow_fill
                        yellow_cell_count += 1

        out_file_path = file_path.parent / "KetQuaNhap.xlsx"
        # Tránh ghi đè nếu file nguồn trùng tên KetQuaNhap.xlsx
        if file_path.resolve() == out_file_path.resolve():
            out_file_path = file_path.parent / "KetQuaNhap_Loc.xlsx"

        wb_out.save(out_file_path)

        print("\n" + "=" * 80)
        print(" HOÀN THÀNH XỬ LÝ FILE KẾT QUẢ NHẬP ".center(80, "="))
        print(f"✅ Đã quét thành công: {processed_sheets_count}/{len(visible_sheets)} sheet từ file '{file_path.name}'.")
        print(f"✅ Tổng số học sinh chưa nhập (TÔ VÀNG NGUYÊN DÒNG): {yellow_rows_count} dòng.")
        print(f"✅ Tổng số học sinh tham chiếu (±5 dòng liền kề): {ref_rows_count} dòng.")
        print(f"✅ Tổng số dòng dữ liệu xuất ra: {len(all_kept_rows)} dòng.")
        print(f"📁 File kết quả đã được lưu tại:")
        print(f"   {out_file_path.absolute()}")
        print("=" * 80)

    except PermissionError:
        print(f"\n❌ Lỗi: Không thể lưu file '{out_file_path.name}' do file đang được mở trong Excel.")
        print("👉 Vui lòng ĐÓNG file trong Excel rồi chạy lại chức năng này.")
    except Exception as e:
        print(f"\n❌ Lỗi khi xuất file Excel: {e}")
    finally:
        if wb_out is not None:
            try:
                wb_out.close()
            except Exception:
                pass


# ==============================================================================
# CÔNG ĐOẠN 4: TÁCH FILE EXCEL THEO TÊN TRƯỜNG ĐÃ HỌC (CỘT N)
# ==============================================================================
def stage4_split_by_school():
    """
    Tách từ một file Excel thành nhiều file dựa theo giá trị cột 'Tên trường đã học' (cột N):
    - Đọc file Excel nguồn được chọn (hỗ trợ kéo thả file hoặc hộp thoại chọn file).
    - Quét toàn bộ các dòng dữ liệu của các sheet hiển thị (visible).
    - Tự động nhận diện cột Tên trường đã học (mặc định cột N - index 13).
    - Gom nhóm các dòng theo từng trường học.
    - Cho phép tùy chọn đánh lại STT liên tục (1..N) cho từng file hoặc giữ nguyên STT gốc.
    - Giữ nguyên cấu trúc 28 cột chuẩn (hoặc 29 cột nếu có cột Kết quả nhập), bảo toàn các ô tô vàng.
    - Xuất mỗi trường thành 1 file Excel riêng biệt trong thư mục 'Tach_Theo_Truong'.
    """
    print("\n" + "=" * 80)
    print(" CÔNG ĐOẠN 4: TÁCH FILE EXCEL THEO TÊN TRƯỜNG ĐÃ HỌC (CỘT N) ".center(80, "="))
    print("=" * 80)

    print("Đang mở cửa sổ chọn file Excel nguồn...")
    file_path_str = get_file_path("Chọn file Excel cần tách theo Tên trường đã học:")
    if not file_path_str:
        print("Bạn đã hủy chọn file.")
        return

    file_path = Path(file_path_str)
    if not file_path.is_file():
        print(f"❌ Lỗi: File không tồn tại: {file_path}")
        return

    print(f"File nguồn đã chọn: {file_path.absolute()}")

    if not TEMPLATE_PATH.exists():
        print(f"❌ Lỗi: Không tìm thấy file mẫu '{TEMPLATE_PATH.name}' tại: {TEMPLATE_PATH.parent}")
        return

    try:
        with open(TEMPLATE_PATH, "rb") as f:
            template_bytes = f.read()
        wb_tpl = openpyxl.load_workbook(io.BytesIO(template_bytes))
        ws_tpl = wb_tpl.active
        cached_fonts_28 = get_cached_row_fonts(ws_tpl, ws_tpl.max_row, max_col=28)
        wb_tpl.close()
    except Exception as e:
        print(f"❌ Lỗi khi đọc file mẫu template.xlsx: {e}")
        return

    print("\n🔢 TÙY CHỌN SỐ THỨ TỰ (STT):")
    print("   [C] - Đánh lại STT từ 1..N cho từng trường (Khuyên dùng)")
    print("   [K] - Giữ nguyên STT ban đầu từ file nguồn")
    renumber_choice = input("👉 Bạn có muốn đánh lại STT (1..N) cho từng file không? [C/k]: ").strip().lower()
    re_stt = (renumber_choice not in ["k", "khong", "không", "n", "no"])
    if re_stt:
        print("-> Đã chọn: Tự động đánh lại STT liên tục (1..N) cho từng trường.")
    else:
        print("-> Đã chọn: Giữ nguyên STT gốc từ file nguồn.")

    yellow_fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")

    print("\nĐang mở và đọc dữ liệu từ file Excel (vui lòng chờ)...")
    try:
        wb_in = openpyxl.load_workbook(file_path, data_only=True)
    except Exception as e:
        print(f"❌ Lỗi khi mở file Excel '{file_path.name}': {e}")
        return

    visible_sheets = [ws for ws in wb_in.worksheets if ws.sheet_state == 'visible']
    if not visible_sheets:
        print("⚠️ File Excel không có sheet nào hiển thị dữ liệu.")
        return

    print(f"📁 Tìm thấy {len(visible_sheets)} sheet cần quét dữ liệu:")

    # Gom nhóm dữ liệu theo tên trường: key_norm -> dict
    school_groups = {}
    total_students_found = 0
    processed_sheets_count = 0

    for s_idx, ws_in in enumerate(visible_sheets, start=1):
        print(f"  [{s_idx}/{len(visible_sheets)}] Quét Sheet: '{ws_in.title}' ...", end=" ")

        # Nhận diện dòng tiêu đề và vị trí cột "Tên trường đã học"
        header_row_idx = 2
        school_col_idx = 13  # Mặc định cột N (0-indexed: 13, tức cột thứ 14)
        sheet_has_col_ac = False

        for r in range(1, min(30, ws_in.max_row + 1)):
            row_vals = [clean_whitespace(str(ws_in.cell(row=r, column=c).value or '')) for c in range(1, min(ws_in.max_column + 1, 35))]
            row_lower = [v.lower() for v in row_vals]
            has_stt = any(s in ["stt", "số tt", "số thứ tự", "stt.", "tt", "tt."] for s in row_lower[:5])
            has_name = any(("họ" in s and "tên" in s) for s in row_lower)

            if has_stt or has_name:
                header_row_idx = r
                for c_idx, h in enumerate(row_lower):
                    if "tên trường" in h or "trường đã học" in h or "trường học" in h:
                        school_col_idx = c_idx
                        break
                if any("kết quả nhập" in h for h in row_lower) or (len(row_lower) >= 29 and "kết quả" in row_lower[28]):
                    sheet_has_col_ac = True
                break

        max_cols = 29 if sheet_has_col_ac else 28
        sheet_student_count = 0
        empty_consecutive_rows = 0

        for r in range(header_row_idx + 1, ws_in.max_row + 1):
            cells = [ws_in.cell(row=r, column=c) for c in range(1, max_cols + 1)]

            # Bỏ qua dòng trống
            is_empty = not any(c.value is not None and str(c.value).strip() != "" for c in cells[:15])
            if is_empty:
                empty_consecutive_rows += 1
                if empty_consecutive_rows >= 25:
                    break
                continue
            else:
                empty_consecutive_rows = 0

            # Bỏ qua dòng tổng kết / ghi chú chân trang
            if is_summary_or_footer_row(cells):
                continue

            cell_a_val = cells[0].value
            name_val = cells[4].value if len(cells) > 4 else None
            name_str = clean_whitespace(str(name_val)) if name_val is not None else ""
            if name_str.lower() in ["họ và tên", "họ tên", "họ, chữ đệm và tên"]:
                continue

            has_letters = bool(re.search(r'[a-zA-ZÀ-ỹ]', name_str))

            # Dòng hợp lệ phải có STT hợp lệ hoặc Họ tên có chữ cái
            if not is_valid_stt(cell_a_val) and not has_letters:
                continue

            row_vals = [clean_whitespace(c.value) for c in cells]
            while len(row_vals) < max_cols:
                row_vals.append(None)

            # Lấy thông tin màu vàng của từng ô
            row_yellows = []
            for col_idx, c in enumerate(cells):
                is_yel = is_cell_yellow(c) or (col_idx > 0 and is_cell_empty(c.value))
                row_yellows.append(is_yel)

            is_row_all_yellow = sum(1 for c in cells if is_cell_yellow(c)) >= 20

            # Lấy giá trị cột Tên trường đã học
            raw_school = row_vals[school_col_idx] if len(row_vals) > school_col_idx else None
            school_str = clean_whitespace(str(raw_school)) if raw_school is not None else ""

            if not school_str or school_str.lower() in ("none", "nan", "null", "<na>"):
                school_key = "Chưa xác định trường"
                display_name = "Chưa xác định trường"
            else:
                school_key = school_str
                display_name = school_str

            norm_key = school_key.lower().strip()
            if norm_key not in school_groups:
                school_groups[norm_key] = {
                    "display_name": display_name,
                    "rows": [],
                    "has_col_ac": sheet_has_col_ac
                }
            if sheet_has_col_ac:
                school_groups[norm_key]["has_col_ac"] = True

            school_groups[norm_key]["rows"].append({
                "values": row_vals,
                "yellows": row_yellows,
                "is_row_highlight": is_row_all_yellow
            })

            sheet_student_count += 1
            total_students_found += 1

        processed_sheets_count += 1
        col_letter = get_col_letter(school_col_idx)
        print(f"-> Trích xuất được {sheet_student_count:>4} học sinh (Cột trường: {col_letter}).")

    try:
        wb_in.close()
    except Exception:
        pass

    if total_students_found == 0:
        print("\n⚠️ Không tìm thấy dòng dữ liệu học sinh nào trong file Excel đã chọn!")
        return

    print("\n" + "-" * 80)
    print(f"TỔNG KẾT DỮ LIỆU: Tìm thấy {total_students_found} học sinh thuộc {len(school_groups)} trường học khác nhau.")
    print("-" * 80)

    # Thư mục xuất file tách
    output_dir = file_path.parent / "Tach_Theo_Truong"
    output_dir.mkdir(parents=True, exist_ok=True)

    # Chuẩn bị tên file cho từng trường (tránh trùng lặp sau sanitize)
    used_filenames = {}
    for norm_key, grp in school_groups.items():
        base_name = sanitize_filename(grp["display_name"])
        if base_name in used_filenames:
            used_filenames[base_name] += 1
            file_name = f"{base_name}_{used_filenames[base_name]}.xlsx"
        else:
            used_filenames[base_name] = 1
            file_name = f"{base_name}.xlsx"
        grp["output_filename"] = file_name

    # Font cache cho cả trường hợp 28 cột và 29 cột
    cached_fonts_29 = list(cached_fonts_28)
    cached_fonts_29.append(Font(name=cached_fonts_28[0].name, size=cached_fonts_28[0].size) if cached_fonts_28 and cached_fonts_28[0] else None)

    created_files_summary = []
    error_count = 0

    print(f"\n🚀 Đang xuất các file Excel vào thư mục 'Tach_Theo_Truong'...")

    for idx, (norm_key, grp) in enumerate(school_groups.items(), start=1):
        display_name = grp["display_name"]
        file_name = grp["output_filename"]
        out_file_path = output_dir / file_name
        has_ac = grp.get("has_col_ac", False)
        target_max_col = 29 if has_ac else 28
        cached_fonts = cached_fonts_29 if has_ac else cached_fonts_28

        try:
            wb_out = openpyxl.load_workbook(io.BytesIO(template_bytes))
            ws_out = wb_out.active
            ws_out.title = "Data"

            if has_ac:
                ws_out.cell(row=2, column=29, value="Kết quả nhập")
                cell_ab2 = ws_out.cell(row=2, column=28)
                cell_ac2 = ws_out.cell(row=2, column=29)
                if cell_ab2.has_style:
                    if cell_ab2.font:
                        cell_ac2.font = copy.copy(cell_ab2.font)
                    if cell_ab2.fill:
                        cell_ac2.fill = copy.copy(cell_ab2.fill)
                    if cell_ab2.border:
                        cell_ac2.border = copy.copy(cell_ab2.border)
                    if cell_ab2.alignment:
                        cell_ac2.alignment = copy.copy(cell_ab2.alignment)
                try:
                    ws_out.unmerge_cells("A1:AB1")
                    ws_out.merge_cells("A1:AC1")
                except Exception:
                    pass
                ws_out.column_dimensions['AC'].width = 25

            for r_idx, item in enumerate(grp["rows"], start=1):
                r_vals = list(item["values"])
                while len(r_vals) < target_max_col:
                    r_vals.append(None)
                r_vals = r_vals[:target_max_col]

                if re_stt:
                    r_vals[0] = r_idx
                else:
                    if isinstance(r_vals[0], float) and r_vals[0].is_integer():
                        r_vals[0] = int(r_vals[0])
                    elif isinstance(r_vals[0], str) and r_vals[0].isdigit():
                        r_vals[0] = int(r_vals[0])

                ws_out.append(r_vals)
                curr_row = ws_out.max_row
                apply_cached_fonts(ws_out, curr_row, cached_fonts)

                if item["is_row_highlight"]:
                    for c in range(1, target_max_col + 1):
                        ws_out.cell(row=curr_row, column=c).fill = yellow_fill
                else:
                    for c_idx, is_yel in enumerate(item["yellows"][:target_max_col]):
                        col_pos = c_idx + 1
                        if col_pos > 1 and is_yel:
                            ws_out.cell(row=curr_row, column=col_pos).fill = yellow_fill

            wb_out.save(out_file_path)
            created_files_summary.append({
                "school": display_name,
                "file": file_name,
                "count": len(grp["rows"])
            })
            print(f"  [{idx:>2}/{len(school_groups)}] ✅ {display_name} ({len(grp['rows'])} học sinh) -> {file_name}")

        except PermissionError:
            print(f"  [{idx:>2}/{len(school_groups)}] ❌ Lỗi: File '{file_name}' đang mở trong Excel. Không thể ghi đè!")
            error_count += 1
        except Exception as e:
            print(f"  [{idx:>2}/{len(school_groups)}] ❌ Lỗi khi xuất file '{file_name}': {e}")
            error_count += 1
        finally:
            try:
                wb_out.close()
            except Exception:
                pass

    print("\n" + "=" * 80)
    print(" HOÀN THÀNH TÁCH FILE THEO TÊN TRƯỜNG ".center(80, "="))
    print(f"✅ Đã quét thành công: {processed_sheets_count}/{len(visible_sheets)} sheet từ file '{file_path.name}'.")
    print(f"✅ Tổng số học sinh đã tách: {total_students_found} dòng dữ liệu.")
    print(f"✅ Số file Excel xuất thành công: {len(created_files_summary)}/{len(school_groups)} file.")
    if error_count > 0:
        print(f"⚠️ Có {error_count} file gặp lỗi khi lưu (vui lòng đóng file nếu đang mở trong Excel và chạy lại).")
    print(f"📁 Thư mục lưu các file đã tách:")
    print(f"   {output_dir.absolute()}")
    print("=" * 80)


# ==============================================================================
# MENU ĐIỀU KHIỂN CHÍNH
# ==============================================================================
def main():
    while True:
        print("\n" + "=" * 75)
        print("     CÔNG CỤ XỬ LÝ & SỐ HÓA VĂN BẰNG CHỨNG CHỈ THCS".center(75))
        print("=" * 75)
        print("  [1] Chuẩn hóa hàng tiêu đề và dữ liệu")
        print("      -> Quét từng file nguồn (thứ tự ABC) & chuẩn hóa cấu trúc 28 cột")
        print("      -> Chuẩn hóa ngày sinh và ngày cấp bằng (nhắc nhập tay nếu ngày chưa chuẩn)")
        print("      -> Tự động tô màu vàng vào các ô/cột bị thiếu dữ liệu")
        print("      -> Xuất các file Excel lẻ đã chuẩn hóa vào thư mục 'Da_Chuan_Hoa'")
        print()
        print("  [2] Quét các file đã chuẩn hóa và tạo lại file Excel tổng hợp")
        print("      -> Quét toàn bộ file trong thư mục 'Da_Chuan_Hoa'")
        print("      -> Đánh lại STT liên tục 1..N, giữ nguyên tô vàng các ô thiếu")
        print("      -> Giữ nguyên ngày sinh (không cần chuẩn hóa lại)")
        print("      -> Xuất ra file Excel tổng hợp: <Tên_Thư_Mục>_TongHop.xlsx")
        print()
        print("  [3] Lọc & ghép dữ liệu chưa nhập từ file Kết quả nhập")
        print("      -> Xử lý 1 file Excel chứa nhiều sheet có cùng cấu trúc")
        print("      -> Hàng tiêu đề chuẩn cột Tên đơn vị là AB, Kết quả nhập là AC")
        print("      -> Giữ nguyên STT và dữ liệu gốc (không đánh lại STT)")
        print("      -> TÔ VÀNG TOÀN BỘ DÒNG có Kết quả nhập KHÁC 'Đã nhập'")
        print("      -> Giữ kèm thêm 5 dòng liền trước và 5 dòng liền sau để tham chiếu")
        print("      -> Ghép vào 1 sheet duy nhất & xuất file 'KetQuaNhap.xlsx' cùng thư mục")
        print()
        print("  [4] Tách 1 file Excel thành nhiều file theo Tên trường đã học (cột N)")
        print("      -> Cho phép chọn 1 file Excel nguồn (hỗ trợ kéo thả hoặc hộp thoại)")
        print("      -> Tự động gom nhóm học sinh theo giá trị tại cột Tên trường đã học")
        print("      -> Tùy chọn đánh lại STT liên tục (1..N) cho từng trường")
        print("      -> Bảo toàn định dạng 28 cột chuẩn và các ô tô màu vàng")
        print("      -> Xuất các file Excel riêng lẻ vào thư mục 'Tach_Theo_Truong'")
        print()
        print("  [0] Thoát chương trình")
        print("=" * 75)

        choice = input("👉 Chọn chức năng [0-4]: ").strip()

        try:
            if choice == "1":
                stage1_normalize_individual_files()
            elif choice == "2":
                stage2_merge_normalized_files()
            elif choice == "3":
                stage3_filter_unimported_records()
            elif choice == "4":
                stage4_split_by_school()
            elif choice in ["0", "q", "exit"]:
                print("\nĐã thoát chương trình. Tạm biệt!\n")
                break
            else:
                print("\n⚠️ Lựa chọn không hợp lệ! Vui lòng nhập số từ 0 đến 4.")
        except KeyboardInterrupt:
            print("\n\n⚠️ Đã hủy thao tác hiện tại theo yêu cầu người dùng (quay lại Menu chính).")

        try:
            input("\nNhấn Enter để quay lại Menu chính...")
        except KeyboardInterrupt:
            print("\n")

if __name__ == "__main__":
    main()
