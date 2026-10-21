from __future__ import annotations

import os
import sys
import re
from pathlib import Path
import openpyxl

def find_data_file(filename: str) -> Path:
    """Tìm đường dẫn tệp tài nguyên / danh mục tra cứu theo thứ tự ưu tiên:
    1. Thư mục data đặt bên cạnh file thực thi (.exe trên Windows, .app trên macOS).
    2. Thư mục làm việc hiện tại (cwd / data hoặc cwd).
    3. Thư mục bundle đóng gói sẵn của PyInstaller (_MEIPASS / data hoặc _MEIPASS).
    4. Thư mục data trong mã nguồn dự án.
    Điều này cho phép người dùng tùy biến Danh_Muc_Tra_Cuu.xlsx bên ngoài mà không cần đóng gói lại phần mềm.
    """
    candidates: list[Path] = []

    # 1. Khi chạy file đóng gói (PyInstaller frozen)
    if getattr(sys, 'frozen', False):
        exe_path = Path(sys.executable).resolve()
        # Đối với ứng dụng macOS .app bundle: vd SohoaTHCS.app/Contents/MacOS/SohoaTHCS
        # Thư mục chứa .app là cha cấp 4 của binary
        if len(exe_path.parents) >= 4 and exe_path.parent.name == "MacOS" and exe_path.parents[1].name == "Contents":
            app_dir = exe_path.parents[3]
            candidates.extend([app_dir / "data", app_dir])

        # Đối với file .exe trên Windows hoặc standalone binary trên Linux/macOS
        candidates.extend([exe_path.parent / "data", exe_path.parent])

    # 2. Thư mục làm việc hiện tại
    cwd = Path.cwd().resolve()
    candidates.extend([cwd / "data", cwd])

    # 3. Khi chạy bản đóng gói PyInstaller (_MEIPASS)
    if getattr(sys, 'frozen', False) and hasattr(sys, '_MEIPASS'):
        base = Path(sys._MEIPASS)
        candidates.extend([base / "data", base])

    # 4. Khi chạy từ mã nguồn
    cur = Path(__file__).resolve().parent.parent / "data"
    candidates.extend([cur, Path(__file__).resolve().parent.parent])

    # Quét theo đúng thứ tự ưu tiên
    seen = set()
    for d in candidates:
        try:
            resolved_dir = d.resolve()
            if resolved_dir in seen or not resolved_dir.is_dir():
                continue
            seen.add(resolved_dir)
            target = resolved_dir / filename
            if target.is_file():
                return target
        except Exception:
            continue

    # Fallback cuối cùng nếu không tìm thấy tệp
    return get_data_dir() / filename

def get_data_dir() -> Path:
    """Xác định đường dẫn thư mục data chứa tài nguyên biểu mẫu (ưu tiên thư mục chứa Danh_Muc_Tra_Cuu.xlsx)."""
    p = find_data_file("Danh_Muc_Tra_Cuu.xlsx")
    if p.is_file():
        return p.parent

    if getattr(sys, 'frozen', False) and hasattr(sys, '_MEIPASS'):
        base = Path(sys._MEIPASS)
        if (base / "data").exists():
            return base / "data"
        return base

    cur = Path(__file__).resolve().parent.parent / "data"
    if cur.exists():
        return cur
    return Path.cwd()

def get_persistent_data_dir() -> Path:
    """
    Xác định thư mục dữ liệu người dùng bền vững (Persistent Data Directory).
    Dùng cho các tệp CẦN GHI VÀ LƯU TRỮ LÂU DÀI:
    - CSDL SQLite: vbcc_database.db
    - Thư mục sao lưu: backups/
    - Cấu hình ánh xạ cột: mapping_profiles.json
    
    Quy tắc xác định:
    1. Khi chạy từ mã nguồn (Development / Chay_Chuong_Trinh.command):
       Trả về thư mục <gốc_dự_án>/data.
    2. Khi chạy ứng dụng đóng gói độc lập (PyInstaller Frozen):
       - Ưu tiên 1 (Portable Mode): Thư mục `data/` đặt ngay cạnh tệp thực thi (.exe trên Windows, hoặc cạnh .app trên macOS).
         Nếu thư mục này có quyền ghi (writable), sử dụng ngay để dữ liệu mang tính di động.
       - Ưu tiên 2 (Standard AppData / Application Support):
         Nếu thư mục chứa ứng dụng là chỉ đọc (ví dụ macOS /Applications hoặc Windows C:\\Program Files):
         + macOS: ~/Library/Application Support/SohoaTHCS/data
         + Windows: %APPDATA%\\SohoaTHCS\\data
         + Linux: ~/.local/share/SohoaTHCS/data
    """
    # 1. Chạy từ mã nguồn
    if not getattr(sys, 'frozen', False):
        cur = Path(__file__).resolve().parent.parent / "data"
        cur.mkdir(parents=True, exist_ok=True)
        return cur

    # 2. Chạy từ ứng dụng đóng gói độc lập
    exe_path = Path(sys.executable).resolve()

    candidate_portable: Path | None = None
    if sys.platform == "darwin" and "Contents/MacOS" in str(exe_path):
        # exe_path: /path/to/SohoaTHCS.app/Contents/MacOS/SohoaTHCS
        app_bundle_dir = exe_path.parents[3]  # Thư mục chứa .app
        if str(app_bundle_dir) not in ("/Applications", "/System/Applications") and not str(app_bundle_dir).startswith("/Volumes/"):
            candidate_portable = app_bundle_dir / "data"
    else:
        # Windows .exe hoặc Linux standalone binary
        candidate_portable = exe_path.parent / "data"

    # Thử kiểm tra quyền ghi tại thư mục portable
    if candidate_portable is not None:
        try:
            candidate_portable.mkdir(parents=True, exist_ok=True)
            test_file = candidate_portable / ".perm_check"
            test_file.touch()
            test_file.unlink()
            return candidate_portable
        except (PermissionError, OSError):
            pass

    # Fallback chuẩn theo hệ điều hành (Đảm bảo luôn luôn ghi được 100%)
    if sys.platform == "darwin":
        fallback_dir = Path.home() / "Library" / "Application Support" / "SohoaTHCS" / "data"
    elif sys.platform == "win32":
        appdata = os.getenv("APPDATA")
        if appdata:
            fallback_dir = Path(appdata) / "SohoaTHCS" / "data"
        else:
            fallback_dir = Path.home() / "AppData" / "Roaming" / "SohoaTHCS" / "data"
    else:
        xdg_data = os.getenv("XDG_DATA_HOME")
        if xdg_data:
            fallback_dir = Path(xdg_data) / "SohoaTHCS" / "data"
        else:
            fallback_dir = Path.home() / ".local" / "share" / "SohoaTHCS" / "data"

    fallback_dir.mkdir(parents=True, exist_ok=True)
    return fallback_dir

class LookupManager:
    """Quản lý nạp và tra cứu từ tệp Danh_Muc_Tra_Cuu.xlsx."""
    def __init__(self, excel_path: Path | str | None = None):
        if excel_path is None:
            self.excel_path = find_data_file("Danh_Muc_Tra_Cuu.xlsx")
        else:
            self.excel_path = Path(excel_path)
            
        self.units: dict[str, dict] = {}  # {unit_key: {"key": ..., "code": ..., "name": ..., "commune": ..., "schools": [...], "school_map": {...}}}
        self.unit_headers: list[str] = []
        self.unit_names_list: list[str] = []
        self.all_schools_list: list[str] = []
        self.all_schools_display_list: list[str] = []
        self.school_details_map: dict[str, dict] = {}
        self.gender_map: dict[str, str] = {}
        self.ethnic_map: dict[str, str] = {}
        self.nationality_map: dict[str, str] = {}
        self.birthplace_map: dict[str, str] = {}
        
        self.load_all()

    def load_all(self):
        """Nạp toàn bộ 5 sheet từ tệp Danh_Muc_Tra_Cuu.xlsx vào bộ nhớ."""
        if not self.excel_path.exists():
            raise FileNotFoundError(f"Không tìm thấy tệp danh mục tra cứu: {self.excel_path}")

        wb = openpyxl.load_workbook(self.excel_path, data_only=True)
        try:
            # 1. Sheet DS_Don_Vi
            if "DS_Don_Vi" in wb.sheetnames:
                ws = wb["DS_Don_Vi"]
                # Đọc dòng tiêu đề (Dòng 1) để nhận diện vị trí các cột linh hoạt
                header_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), None)
                self.unit_headers = [str(c).strip() for c in header_row if c] if header_row else []
                
                # Mặc định theo index chuẩn: 0: Mã trường, 1: Tên trường, 2: Mã đơn vị, 3: Tên đơn vị, 4: Xã/phường
                sch_code_idx = 0
                sch_name_idx = 1
                u_code_idx = 2
                u_name_idx = 3
                commune_idx = 4
                extra_col_map = {}

                if header_row:
                    for c_idx, col_name in enumerate(header_row):
                        if not col_name:
                            continue
                        cn_lower = str(col_name).strip().lower()
                        if "mã trường" in cn_lower:
                            sch_code_idx = c_idx
                        elif "tên trường" in cn_lower:
                            sch_name_idx = c_idx
                        elif "mã đơn vị" in cn_lower:
                            u_code_idx = c_idx
                        elif "tên đơn vị" in cn_lower:
                            u_name_idx = c_idx
                        elif any(w in cn_lower for w in ["xã/phường", "xã / phường", "tên xã", "tên phường", "xã", "phường"]) and "mã" not in cn_lower:
                            commune_idx = c_idx
                        else:
                            extra_col_map[str(col_name).strip()] = c_idx

                for row in ws.iter_rows(min_row=2, values_only=True):
                    if not row:
                        continue
                    sch_code = str(row[sch_code_idx] or '').strip() if len(row) > sch_code_idx and row[sch_code_idx] is not None else ''
                    sch_name = str(row[sch_name_idx] or '').strip() if len(row) > sch_name_idx and row[sch_name_idx] is not None else ''
                    u_code = str(row[u_code_idx] or '').strip() if len(row) > u_code_idx and row[u_code_idx] is not None else ''
                    u_name = str(row[u_name_idx] or '').strip() if len(row) > u_name_idx and row[u_name_idx] is not None else ''
                    commune = str(row[commune_idx] or '').strip() if len(row) > commune_idx and row[commune_idx] is not None else ''

                    sch_code = re.sub(r'\.0+$', '', sch_code)
                    u_code = re.sub(r'\.0+$', '', u_code)

                    # Thu thập các cột người dùng bổ sung thêm (nếu có)
                    extra_data = {}
                    for col_title, c_i in extra_col_map.items():
                        if len(row) > c_i and row[c_i] is not None and str(row[c_i]).strip():
                            val_clean = str(row[c_i]).strip()
                            val_clean = re.sub(r'\.0+$', '', val_clean)
                            extra_data[col_title] = val_clean

                    if u_name:
                        # Tạo khóa định danh duy nhất hiển thị cả Tên đơn vị, Xã/phường và Mã đơn vị
                        if commune:
                            unit_key = f"{u_name} - {commune} [Mã: {u_code}]"
                        else:
                            unit_key = f"{u_name} [Mã: {u_code}]"

                        if unit_key not in self.units:
                            self.units[unit_key] = {
                                "key": unit_key,
                                "code": u_code,
                                "name": u_name,
                                "commune": commune,
                                "district_code": extra_data.get("Mã huyện cũ", ""),
                                "district_name": extra_data.get("Tên Huyện cũ", ""),
                                "province_name": extra_data.get("Tên tỉnh cũ", ""),
                                "province_code": extra_data.get("Mã tỉnh cũ", ""),
                                "commune_code": extra_data.get("Mã xã phường", ""),
                                "schools": [],
                                "school_map": {},
                                "extra": extra_data
                            }
                        else:
                            if extra_data:
                                self.units[unit_key].setdefault("extra", {}).update(extra_data)
                                if not self.units[unit_key].get("district_code") and "Mã huyện cũ" in extra_data:
                                    self.units[unit_key]["district_code"] = extra_data["Mã huyện cũ"]
                                if not self.units[unit_key].get("district_name") and "Tên Huyện cũ" in extra_data:
                                    self.units[unit_key]["district_name"] = extra_data["Tên Huyện cũ"]
                                if not self.units[unit_key].get("province_name") and "Tên tỉnh cũ" in extra_data:
                                    self.units[unit_key]["province_name"] = extra_data["Tên tỉnh cũ"]
                                if not self.units[unit_key].get("province_code") and "Mã tỉnh cũ" in extra_data:
                                    self.units[unit_key]["province_code"] = extra_data["Mã tỉnh cũ"]
                                if not self.units[unit_key].get("commune_code") and "Mã xã phường" in extra_data:
                                    self.units[unit_key]["commune_code"] = extra_data["Mã xã phường"]

                        u_dict = self.units[unit_key]
                        if sch_name and sch_name not in u_dict["schools"]:
                            u_dict["schools"].append(sch_name)
                        if sch_name and sch_code:
                            u_dict["school_map"][sch_name.lower()] = sch_code

                    if sch_name:
                        dist_name = extra_data.get("Tên Huyện cũ", "")
                        prov_name = extra_data.get("Tên tỉnh cũ", "")
                        if dist_name and prov_name:
                            disp_name = f"{sch_name} — [{dist_name}, {prov_name}]"
                        elif dist_name or prov_name:
                            disp_name = f"{sch_name} — [{dist_name or prov_name}]"
                        else:
                            disp_name = sch_name

                        sch_detail = {
                            "school_code": sch_code,
                            "school_name": sch_name,
                            "display_name": disp_name,
                            "unit_code": u_code,
                            "unit_name": u_name,
                            "commune": commune,
                            "commune_code": extra_data.get("Mã xã phường", ""),
                            "district_code": extra_data.get("Mã huyện cũ", ""),
                            "district_name": dist_name,
                            "province_name": prov_name,
                        }
                        self.school_details_map[sch_name.lower()] = sch_detail
                        self.school_details_map[disp_name.lower()] = sch_detail
                            
                self.unit_names_list = sorted(list(self.units.keys()))
                self.all_schools_list = sorted(list({v["school_name"] for v in self.school_details_map.values()}))
                self.all_schools_display_list = sorted(list({v["display_name"] for v in self.school_details_map.values()}))

            # 2. Sheet Gioi_Tinh
            if "Gioi_Tinh" in wb.sheetnames:
                ws = wb["Gioi_Tinh"]
                for row in ws.iter_rows(min_row=2, values_only=True):
                    if not row or len(row) < 2:
                        continue
                    code = str(row[0] or '').strip()
                    name = str(row[1] or '').strip()
                    if code and name:
                        code_clean = re.sub(r'\.0+$', '', code)
                        self.gender_map[code_clean] = name

            # 3. Sheet Dan_Toc
            if "Dan_Toc" in wb.sheetnames:
                ws = wb["Dan_Toc"]
                for row in ws.iter_rows(min_row=2, values_only=True):
                    if not row or len(row) < 2:
                        continue
                    code = str(row[0] or '').strip()
                    name = str(row[1] or '').strip()
                    if code and name:
                        code_clean = re.sub(r'\.0+$', '', code)
                        self.ethnic_map[code_clean] = name

            # 4. Sheet Quoc_Tich
            if "Quoc_Tich" in wb.sheetnames:
                ws = wb["Quoc_Tich"]
                for row in ws.iter_rows(min_row=2, values_only=True):
                    if not row or len(row) < 2:
                        continue
                    code = str(row[0] or '').strip()
                    name = str(row[1] or '').strip()
                    if code and name:
                        code_clean = re.sub(r'\.0+$', '', code)
                        self.nationality_map[code_clean] = name

            # 5. Sheet Noi_Sinh
            if "Noi_Sinh" in wb.sheetnames:
                ws = wb["Noi_Sinh"]
                for row in ws.iter_rows(min_row=2, values_only=True):
                    if not row or len(row) < 2:
                        continue
                    code = str(row[0] or '').strip()
                    name = str(row[1] or '').strip()
                    if code and name:
                        code_clean = re.sub(r'\.0+$', '', code)
                        self.birthplace_map[code_clean] = name
        finally:
            wb.close()

    def get_unit_info(self, unit_identifier: str) -> dict | None:
        """Lấy thông tin đơn vị từ unit_key, tên đơn vị hoặc mã đơn vị."""
        if not unit_identifier:
            return None
        u_clean = unit_identifier.strip()
        # 1. Khớp trực tiếp unit_key
        if u_clean in self.units:
            return self.units[u_clean]

        # 2. Tìm theo tên đơn vị hoặc mã đơn vị
        for key, info in self.units.items():
            if key.lower() == u_clean.lower() or info["name"].lower() == u_clean.lower():
                return info
            if info["code"] == u_clean:
                return info

        return None

    def get_school_code(self, unit_info: dict, school_name: str) -> str:
        """Tra cứu mã trường trực thuộc đơn vị từ tên trường đã học."""
        if not unit_info or not school_name:
            return ""
        s_map = unit_info.get("school_map", {})
        s_clean = school_name.strip().lower()
        if s_clean in s_map:
            return s_map[s_clean]

        # Khớp mềm (bỏ khoảng trắng và ký tự đặc biệt)
        s_alphanumeric = re.sub(r'\W+', '', s_clean)
        for k, code in s_map.items():
            k_alphanumeric = re.sub(r'\W+', '', k)
            if s_alphanumeric == k_alphanumeric:
                return code

        # Tìm kiếm chuỗi con
        for k, code in s_map.items():
            if k in s_clean or s_clean in k:
                return code

        return ""

    def map_gender(self, val) -> str:
        """Ánh xạ giới tính từ mã hoặc làm sạch chuỗi."""
        if val is None:
            return ""
        s = str(val).strip()
        if not s or s.lower() in ("none", "nan", "null"):
            return ""
        s_clean = re.sub(r'\.0+$', '', s)
        if s_clean in self.gender_map:
            return self.gender_map[s_clean]
        s_lower = s_clean.lower()
        if s_lower == "nam":
            return "Nam"
        elif s_lower in ("nữ", "nu"):
            return "Nữ"
        elif s_lower in ("khác", "khac"):
            return "Khác"
        return s_clean

    def map_ethnic(self, val) -> str:
        """Ánh xạ dân tộc từ mã hoặc trả về chuỗi gốc."""
        if val is None:
            return ""
        s = str(val).strip()
        if not s or s.lower() in ("none", "nan", "null"):
            return ""
        s_clean = re.sub(r'\.0+$', '', s)
        if s_clean in self.ethnic_map:
            return self.ethnic_map[s_clean]
        return s_clean

    def map_nationality(self, val) -> str:
        """Ánh xạ quốc tịch từ mã hoặc trả về chuỗi gốc."""
        if val is None:
            return ""
        s = str(val).strip()
        if not s or s.lower() in ("none", "nan", "null"):
            return ""
        s_clean = re.sub(r'\.0+$', '', s)
        if s_clean in self.nationality_map:
            return self.nationality_map[s_clean]
        if s_clean.lower() in ("việt nam", "viet nam", "vn"):
            return "Việt Nam"
        return s_clean

    def map_birthplace(self, val) -> str:
        """Ánh xạ nơi sinh từ mã hoặc trả về chuỗi gốc."""
        if val is None:
            return ""
        s = str(val).strip()
        if not s or s.lower() in ("none", "nan", "null"):
            return ""
        s_clean = re.sub(r'\.0+$', '', s)
        if s_clean in self.birthplace_map:
            return self.birthplace_map[s_clean]
        return s_clean

    def get_all_schools(self) -> list[str]:
        """Lấy toàn bộ danh sách tên trường chuẩn (không phân biệt đơn vị)."""
        return self.all_schools_list

    def get_all_schools_display(self) -> list[str]:
        """Lấy toàn bộ danh sách tên trường hiển thị kèm [Huyện cũ, Tỉnh cũ]."""
        return self.all_schools_display_list

    def parse_school_clean_name(self, name_or_display: str) -> str:
        """Tách lấy tên trường thuần túy nếu người dùng chọn chuỗi dạng 'Tên trường — [Huyện, Tỉnh]'."""
        if not name_or_display:
            return ""
        s = str(name_or_display).strip()
        if " — [" in s and s.endswith("]"):
            return s.split(" — [")[0].strip()
        elif " - [" in s and s.endswith("]"):
            return s.split(" - [")[0].strip()
        return s

    def get_school_details(self, school_identifier: str) -> dict | None:
        """
        Tra cứu chi tiết trường học gồm: school_code, unit_code, unit_name,
        commune, commune_code, district_code, district_name, province_name.
        Hỗ trợ tìm theo tên trường thuần, tên hiển thị hoặc tìm mờ.
        """
        if not school_identifier:
            return None
        clean_name = self.parse_school_clean_name(school_identifier).lower()
        if clean_name in self.school_details_map:
            return self.school_details_map[clean_name]

        # Khớp mềm (bỏ khoảng trắng và ký tự đặc biệt)
        s_alphanumeric = re.sub(r'\W+', '', clean_name)
        for k, info in self.school_details_map.items():
            if s_alphanumeric == re.sub(r'\W+', '', k):
                return info

        # Tìm kiếm chuỗi con
        for k, info in self.school_details_map.items():
            if k in clean_name or clean_name in k:
                return info

        return None
