from __future__ import annotations

import io
import re
import copy
import unicodedata
from datetime import datetime
from pathlib import Path
from typing import Any, Callable

import openpyxl
from openpyxl.styles import Font

from src.lookup_manager import find_data_file
from src.normalizer import parse_nien_khoa

TEXT_COLS_STUDENT = {2, 3, 6, 7, 8, 9, 19, 24, 27}
TEXT_COLS_DOC = {2, 4, 6, 7, 8}

def sanitize_filename(name: str) -> str:
    """Loại bỏ các ký tự cấm của hệ điều hành, bảo toàn tiếng Việt có dấu."""
    if not name:
        return "Chua_Xac_Dinh"
    s = unicodedata.normalize('NFC', str(name)).strip()
    s = re.sub(r'[\\/*?:"<>|]', '_', s)
    s = re.sub(r'\s+', ' ', s).strip()
    s = re.sub(r'_+', '_', s).strip('._ ')
    return s if s else "Chua_Xac_Dinh"

def parse_cohort_grad_year(cohort_str: str) -> int | None:
    """
    Trích xuất năm ra trường (tốt nghiệp) từ chuỗi niên khóa.
    Ví dụ: '2016-2020' -> 2020, '2020' -> 2020, 'Khóa 2018-2022' -> 2022.
    """
    if not cohort_str:
        return None
    _, _, y2, _ = parse_nien_khoa(cohort_str)
    if y2 and isinstance(y2, int):
        return y2
    match = re.findall(r'\b(19\d{2}|20\d{2})\b', str(cohort_str))
    if match:
        return int(match[-1])
    return None

def analyze_cohorts(cohorts: list[str]) -> dict[str, Any]:
    """
    Phân tích danh sách niên khóa được chọn để xác định tính liên tục:
    - 1 năm: is_single = True
    - Nhiều năm liên tiếp: is_consecutive = True
    - Nhiều năm không liên tiếp (rời rạc): is_consecutive = False
    """
    if not cohorts:
        return {
            "selected_cohorts": [],
            "year_to_cohorts": {},
            "sorted_years": [],
            "is_single": False,
            "is_consecutive": False,
            "min_year": 0,
            "max_year": 0
        }

    year_to_cohorts: dict[int, list[str]] = {}
    for c in cohorts:
        y = parse_cohort_grad_year(c)
        if y is not None:
            year_to_cohorts.setdefault(y, []).append(c)

    sorted_years = sorted(year_to_cohorts.keys())
    if not sorted_years:
        return {
            "selected_cohorts": cohorts,
            "year_to_cohorts": {},
            "sorted_years": [],
            "is_single": False,
            "is_consecutive": False,
            "min_year": 0,
            "max_year": 0
        }

    num_years = len(sorted_years)
    is_single = (num_years == 1)
    min_year = sorted_years[0]
    max_year = sorted_years[-1]

    if is_single:
        is_consecutive = True
    else:
        # Kiểm tra tính liên tục: khoảng cách giữa các năm liền kề luôn là 1
        is_consecutive = all(
            sorted_years[i] - sorted_years[i - 1] == 1
            for i in range(1, num_years)
        )

    return {
        "selected_cohorts": cohorts,
        "year_to_cohorts": year_to_cohorts,
        "sorted_years": sorted_years,
        "is_single": is_single,
        "is_consecutive": is_consecutive,
        "min_year": min_year,
        "max_year": max_year
    }

class ExportManager:
    """
    Bộ điều phối xuất dữ liệu và báo cáo Excel tập trung cho Phân hệ 3:
    1. Xuất Biểu mẫu chuẩn 28 cột của Bộ/Sở theo 4 cấp (Huyện cũ / Xã phường / Đơn vị / Trường).
    2. Xuất kèm tệp Danh sách tài liệu số hóa chuẩn 11 cột (Mau_DS_Tai_Lieu_So_Hoa.xlsx).
    3. Xuất tệp Danh sách lỗi / Cảnh báo kèm 3 cột vết nguồn gốc.
    """
    def __init__(self, lookup_mgr: Any = None):
        self.lookup_mgr = lookup_mgr
        self.template_student_path = find_data_file("Mau_DS_Cap_Bang_THCS.xlsx")
        self.template_doc_path = find_data_file("Mau_DS_Tai_Lieu_So_Hoa.xlsx")

        self.template_student_bytes: bytes | None = None
        self.cached_student_fonts: list[Font | None] = []
        self._load_student_template()

        self.template_doc_bytes: bytes | None = None
        self.cached_doc_fonts: list[Font | None] = []
        self._load_doc_template()

    def _load_student_template(self):
        if self.template_student_path.exists():
            with open(self.template_student_path, "rb") as f:
                self.template_student_bytes = f.read()
            wb = openpyxl.load_workbook(io.BytesIO(self.template_student_bytes))
            ws = wb.active
            self.cached_student_fonts = []
            for col in range(1, 29):
                cell = ws.cell(row=2, column=col)
                if cell.has_style and cell.font:
                    self.cached_student_fonts.append(Font(name=cell.font.name, size=cell.font.size))
                else:
                    self.cached_student_fonts.append(None)
            wb.close()

    def _load_doc_template(self):
        if self.template_doc_path.exists():
            with open(self.template_doc_path, "rb") as f:
                self.template_doc_bytes = f.read()
            wb = openpyxl.load_workbook(io.BytesIO(self.template_doc_bytes))
            ws = wb.active
            self.cached_doc_fonts = []
            for col in range(1, 12):
                cell = ws.cell(row=2, column=col)
                if cell.has_style and cell.font:
                    self.cached_doc_fonts.append(Font(name=cell.font.name, size=cell.font.size))
                else:
                    self.cached_doc_fonts.append(None)
            wb.close()

    def _create_student_workbook(self, sheet_title="Data") -> tuple[openpyxl.Workbook, Any]:
        """Tạo workbook chuẩn từ template 28 cột."""
        if self.template_student_bytes:
            wb = openpyxl.load_workbook(io.BytesIO(self.template_student_bytes))
            ws = wb.active
            ws.title = sheet_title
            if ws.max_row > 2:
                ws.delete_rows(3, ws.max_row - 2)
            return wb, ws
        else:
            # Fallback nếu không có template
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = sheet_title
            ws.append(["DANH SÁCH BẰNG TỐT NGHIỆP TRUNG HỌC CƠ SỞ"])
            ws.append([
                "STT", "Số hiệu bằng", "Số vào sổ gốc cấp văn bằng", "Tên văn bằng", "Họ, chữ đệm và tên",
                "Mã người học", "Số định danh cá nhân", "Hộ chiếu", "Ngày, tháng, năm sinh", "Giới tính",
                "Dân tộc", "Quốc tịch", "Nơi sinh", "Tên trường đã học", "Niên khóa",
                "Điểm các môn thi", "Hội đồng thi", "Địa danh nơi cơ quan cấp bằng đặt trụ sở",
                "Ngày tháng năm cấp bằng", "Tên cơ quan cấp bằng", "Chức danh người ký bằng",
                "Họ, chữ đệm, tên người ký bằng", "Tình trạng văn bằng", "Ghi chú", "Trạng thái số hóa",
                "Thông tin chưa bảo đảm đầy đủ, chính xác và nguyên nhân", "Mã đơn vị", "Tên đơn vị"
            ])
            return wb, ws

    def _create_doc_workbook(self, sheet_title="Data") -> tuple[openpyxl.Workbook, Any]:
        """Tạo workbook chuẩn từ template Danh sách tài liệu số hóa (11 cột)."""
        if self.template_doc_bytes:
            wb = openpyxl.load_workbook(io.BytesIO(self.template_doc_bytes))
            ws = wb.active
            ws.title = sheet_title
            if ws.max_row > 2:
                ws.delete_rows(3, ws.max_row - 2)
            return wb, ws
        else:
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = sheet_title
            ws.append(["DANH SÁCH TÀI LIỆU SỐ HÓA"])
            ws.append([
                "STT", "Mã đơn vị", "Tên đơn vị", "Mã sổ gốc", "Tên tài liệu",
                "Từ năm", "Đến năm", "Ngày lập sổ", "Số trang", "Tình trạng tài liệu", "Thông tư"
            ])
            return wb, ws

    def determine_export_level(
        self,
        branch: str = "branch_history",
        huyen_cu: str | None = None,
        ten_truong: str | None = None,
        xa_phuong: str | None = None,
        ten_don_vi: str | None = None,
        sample_student: dict | None = None
    ) -> dict[str, Any]:
        """
        Xác định cấp xuất và quy tắc ánh xạ Cột 27, 28:
        - 1. Huyện cũ: Tên huyện = cụ thể, Tên trường = 'Tất cả'
             Cột 27 = ma_huyen_cu, Cột 28 = huyen_cu
        - 2. Xã/Phường: Xã = cụ thể, Đơn vị = 'Tất cả'
             Cột 27 = ma_xa_phuong, Cột 28 = xa_phuong
        - 3. Đơn vị: Xã = cụ thể, Đơn vị = cụ thể
             Cột 27 và 28 giữ nguyên
        - 4. Trường học: Huyện = cụ thể, Trường = cụ thể
             Cột 27 và 28 giữ nguyên
        """
        st = sample_student or {}

        if branch == "branch_history":
            if not ten_truong or ten_truong == "Tất cả":
                # Cấp 1: Huyện cũ
                h_name = huyen_cu or st.get("huyen_cu", "") or "Huyen_Cu"
                h_code = st.get("ma_huyen_cu", "")
                if not h_code and self.lookup_mgr:
                    for u in self.lookup_mgr.units.values():
                        if u.get("district_name") == h_name and u.get("district_code"):
                            h_code = u["district_code"]
                            break
                return {
                    "level": "HUYEN_CU",
                    "level_title": "Cấp Huyện cũ",
                    "base_name": sanitize_filename(h_name),
                    "override_unit_code": h_code,
                    "override_unit_name": h_name,
                    "doc_entity_type": "Huyện",
                    "doc_entity_name": h_name,
                    "doc_entity_code": h_code,
                    "school_code": ""
                }
            else:
                # Cấp 4: Trường học
                sch_name = ten_truong or st.get("ten_truong", "")
                sch_code = st.get("ma_truong", "")
                if not sch_code and self.lookup_mgr:
                    details = self.lookup_mgr.get_school_details(sch_name)
                    if details:
                        sch_code = details.get("school_code", "")

                u_code = st.get("ma_don_vi", "")
                u_name = st.get("ten_don_vi", "")
                if (not u_code or not u_name) and self.lookup_mgr:
                    details = self.lookup_mgr.get_school_details(sch_name)
                    if details:
                        u_code = u_code or details.get("unit_code", "")
                        u_name = u_name or details.get("unit_name", "")

                return {
                    "level": "TRUONG",
                    "level_title": "Cấp Trường học",
                    "base_name": sanitize_filename(sch_name),
                    "override_unit_code": None,  # Giữ nguyên
                    "override_unit_name": None,  # Giữ nguyên
                    "doc_entity_type": "Trường",
                    "doc_entity_name": sch_name,
                    "doc_entity_code": sch_code or u_code,
                    "school_code": sch_code
                }
        else:
            # Nhánh Địa bàn mới
            if not ten_don_vi or ten_don_vi == "Tất cả":
                # Cấp 2: Xã/phường mới
                x_name = xa_phuong or st.get("xa_phuong", "") or "Xa_Phuong"
                x_code = st.get("ma_xa_phuong", "")
                if not x_code and self.lookup_mgr:
                    for u in self.lookup_mgr.units.values():
                        if u.get("commune") == x_name and u.get("commune_code"):
                            x_code = u["commune_code"]
                            break
                return {
                    "level": "XA_PHUONG",
                    "level_title": "Cấp Xã/Phường mới",
                    "base_name": sanitize_filename(x_name),
                    "override_unit_code": x_code,
                    "override_unit_name": x_name,
                    "doc_entity_type": "Xã/Phường",
                    "doc_entity_name": x_name,
                    "doc_entity_code": x_code,
                    "school_code": ""
                }
            else:
                # Cấp 3: Đơn vị quản lý
                u_name = ten_don_vi or st.get("ten_don_vi", "")
                u_code = st.get("ma_don_vi", "")
                if not u_code and self.lookup_mgr:
                    for u in self.lookup_mgr.units.values():
                        if u.get("name") == u_name and u.get("code"):
                            u_code = u["code"]
                            break
                return {
                    "level": "DON_VI",
                    "level_title": "Cấp Đơn vị quản lý",
                    "base_name": sanitize_filename(u_name),
                    "override_unit_code": None,  # Giữ nguyên
                    "override_unit_name": None,  # Giữ nguyên
                    "doc_entity_type": "Đơn vị",
                    "doc_entity_name": u_name,
                    "doc_entity_code": u_code,
                    "school_code": ""
                }

    def generate_doc_title_and_code(
        self,
        level_info: dict[str, Any],
        min_year: int,
        max_year: int,
        is_single_year: bool
    ) -> tuple[str, str]:
        """
        Sinh Mã sổ gốc và Tên tài liệu theo quy định:
        - Nếu xuất theo Huyện cũ, Xã/phường, Đơn vị:
          + Liên tiếp: Mã = {Mã đơn vị}_{Năm nhỏ}-{Năm lớn}, Tên = "Sổ gốc cấp bằng THCS {Tên} từ {Năm nhỏ} đến {Năm lớn}"
          + Từng năm: Mã = {Mã đơn vị}_{Năm}, Tên = "Sổ gốc cấp bằng THCS {Tên} năm {Năm}"
        - Nếu xuất theo Trường:
          + Liên tiếp: Mã = {Mã trường}_{Năm nhỏ}-{Năm lớn}, Tên = "Sổ gốc cấp bằng THCS trường {Tên trường} từ {Năm nhỏ} đến {Năm lớn}"
          + Từng năm: Mã = {Mã trường}_{Năm}, Tên = "Sổ gốc cấp bằng THCS trường {Tên trường} năm {Năm}"
        """
        lvl = level_info.get("level")
        entity_name = level_info.get("doc_entity_name", "")

        if lvl == "TRUONG":
            code = level_info.get("school_code") or level_info.get("doc_entity_code") or "TRUONG"
            if is_single_year or min_year == max_year:
                so_goc = f"{code}_{min_year}"
                doc_title = f"Sổ gốc cấp bằng THCS trường {entity_name} năm {min_year}"
            else:
                so_goc = f"{code}_{min_year}-{max_year}"
                doc_title = f"Sổ gốc cấp bằng THCS trường {entity_name} từ {min_year} đến {max_year}"
        else:
            code = level_info.get("doc_entity_code") or "DONVI"
            if is_single_year or min_year == max_year:
                so_goc = f"{code}_{min_year}"
                doc_title = f"Sổ gốc cấp bằng THCS {entity_name} năm {min_year}"
            else:
                so_goc = f"{code}_{min_year}-{max_year}"
                doc_title = f"Sổ gốc cấp bằng THCS {entity_name} từ {min_year} đến {max_year}"

        return so_goc, doc_title

    def _clean_cell_value(self, val: Any) -> Any:
        if val is None:
            return None
        if isinstance(val, str) and val.startswith("'"):
            return val[1:]
        return val

    def _write_student_rows(
        self,
        ws: Any,
        students: list[dict],
        override_code: str | None = None,
        override_name: str | None = None
    ) -> int:
        """
        Ghi danh sách học sinh vào sheet 28 cột:
        - Đánh lại STT 1..N
        - Áp dụng override Cột 27, 28 nếu có
        - Ép định dạng Text @ cho các cột text (bảo toàn số 0 ở đầu)
        """
        for idx, st in enumerate(students, start=1):
            row_data = [
                idx,
                self._clean_cell_value(st.get("so_hieu_bang")),
                self._clean_cell_value(st.get("so_vao_so")),
                st.get("ten_van_bang") or "Bằng tốt nghiệp trung học cơ sở",
                self._clean_cell_value(st.get("ho_ten")),
                self._clean_cell_value(st.get("ma_nguoi_hoc")),
                self._clean_cell_value(st.get("cccd")),
                self._clean_cell_value(st.get("ho_chieu")),
                self._clean_cell_value(st.get("ngay_sinh")),
                st.get("gioi_tinh"),
                st.get("dan_toc"),
                st.get("quoc_tich") or "Việt Nam",
                st.get("noi_sinh"),
                st.get("ten_truong"),
                st.get("nien_khoa"),
                st.get("diem_thi"),
                st.get("hoi_dong_thi"),
                st.get("dia_danh"),
                self._clean_cell_value(st.get("ngay_cap_bang")),
                st.get("co_quan_cap_bang"),
                st.get("chuc_danh_nguoi_ky"),
                st.get("nguoi_ky"),
                st.get("tinh_trang_van_bang") or "Cấp mới",
                st.get("ghi_chu"),
                st.get("trang_thai"),
                st.get("ly_do_loi"),
                self._clean_cell_value(override_code if override_code is not None else st.get("ma_don_vi")),
                self._clean_cell_value(override_name if override_name is not None else st.get("ten_don_vi"))
            ]

            ws.append(row_data)
            c_row = ws.max_row
            row_cells = ws[c_row]

            for c_idx in range(1, 29):
                cell = row_cells[c_idx - 1]
                if self.cached_student_fonts and c_idx - 1 < len(self.cached_student_fonts) and self.cached_student_fonts[c_idx - 1]:
                    cell.font = self.cached_student_fonts[c_idx - 1]

                if c_idx in TEXT_COLS_STUDENT:
                    cell.number_format = '@'
                    if cell.value is not None:
                        cell.data_type = 's'
                        cell.value = str(cell.value)

        # Freeze panes và auto-filter
        ws.freeze_panes = "A3"
        ws.auto_filter.ref = f"A2:AB{max(2, ws.max_row)}"
        return len(students)

    def _write_error_rows(
        self,
        ws: Any,
        error_students: list[dict],
        override_code: str | None = None,
        override_name: str | None = None
    ) -> int:
        """Ghi danh sách học sinh bị lỗi: 28 cột chuẩn + 3 cột vết nguồn gốc (AC..AE)."""
        # Thêm 3 cột vết nguồn
        ws.cell(row=2, column=29, value="Tệp nguồn gốc")
        ws.cell(row=2, column=30, value="Sheet nguồn")
        ws.cell(row=2, column=31, value="Dòng Excel gốc")

        for col_c in range(29, 32):
            c_ref = ws.cell(row=2, column=col_c)
            if self.cached_student_fonts and self.cached_student_fonts[0]:
                c_ref.font = copy.copy(self.cached_student_fonts[0])

        for idx, st in enumerate(error_students, start=1):
            row_data = [
                idx,
                self._clean_cell_value(st.get("so_hieu_bang")),
                self._clean_cell_value(st.get("so_vao_so")),
                st.get("ten_van_bang") or "Bằng tốt nghiệp trung học cơ sở",
                self._clean_cell_value(st.get("ho_ten")),
                self._clean_cell_value(st.get("ma_nguoi_hoc")),
                self._clean_cell_value(st.get("cccd")),
                self._clean_cell_value(st.get("ho_chieu")),
                self._clean_cell_value(st.get("ngay_sinh")),
                st.get("gioi_tinh"),
                st.get("dan_toc"),
                st.get("quoc_tich") or "Việt Nam",
                st.get("noi_sinh"),
                st.get("ten_truong"),
                st.get("nien_khoa"),
                st.get("diem_thi"),
                st.get("hoi_dong_thi"),
                st.get("dia_danh"),
                self._clean_cell_value(st.get("ngay_cap_bang")),
                st.get("co_quan_cap_bang"),
                st.get("chuc_danh_nguoi_ky"),
                st.get("nguoi_ky"),
                st.get("tinh_trang_van_bang"),
                st.get("ghi_chu"),
                st.get("trang_thai"),
                st.get("ly_do_loi"),
                self._clean_cell_value(override_code if override_code is not None else st.get("ma_don_vi")),
                self._clean_cell_value(override_name if override_name is not None else st.get("ten_don_vi")),
                st.get("source_file", ""),
                st.get("source_sheet", ""),
                st.get("source_row", "")
            ]
            ws.append(row_data)
            c_row = ws.max_row
            row_cells = ws[c_row]

            for c_idx in range(1, 32):
                cell = row_cells[c_idx - 1]
                if self.cached_student_fonts and c_idx - 1 < len(self.cached_student_fonts) and self.cached_student_fonts[c_idx - 1]:
                    cell.font = self.cached_student_fonts[c_idx - 1]
                if c_idx in TEXT_COLS_STUDENT or c_idx in (29, 30, 31):
                    cell.number_format = '@'
                    if cell.value is not None:
                        cell.data_type = 's'
                        cell.value = str(cell.value)

        ws.freeze_panes = "A3"
        ws.auto_filter.ref = f"A2:AE{max(2, ws.max_row)}"
        return len(error_students)

    def _write_doc_table_rows(
        self,
        ws: Any,
        doc_records: list[dict[str, Any]]
    ):
        """
        Ghi các dòng vào tệp Danh sách tài liệu số hóa:
        11 cột: STT, Mã đơn vị, Tên đơn vị, Mã sổ gốc, Tên tài liệu, Từ năm, Đến năm, Ngày lập sổ, Số trang, Tình trạng, Thông tư.
        """
        for idx, rec in enumerate(doc_records, start=1):
            row_vals = [
                idx,
                rec.get("ma_don_vi", ""),
                rec.get("ten_don_vi", ""),
                rec.get("ma_so_goc", ""),
                rec.get("ten_tai_lieu", ""),
                rec.get("tu_nam", ""),
                rec.get("den_nam", ""),
                rec.get("ngay_lap_so", ""),
                rec.get("so_trang", ""),
                rec.get("tinh_trang", "Nguyên vẹn"),
                rec.get("thong_tu", "")
            ]
            ws.append(row_vals)
            c_row = ws.max_row
            row_cells = ws[c_row]

            for c_idx in range(1, 12):
                cell = row_cells[c_idx - 1]
                if self.cached_doc_fonts and c_idx - 1 < len(self.cached_doc_fonts) and self.cached_doc_fonts[c_idx - 1]:
                    cell.font = self.cached_doc_fonts[c_idx - 1]
                if c_idx in TEXT_COLS_DOC:
                    cell.number_format = '@'
                    if cell.value is not None:
                        cell.data_type = 's'
                        cell.value = str(cell.value)

        ws.freeze_panes = "A3"
        ws.auto_filter.ref = f"A2:K{max(2, ws.max_row)}"

    def execute_export(
        self,
        db_manager: Any,
        output_dir: Path | str,
        branch: str = "branch_history",
        huyen_cu: str | None = None,
        ten_truong: str | None = None,
        xa_phuong: str | None = None,
        ten_don_vi: str | None = None,
        selected_cohorts: list[str] | None = None,
        export_error_files: bool = True,
        progress_callback: Callable[[int, int, str], None] | None = None
    ) -> dict[str, Any]:
        """
        Thực thi toàn bộ luồng xuất báo cáo theo yêu cầu:
        1. Phân tích tính liên tục của Niên khóa (1 năm / liên tiếp / rời rạc).
        2. Xác định Cấp xuất (Huyện cũ / Xã / Đơn vị / Trường) & Quy tắc Cột 27, 28.
        3. Xuất file(s) học sinh:
           - 1 năm hoặc liên tiếp -> 1 file Excel duy nhất
           - Rời rạc -> K file Excel theo từng năm
        4. Xuất 1 tệp Danh sách tài liệu số hóa DUY NHẤT:
           - 1 năm hoặc liên tiếp -> Ghi 1 dòng duy nhất ứng với file Excel đó
           - Rời rạc -> Ghi K dòng, mỗi dòng ứng với 1 file Excel đã xuất
        5. Xuất file lỗi (nếu có hồ sơ lỗi và được bật).
        """
        out_dir = Path(output_dir)
        out_dir.mkdir(parents=True, exist_ok=True)

        selected_cohorts = selected_cohorts or []
        analysis = analyze_cohorts(selected_cohorts)
        sorted_years = analysis["sorted_years"]
        is_single_year = analysis["is_single"]
        is_consecutive = analysis["is_consecutive"]
        min_year = analysis["min_year"]
        max_year = analysis["max_year"]

        if progress_callback:
            progress_callback(5, 100, "Đang truy vấn CSDL học sinh...")

        # 1. Truy vấn học sinh từ CSDL
        all_students = db_manager.query_students_for_export(
            huyen_cu=huyen_cu if branch == "branch_history" else None,
            ten_truong=ten_truong if branch == "branch_history" else None,
            xa_phuong=xa_phuong if branch == "branch_new_admin" else None,
            ten_don_vi=ten_don_vi if branch == "branch_new_admin" else None,
            cohorts=selected_cohorts if selected_cohorts else None,
            only_errors=False,
            sort_mode="stt_school_cohort"
        )

        sample_student = all_students[0] if all_students else {}
        level_info = self.determine_export_level(
            branch=branch,
            huyen_cu=huyen_cu,
            ten_truong=ten_truong,
            xa_phuong=xa_phuong,
            ten_don_vi=ten_don_vi,
            sample_student=sample_student
        )

        base_name = level_info["base_name"]
        override_code = level_info["override_unit_code"]
        override_name = level_info["override_unit_name"]
        curr_date_str = datetime.now().strftime("%d/%m/%Y")

        results = {
            "student_files": [],
            "doc_files": [],
            "error_files": [],
            "total_students": len(all_students),
            "level_info": level_info,
            "cohort_analysis": analysis
        }

        # Nhóm học sinh theo năm ra trường để phục vụ trường hợp rời rạc
        students_by_grad_year: dict[int, list[dict]] = {}
        for st in all_students:
            y = parse_cohort_grad_year(st.get("nien_khoa", "")) or 0
            students_by_grad_year.setdefault(y, []).append(st)

        # ======================================================================
        # TRƯỜNG HỢP 1: CHỌN 1 NĂM HOẶC NHIỀU NĂM LIÊN TIẾP (1 FILE DUY NHẤT)
        # ======================================================================
        if is_single_year or is_consecutive:
            suffix = f"{min_year}" if is_single_year or min_year == max_year else f"{min_year}_{max_year}"
            student_file_name = f"{base_name}_{suffix}.xlsx"
            student_file_path = out_dir / student_file_name

            if progress_callback:
                progress_callback(30, 100, f"Đang xuất file dữ liệu: {student_file_name}...")

            # 1. Xuất file dữ liệu học sinh
            wb_st, ws_st = self._create_student_workbook("Data")
            self._write_student_rows(ws_st, all_students, override_code, override_name)
            wb_st.save(student_file_path)
            wb_st.close()
            results["student_files"].append(student_file_path)

            # 2. Xuất file Danh sách tài liệu số hóa (CHỈ GHI 1 DÒNG DUY NHẤT)
            doc_file_name = f"{base_name}_{suffix}_DS_TL_So_Hoa.xlsx"
            doc_file_path = out_dir / doc_file_name

            if progress_callback:
                progress_callback(70, 100, f"Đang xuất file tài liệu số hóa: {doc_file_name}...")

            so_goc, doc_title = self.generate_doc_title_and_code(
                level_info, min_year, max_year, is_single_year
            )
            doc_record = {
                "ma_don_vi": override_code if override_code is not None else sample_student.get("ma_don_vi", ""),
                "ten_don_vi": override_name if override_name is not None else sample_student.get("ten_don_vi", ""),
                "ma_so_goc": so_goc,
                "ten_tai_lieu": doc_title,
                "tu_nam": str(min_year),
                "den_nam": str(max_year),
                "ngay_lap_so": curr_date_str,
                "so_trang": "",
                "tinh_trang": "Nguyên vẹn",
                "thong_tu": ""
            }

            wb_doc, ws_doc = self._create_doc_workbook("Data")
            self._write_doc_table_rows(ws_doc, [doc_record])
            wb_doc.save(doc_file_path)
            wb_doc.close()
            results["doc_files"].append(doc_file_path)

            # 3. Xuất file lỗi (nếu có)
            if export_error_files:
                err_students = [
                    s for s in all_students
                    if s.get("trang_thai") != "Đã số hóa đầy đủ, chính xác" or (s.get("ly_do_loi") and str(s.get("ly_do_loi")).strip())
                ]
                if err_students:
                    err_file_name = f"{base_name}_DS_Loi_{suffix}.xlsx"
                    err_file_path = out_dir / err_file_name
                    wb_err, ws_err = self._create_student_workbook("DS_Loi")
                    self._write_error_rows(ws_err, err_students, override_code, override_name)
                    wb_err.save(err_file_path)
                    wb_err.close()
                    results["error_files"].append(err_file_path)

        # ======================================================================
        # TRƯỜNG HỢP 2: CHỌN NHIỀU NĂM KHÔNG LIÊN TIẾP (RỜI RẠC)
        # -> Xuất mỗi năm 1 file Excel học sinh
        # -> Vẫn chỉ tạo 1 file Danh sách tài liệu số hóa duy nhất (mỗi dòng ứng với 1 file)
        # ======================================================================
        else:
            doc_records = []
            total_y = len(sorted_years)

            for idx_y, yr in enumerate(sorted_years, start=1):
                yr_students = students_by_grad_year.get(yr, [])
                pct = int(10 + (idx_y / total_y) * 60)
                if progress_callback:
                    progress_callback(pct, 100, f"Đang xuất file năm {yr} ({idx_y}/{total_y})...")

                student_file_name = f"{base_name}_{yr}.xlsx"
                student_file_path = out_dir / student_file_name

                # Xuất file học sinh của năm yr
                wb_st, ws_st = self._create_student_workbook("Data")
                self._write_student_rows(ws_st, yr_students, override_code, override_name)
                wb_st.save(student_file_path)
                wb_st.close()
                results["student_files"].append(student_file_path)

                # Chuẩn bị 1 dòng cho file tài liệu số hóa chung
                so_goc, doc_title = self.generate_doc_title_and_code(
                    level_info, yr, yr, is_single_year=True
                )
                doc_records.append({
                    "ma_don_vi": override_code if override_code is not None else sample_student.get("ma_don_vi", ""),
                    "ten_don_vi": override_name if override_name is not None else sample_student.get("ten_don_vi", ""),
                    "ma_so_goc": so_goc,
                    "ten_tai_lieu": doc_title,
                    "tu_nam": str(yr),
                    "den_nam": str(yr),
                    "ngay_lap_so": curr_date_str,
                    "so_trang": "",
                    "tinh_trang": "Nguyên vẹn",
                    "thong_tu": ""
                })

                # Xuất file lỗi cho năm yr nếu có
                if export_error_files:
                    yr_err = [
                        s for s in yr_students
                        if s.get("trang_thai") != "Đã số hóa đầy đủ, chính xác" or (s.get("ly_do_loi") and str(s.get("ly_do_loi")).strip())
                    ]
                    if yr_err:
                        err_file_name = f"{base_name}_DS_Loi_{yr}.xlsx"
                        err_file_path = out_dir / err_file_name
                        wb_err, ws_err = self._create_student_workbook("DS_Loi")
                        self._write_error_rows(ws_err, yr_err, override_code, override_name)
                        wb_err.save(err_file_path)
                        wb_err.close()
                        results["error_files"].append(err_file_path)

            # Xuất 1 TỆP TÀI LIỆU SỐ HÓA DUY NHẤT CHỨA NHIỀU DÒNG (mỗi dòng 1 năm)
            if progress_callback:
                progress_callback(85, 100, "Đang xuất file Danh sách tài liệu số hóa chung...")

            doc_file_name = f"{base_name}_DS_TL_So_Hoa.xlsx"
            doc_file_path = out_dir / doc_file_name
            wb_doc, ws_doc = self._create_doc_workbook("Data")
            self._write_doc_table_rows(ws_doc, doc_records)
            wb_doc.save(doc_file_path)
            wb_doc.close()
            results["doc_files"].append(doc_file_path)

        if progress_callback:
            progress_callback(100, 100, "Xuất dữ liệu hoàn tất 100%!")

        return results
