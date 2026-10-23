from __future__ import annotations

import re
from datetime import datetime
from pathlib import Path
from typing import Callable, Any

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from src.lookup_manager import LookupManager
from src.custom_excel_manager import _is_index_row
from src.normalizer import (
    clean_whitespace,
    vietnamese_title_case,
    format_date_to_text,
    normalize_ethnic,
    normalize_nationality,
    normalize_birthplace,
    auto_convert_legacy_encoding,
)


def normalize_cohort_str(raw: Any) -> str:
    """Chuẩn hoá chuỗi niên khoá, chuyển các dấu gạch ngang lạ về '-' chuẩn."""
    if raw is None:
        return ""
    s = str(raw).strip()
    s = re.sub(r'[\u2013\u2014\u2212]', '-', s)
    s = re.sub(r'\s*-\s*', '-', s)
    return s


def parse_and_validate_date(raw_val: Any) -> tuple[str, str | None, int | None]:
    """
    Phân tích và kiểm tra ngày tháng:
    Hỗ trợ loại bỏ dấu nháy đơn/kép ('19/05/1969, 19/'05/1969...) trước khi chuẩn hoá.
    Trả về: (formatted_val, status, year)
    - status: None (hợp lệ), 'LỖI', hoặc 'CẢNH BÁO'
    - year: Năm (int) nếu phân tích được, hoặc None
    """
    if raw_val is None or str(raw_val).strip() == "":
        return "", "TRỐNG", None

    # Loại bỏ mọi dấu nháy rải rác trước khi chuẩn hoá (ví dụ '19/05/1969 hay 19/'05/1969)
    if isinstance(raw_val, str):
        raw_val = re.sub(r"['\"`’‘“”]", "", raw_val).strip()

    # Định dạng qua format_date_to_text
    formatted = format_date_to_text(raw_val)
    if not formatted or formatted.strip() == "":
        return "", "LỖI", None

    formatted = formatted.strip()

    # Kiểm tra khớp cấu trúc dd/MM/yyyy (ngày và tháng 1 hoặc 2 chữ số, năm 4 chữ số)
    m = re.match(r"^(\d{1,2})/(\d{1,2})/(\d{4})$", formatted)
    if not m:
        return formatted, "LỖI", None

    day = int(m.group(1))
    month = int(m.group(2))
    year = int(m.group(3))
    std_formatted = f"{day:02d}/{month:02d}/{year:04d}"

    # Kiểm tra tính logic ngày tháng thực tế (lịch chuẩn)
    try:
        dt = datetime(year, month, day)
        return std_formatted, None, dt.year
    except ValueError:
        # Sai logic ngày tháng (ví dụ 30/02/2011, 05/31/1999...) -> CẢNH BÁO
        return std_formatted, "CẢNH BÁO", year


def validate_student_record(
    row_28: list[Any],
    lookup_mgr: LookupManager | None = None
) -> tuple[list[Any], dict[str, str | None], str]:
    """
    Chuẩn hoá dữ liệu 28 cột và kiểm tra lỗi / cảnh báo cho 1 bản ghi học sinh:
    - row_28: Danh sách 28 ô dữ liệu (từ Cột 1 đến Cột 28)
    Trả về:
    - norm_row_28: Dữ liệu 28 cột đã chuẩn hoá
    - col_status: dict chứa trạng thái của 4 cột ('ho_ten', 'ngay_sinh', 'nien_khoa', 'ngay_cap')
    - summary_status: 'LỖI' | 'CẢNH BÁO' | ''
    """
    # Đảm bảo đủ 28 phần tử
    data = list(row_28)
    if len(data) < 28:
        data.extend([""] * (28 - len(data)))
    data = data[:28]

    # BƯỚC 1: CHUẨN HOÁ DỮ LIỆU
    for i in range(28):
        val = data[i]
        if isinstance(val, str):
            data[i] = auto_convert_legacy_encoding(val)

    # Cột 2 (B, index 1): Số hiệu bằng -> làm sạch khoảng trắng
    data[1] = clean_whitespace(data[1])

    # Cột 3 (C, index 2): Số vào sổ gốc cấp văn bằng -> làm sạch khoảng trắng
    data[2] = clean_whitespace(data[2])

    # Cột 4 (D, index 3): Tên văn bằng -> Luôn tự động điền 'Bằng tốt nghiệp trung học cơ sở'
    data[3] = "Bằng tốt nghiệp trung học cơ sở"

    # Cột 5 (E, index 4): Họ và tên -> Title Case
    name_raw = clean_whitespace(data[4])
    data[4] = vietnamese_title_case(name_raw) if name_raw else ""

    # Cột 6 (F, index 5): Mã người học -> làm sạch khoảng trắng
    data[5] = clean_whitespace(data[5])

    # Cột 7 (G, index 6): Số định danh cá nhân (CCCD) -> làm sạch khoảng trắng
    cccd_clean = clean_whitespace(data[6])
    if isinstance(cccd_clean, str):
        cccd_clean = re.sub(r'\s+', '', cccd_clean)
    data[6] = cccd_clean

    # Cột 9 (I, index 8): Ngày sinh (loại bỏ dấu nháy trước khi chuẩn hoá)
    dob_raw = data[8]
    dob_formatted, dob_status, birth_year = parse_and_validate_date(dob_raw)
    data[8] = dob_formatted

    # Cột 10 (J, index 9): Giới tính
    if data[9] is not None and str(data[9]).strip() != "":
        g_str = str(data[9]).strip()
        data[9] = lookup_mgr.map_gender(g_str) if lookup_mgr else clean_whitespace(g_str)

    # Cột 11 (K, index 10): Dân tộc
    if data[10] is not None and str(data[10]).strip() != "":
        data[10] = normalize_ethnic(str(data[10]), lookup_mgr)

    # Cột 12 (L, index 11): Quốc tịch
    if data[11] is not None and str(data[11]).strip() != "":
        data[11] = normalize_nationality(str(data[11]), lookup_mgr)

    # Cột 13 (M, index 12): Nơi sinh
    if data[12] is not None and str(data[12]).strip() != "":
        data[12] = normalize_birthplace(str(data[12]), lookup_mgr)

    # Cột 15 (O, index 14): Niên khoá
    cohort_raw = normalize_cohort_str(data[14])
    data[14] = cohort_raw

    # Cột 19 (S, index 18): Ngày cấp bằng
    grad_date_raw = data[18]
    grad_formatted, grad_status, grad_year = parse_and_validate_date(grad_date_raw)
    data[18] = grad_formatted

    # Cột 25 (Y, index 24): Trạng thái số hóa
    if not data[24] or str(data[24]).strip() == "":
        data[24] = "Đã số hoá đầy đủ, chính xác"

    # ==============================================================
    # BƯỚC 2: KIỂM TRA LỖI & CẢNH BÁO
    # ==============================================================
    col_status: dict[str, str | None] = {
        "ho_ten": None,
        "ngay_sinh": None,
        "nien_khoa": None,
        "ngay_cap": None
    }

    # 1. Họ và tên: Bị trống -> Lỗi
    if not data[4] or str(data[4]).strip() == "":
        col_status["ho_ten"] = "LỖI"

    # 2. Ngày sinh chuẩn dạng dd/MM/yyyy:
    # - trống -> Lỗi
    # - không đúng dạng dd/MM/yyyy -> Lỗi
    # - sai logic ngày sinh -> Cảnh báo
    if dob_status == "TRỐNG" or dob_status == "LỖI":
        col_status["ngay_sinh"] = "LỖI"
    elif dob_status == "CẢNH BÁO":
        col_status["ngay_sinh"] = "CẢNH BÁO"

    # 3. Cột niên khoá chuẩn dạng YYYY-YYYY:
    # - trống -> Lỗi
    # - không có dạng YYYY-YYYY -> Lỗi
    # - Y_ra - Y_vao khác 4 -> Lỗi
    # - Nếu Y_ra khác Năm sinh + 15 -> Cảnh báo
    y_ra: int | None = None
    if not cohort_raw:
        col_status["nien_khoa"] = "LỖI"
    else:
        m_cohort = re.match(r"^(\d{4})-(\d{4})$", cohort_raw)
        if not m_cohort:
            col_status["nien_khoa"] = "LỖI"
        else:
            y_vao = int(m_cohort.group(1))
            y_ra = int(m_cohort.group(2))
            if y_ra - y_vao != 4:
                col_status["nien_khoa"] = "LỖI"
            elif birth_year is not None and y_ra != (birth_year + 15):
                col_status["nien_khoa"] = "CẢNH BÁO"

    # 4. Cột ngày cấp bằng chuẩn dạng dd/MM/yyyy:
    # - trống -> Cảnh báo
    # - không đúng dạng dd/MM/yyyy -> Lỗi
    # - sai logic ngày -> Cảnh báo
    # - Năm khác Y_ra -> Cảnh báo
    if grad_status == "TRỐNG":
        col_status["ngay_cap"] = "CẢNH BÁO"
    elif grad_status == "LỖI":
        col_status["ngay_cap"] = "LỖI"
    elif grad_status == "CẢNH BÁO":
        col_status["ngay_cap"] = "CẢNH BÁO"
    else:
        # Nếu ngày cấp hợp lệ, kiểm tra năm khác Y_ra
        if y_ra is not None and grad_year is not None and grad_year != y_ra:
            col_status["ngay_cap"] = "CẢNH BÁO"

    # BƯỚC 3: XÁC ĐỊNH TRẠNG THÁI TỔNG HỢP CỦA DÒNG
    all_statuses = list(col_status.values())
    if "LỖI" in all_statuses:
        summary_status = "LỖI"
    elif "CẢNH BÁO" in all_statuses:
        summary_status = "CẢNH BÁO"
    else:
        summary_status = ""

    return data, col_status, summary_status


class DataInspector:
    """
    Bộ kiểm tra và chuẩn hoá dữ liệu theo biểu mẫu chuẩn 28 cột:
    - Đầu vào: Tệp Excel có cấu trúc tương đương 28 cột chuẩn
    - Chuẩn hoá dữ liệu 28 cột
    - Bổ sung 5 cột kiểm tra ở phía sau (từ Cột 29 / AC trở đi):
        + Cột 29 (AC): [TỔNG HỢP]
        + Cột 30 (AD): [KT] Họ tên
        + Cột 31 (AE): [KT] Ngày sinh
        + Cột 32 (AF): [KT] Niên khoá
        + Cột 33 (AG): [KT] Ngày cấp bằng
    - Tô màu trực quan các ô LỖI (Đỏ nhạt) và CẢNH BÁO (Vàng nhạt).
    """

    # Màu tô chuẩn cho Cell Highlight
    FILL_ERROR = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    FONT_ERROR = Font(name="Times New Roman", size=11, color="9C0006", bold=True)

    FILL_WARNING = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
    FONT_WARNING = Font(name="Times New Roman", size=11, color="9C6500", bold=True)

    FILL_HEADER_KT = PatternFill(start_color="E0EBF5", end_color="E0EBF5", fill_type="solid")
    FONT_HEADER_KT = Font(name="Times New Roman", size=11, bold=True, color="003366")

    BORDER_THIN = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )

    def __init__(self, lookup_mgr: LookupManager | None = None):
        self.lookup_mgr = lookup_mgr or LookupManager()

    def detect_header_row_and_columns(self, ws, max_search_rows: int = 15) -> tuple[int, dict[str, int]]:
        """
        Quét và nhận diện dòng tiêu đề chính xác theo mẫu chuẩn 28 cột (Mau_DS_Cap_Bang_THCS.xlsx).
        Đồng thời xác định động vị trí của 4 cột kiểm tra:
        - Họ và tên (chuẩn Cột 5)
        - Ngày sinh (chuẩn Cột 9)
        - Niên khoá (chuẩn Cột 15)
        - Ngày cấp bằng (chuẩn Cột 19)
        """
        best_row = 2
        max_score = -1
        best_col_map = {
            "ho_ten": 5,
            "ngay_sinh": 9,
            "nien_khoa": 15,
            "ngay_cap": 19
        }

        # Danh sách từ khóa đặc trưng nhận diện 28 cột của mẫu chuẩn THCS (Mau_DS_Cap_Bang_THCS.xlsx)
        STANDARD_28_KEYWORDS = (
            "stt", "số hiệu bằng", "số vào sổ gốc cấp văn bằng", "số vào sổ", "tên văn bằng",
            "họ, chữ đệm và tên", "họ và tên", "mã người học", "số định danh cá nhân",
            "số định danh", "hộ chiếu", "ngày, tháng, năm sinh", "ngày sinh", "giới tính",
            "dân tộc", "quốc tịch", "nơi sinh", "tên trường đã học", "niên khóa", "niên khoá",
            "điểm các môn thi", "hội đồng thi", "địa danh nơi cơ quan cấp bằng đặt trụ sở",
            "ngày tháng năm cấp bằng", "ngày cấp bằng", "tên cơ quan cấp bằng",
            "chức danh người ký bằng", "họ, chữ đệm, tên người ký bằng", "tình trạng văn bằng",
            "ghi chú", "trạng thái số hóa", "trạng thái số hoá",
            "thông tin chưa bảo đảm đầy đủ, chính xác và nguyên nhân", "mã đơn vị", "tên đơn vị"
        )

        for r in range(1, min(max_search_rows + 1, ws.max_row + 1)):
            row_vals = [str(ws.cell(r, c).value or "").strip().lower() for c in range(1, min(ws.max_column + 1, 35))]
            row_text = " ".join(row_vals)
            if not row_text:
                continue

            score = sum(1 for kw in STANDARD_28_KEYWORDS if kw in row_text)
            if score > max_score and score >= 3:
                max_score = score
                best_row = r

                # Dò tìm vị trí cụ thể của 4 cột cần kiểm tra trên dòng này
                current_col_map = dict(best_col_map)
                for c_idx, val in enumerate(row_vals, 1):
                    if not val:
                        continue
                    if any(k in val for k in ("họ, chữ đệm", "họ và tên", "họ tên")) and "người ký" not in val:
                        current_col_map["ho_ten"] = c_idx
                    elif any(k in val for k in ("ngày, tháng, năm sinh", "ngày sinh", "năm sinh")):
                        current_col_map["ngay_sinh"] = c_idx
                    elif any(k in val for k in ("niên khóa", "niên khoá")):
                        current_col_map["nien_khoa"] = c_idx
                    elif any(k in val for k in ("ngày tháng năm cấp bằng", "ngày cấp bằng", "ngày cấp")):
                        current_col_map["ngay_cap"] = c_idx
                best_col_map = current_col_map

        return best_row, best_col_map

    def detect_header_row(self, ws, max_search_rows: int = 15) -> int:
        hdr_row, _ = self.detect_header_row_and_columns(ws, max_search_rows)
        return hdr_row

    def process_file(
        self,
        input_file: Path | str,
        output_file: Path | str | None = None,
        progress_callback: Callable[[int, int, str], None] | None = None,
        cancel_token: Callable[[], bool] | None = None
    ) -> dict[str, Any]:
        """
        Kiểm tra và chuẩn hoá file Excel, xuất file kết quả kèm 5 cột kiểm tra và tô màu ô.
        """
        inp_p = Path(input_file)
        if not inp_p.exists():
            return {"status": "error", "message": f"Tệp không tồn tại: {inp_p.name}"}

        out_p = Path(output_file) if output_file else inp_p.parent / f"{inp_p.stem}_DaKiemTra.xlsx"

        try:
            wb = openpyxl.load_workbook(inp_p, data_only=False)
        except Exception as e:
            return {"status": "error", "message": f"Không thể đọc tệp Excel: {e}"}

        ws = wb.active
        hdr_row, col_idx_map = self.detect_header_row_and_columns(ws)
        data_start_row = hdr_row + 1

        # Bỏ qua dòng chú thích số thứ tự cột (1), (2), (3)... nếu có ngay sau tiêu đề
        if data_start_row <= ws.max_row:
            next_row_vals = [ws.cell(data_start_row, c).value for c in range(1, min(ws.max_column + 1, 30))]
            if _is_index_row(next_row_vals):
                data_start_row += 1

        # Tạo tiêu đề cho 5 cột kiểm tra ở dòng tiêu đề
        # Cột 29 (AC), 30 (AD), 31 (AE), 32 (AF), 33 (AG)
        kt_headers = [
            (29, "[TỔNG HỢP]"),
            (30, "[KT] Họ tên"),
            (31, "[KT] Ngày sinh"),
            (32, "[KT] Niên khoá"),
            (33, "[KT] Ngày cấp bằng"),
        ]
        for col_idx, title in kt_headers:
            cell = ws.cell(hdr_row, col_idx, title)
            cell.fill = self.FILL_HEADER_KT
            cell.font = self.FONT_HEADER_KT
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = self.BORDER_THIN
            col_letter = get_column_letter(col_idx)
            ws.column_dimensions[col_letter].width = 16

        total_rows = ws.max_row - data_start_row + 1
        if total_rows <= 0:
            wb.close()
            return {"status": "warning", "message": "Bảng không có dòng dữ liệu học sinh nào."}

        error_count = 0
        warning_count = 0
        valid_count = 0
        processed_rows = 0

        # Ánh xạ cột tương ứng với các trường trong 28 cột (xác định động hoặc chuẩn)
        COL_IDX_MAP = col_idx_map
        COL_KT_MAP = {
            "ho_ten": 30,     # Cột AD
            "ngay_sinh": 31,  # Cột AE
            "nien_khoa": 32,  # Cột AF
            "ngay_cap": 33    # Cột AG
        }

        align_center = Alignment(horizontal="center", vertical="center")

        for row_cells in ws.iter_rows(min_row=data_start_row, max_row=ws.max_row, min_col=1, max_col=33):
            if cancel_token and cancel_token():
                wb.close()
                return {"status": "cancelled", "message": "Người dùng đã huỷ thao tác."}

            processed_rows += 1
            # Đọc 28 ô đầu của dòng
            raw_vals = [row_cells[c].value for c in range(28)]

            # Kiểm tra dòng trống hoàn toàn
            if all(v is None or str(v).strip() == "" for v in raw_vals):
                continue

            # Chuẩn hoá & Kiểm tra
            norm_vals, col_status, summary_status = validate_student_record(raw_vals, self.lookup_mgr)

            # Cập nhật lại giá trị chuẩn hoá vào 28 cột dữ liệu
            for c in range(28):
                row_cells[c].value = norm_vals[c]

            # Ghi kết quả vào 5 cột kiểm tra
            # 1. Cột 29: [TỔNG HỢP] (index 28)
            c_sum = row_cells[28]
            c_sum.value = summary_status
            c_sum.alignment = align_center
            c_sum.border = self.BORDER_THIN
            if summary_status == "LỖI":
                c_sum.fill = self.FILL_ERROR
                c_sum.font = self.FONT_ERROR
                error_count += 1
            elif summary_status == "CẢNH BÁO":
                c_sum.fill = self.FILL_WARNING
                c_sum.font = self.FONT_WARNING
                warning_count += 1
            else:
                valid_count += 1

            # 2. Bốn cột kiểm tra chi tiết & Tô màu trực tiếp ô nguồn
            for field, st in col_status.items():
                kt_idx = COL_KT_MAP[field] - 1
                src_idx = COL_IDX_MAP[field] - 1

                cell_kt = row_cells[kt_idx]
                cell_kt.value = st or ""
                cell_kt.alignment = align_center
                cell_kt.border = self.BORDER_THIN

                cell_src = row_cells[src_idx]

                if st == "LỖI":
                    cell_kt.fill = self.FILL_ERROR
                    cell_kt.font = self.FONT_ERROR
                    cell_src.fill = self.FILL_ERROR
                    cell_src.font = self.FONT_ERROR
                elif st == "CẢNH BÁO":
                    cell_kt.fill = self.FILL_WARNING
                    cell_kt.font = self.FONT_WARNING
                    cell_src.fill = self.FILL_WARNING
                    cell_src.font = self.FONT_WARNING

            if progress_callback and (processed_rows % 500 == 0 or processed_rows == total_rows):
                progress_callback(processed_rows, total_rows, f"Đang kiểm tra dòng {processed_rows:,}/{total_rows:,}...")

        # Lưu file kết quả
        try:
            wb.save(out_p)
            wb.close()
        except Exception as e:
            return {"status": "error", "message": f"Không thể lưu tệp kết quả: {e}"}

        return {
            "status": "success",
            "output_file": out_p,
            "total_rows": processed_rows,
            "error_count": error_count,
            "warning_count": warning_count,
            "valid_count": valid_count
        }
