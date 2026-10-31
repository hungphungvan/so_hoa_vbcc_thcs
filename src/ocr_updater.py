from __future__ import annotations

from concurrent.futures import ThreadPoolExecutor
import json
import os
import queue
import re
import threading
import time
import unicodedata
from pathlib import Path
from typing import Any, Callable

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from src.custom_excel_manager import detect_header_and_columns
from src.normalizer import clean_whitespace, format_date_to_text, vietnamese_title_case
from src.ocr_processor import (
    DEFAULT_BASE_URL,
    DEFAULT_MODEL_NAME,
    OCRConfigManager,
    OCRProcessor,
    extract_table_from_image,
    parse_ocr_json_response,
)
from src.lookup_manager import get_persistent_data_dir
from src.smart_merger import TARGET_COLUMNS, normalize_text_for_match

# Định nghĩa các tiêu chí so khớp dòng học sinh
AVAILABLE_MATCH_FIELDS = [
    {"key": "stt", "label": "STT (Cột 1)", "col_idx": 1},
    {"key": "so_hieu_bang", "label": "Số hiệu bằng (Cột 2)", "col_idx": 2},
    {"key": "so_vao_so", "label": "Số vào sổ (Cột 3)", "col_idx": 3},
    {"key": "ho_ten", "label": "Họ và tên (Cột 5)", "col_idx": 5},
    {"key": "ngay_sinh", "label": "Ngày sinh (Cột 9)", "col_idx": 9},
    {"key": "ten_truong", "label": "Trường đang học (Cột 14)", "col_idx": 14},
]

# Ánh xạ từ tên trường dữ liệu sang chỉ số cột chuẩn 28 (1-based)
FIELD_TO_COL_IDX: dict[str, int] = {
    "stt": 1,
    "so_hieu_bang": 2,
    "so_vao_so": 3,
    "ten_van_bang": 4,
    "ho_ten": 5,
    "ma_nguoi_hoc": 6,
    "cccd": 7,
    "ho_chieu": 8,
    "ngay_sinh": 9,
    "gioi_tinh": 10,
    "dan_toc": 11,
    "quoc_tich": 12,
    "noi_sinh": 13,
    "ten_truong": 14,
    "nien_khoa": 15,
    "diem_thi": 16,
    "xep_loai": 16,
    "hoi_dong_thi": 17,
    "dia_danh": 18,
    "ngay_cap_bang": 19,
    "co_quan_cap_bang": 20,
    "chuc_danh_nguoi_ky": 21,
    "nguoi_ky": 22,
    "tinh_trang_van_bang": 23,
    "ghi_chu": 24,
    "trang_thai": 25,
    "ly_do_loi": 26,
    "ma_don_vi": 27,
    "ten_don_vi": 28,
}

# Từ đồng nghĩa và tên tiếng Việt của các cột
FIELD_SYNONYMS: dict[str, int] = {
    "so thu tu": 1,
    "stt": 1,
    "so hieu": 2,
    "so hieu bang": 2,
    "so vao so": 3,
    "so vao so goc": 3,
    "vao so": 3,
    "ten van bang": 4,
    "ho ten": 5,
    "ho va ten": 5,
    "ten": 5,
    "ho, chu dem va ten": 5,
    "ma nguoi hoc": 6,
    "cccd": 7,
    "dinh danh": 7,
    "so dinh danh": 7,
    "ho chieu": 8,
    "ngay sinh": 9,
    "nam sinh": 9,
    "ngay thang nam sinh": 9,
    "gioi tinh": 10,
    "gt": 10,
    "dan toc": 11,
    "quoc tich": 12,
    "noi sinh": 13,
    "ten truong": 14,
    "truong": 14,
    "truong dang hoc": 14,
    "nien khoa": 15,
    "diem thi": 16,
    "xep loai": 16,
    "xep loai tn": 16,
    "hoi dong thi": 17,
    "dia danh": 18,
    "ngay cap bang": 19,
    "co quan cap bang": 20,
    "chuc danh nguoi ky": 21,
    "nguoi ky": 22,
    "tinh trang": 23,
    "ghi chu": 24,
    "trang thai": 25,
}

DEFAULT_PATCH_PROMPT = """Bạn là hệ thống OCR tốc độ cao chuyên lọc các dòng CÓ CHỮ SỬA ĐỔI MÀU ĐỎ trong danh sách tốt nghiệp THCS.

QUY TẮC LỌC TỐC ĐỘ CAO (BỎ QUA NHANH):
1. BỎ QUA DÒNG VIẾT THÊM (BỔ SUNG): Tuyệt đối KHÔNG lấy các dòng học sinh viết tay bổ sung thêm ở cuối bảng hay chèn mới. CHỈ quan tâm các dòng in máy có sửa đổi bằng mực đỏ.
2. BỎ QUA DẤU TÍCH KIỂM TRA (dấu "v", "✓", chấm tròn đánh dấu đã kiểm tra): Các dấu tích v màu đỏ hoặc bút bi dùng để đánh dấu kiểm đếm KHÔNG phải là sửa đổi. Nếu một dòng chỉ có dấu tích "v" mà không gạch xóa hay viết chữ sửa đổi thì BỎ QUA NGAY.
3. NẾU DÒNG CHỈ CÓ CHỮ IN MÁY BÌNH THƯỜNG: Lập tức bỏ qua ngay, không đọc chi tiết.

QUY TẮC TRÍCH XUẤT 5 TRƯỜNG DỮ LIỆU SỬA ĐỔI (MÀU ĐỎ) Ở MỖI DÒNG CÓ SỬA:
Chỉ lấy dữ liệu sửa đổi thuộc 5 mục sau:
1. Họ và tên ("ho_ten"):
   - Nếu gạch đỏ cả họ tên hoặc chỉ gạch đỏ một phần (họ, đệm, hoặc tên) và viết chữ đỏ sửa lại: hãy ghép phần chữ đỏ mới với phần chữ in không bị gạch để tạo thành HỌ VÀ TÊN ĐẦY ĐỦ sau khi sửa (đồng thời ghi nhận tên in máy ban đầu vào "gia_tri_cu").
   - CHÚ Ý: Chữ viết tay màu đỏ sửa họ và tên có thể viết lệch sang cột bên cạnh, viết ở lề hoặc phía trên/dưới dòng do thiếu chỗ. Hãy nhìn theo vị trí gạch đỏ ở cột Họ tên để lấy đúng chữ đỏ tương ứng.
2. Ngày tháng năm sinh ("ngay_sinh"):
   - Nếu gạch đỏ toàn bộ hoặc chỉ một phần (chỉ gạch ngày, chỉ gạch tháng, hoặc chỉ gạch năm sinh) và viết số màu đỏ sửa lại: hãy ghép phần số đỏ mới với phần số in còn lại để tạo thành NGÀY THÁNG NĂM SINH ĐẦY ĐỦ (dạng dd/mm/yyyy).
   - CHÚ Ý: Số sửa ngày/tháng/năm sinh màu đỏ có thể bị viết tràn sang cột lân cận (như cột Giới tính, Nơi sinh). Hãy chú ý vết gạch đỏ ở cột Ngày sinh để ghép đúng.
3. Giới tính ("gioi_tinh"):
   - Nếu chữ giới tính in sẵn bị gạch đỏ (kể cả không ghi chữ mới bên cạnh), nghĩa là đổi sang giới tính còn lại: ví dụ gạch chéo/gạch ngang chữ "Nam" tức là sửa thành "Nữ" (gia_tri_cu: "Nam", gia_tri_moi: "Nữ"); ngược lại gạch chữ "Nữ" tức là sửa thành "Nam" (gia_tri_cu: "Nữ", gia_tri_moi: "Nam").
4. Dân tộc ("dan_toc"):
   - Thường gạch đỏ dân tộc in sẵn và ghi tên dân tộc mới bằng chữ đỏ ngay bên cạnh (ví dụ gạch "Kinh" ghi "Tày" bên cạnh).
5. Nơi sinh ("noi_sinh"):
   - Thường gạch đỏ nơi sinh in sẵn và ghi nơi sinh mới bằng chữ đỏ bên cạnh.

ĐẦU RA BẮT BUỘC:
- Nếu cả trang KHÔNG có dòng nào sửa đổi thuộc 5 mục trên, trả về ngay lập tức: []
- Nếu có, ở mỗi dòng có sửa đổi BẮT BUỘC đọc luôn cả 3 thông tin nhận diện trên ảnh gồm: Số thứ tự ("stt"), Họ và tên ("ho_ten"), và Ngày tháng năm sinh ("ngay_sinh") của dòng đó (kể cả khi dòng đó chỉ sửa giới tính, dân tộc hay nơi sinh), rồi trả về DUY NHẤT 1 mảng JSON chuẩn (không dùng ```json, không giải thích):
[
  {
    "loai": "SUA_DOI",
    "stt": <số thứ tự in trên dòng đó>,
    "so_hieu_bang": "<số hiệu bằng in trên dòng đó nếu có>",
    "so_vao_so": "<số vào sổ in trên dòng đó nếu có>",
    "ho_ten": "<họ và tên đầy đủ của học sinh trên ảnh (đã ghép phần sửa nếu bị sửa tên)>",
    "ngay_sinh": "<ngày tháng năm sinh đầy đủ dd/mm/yyyy trên ảnh (đã ghép phần sửa nếu bị sửa ngày sinh)>",
    "gioi_tinh": "<Nam hoặc Nữ sau khi sửa>",
    "dan_toc": "<dân tộc sau khi sửa>",
    "noi_sinh": "<nơi sinh sau khi sửa>",
    "cac_truong_sua": [
      {
        "ten_truong": "<chỉ nhận 1 trong 5 giá trị: ho_ten, ngay_sinh, gioi_tinh, dan_toc, noi_sinh>",
        "gia_tri_cu": "<giá trị in máy đầy đủ ban đầu trước khi bị gạch đỏ>",
        "gia_tri_moi": "<giá trị đầy đủ mới sau khi đã áp dụng chữ đỏ sửa đổi>"
      }
    ],
    "can_kiem_tra": false,
    "ly_do_nghi_ngo": null
  }
]"""


def normalize_match_val(val: Any, field_key: str) -> str:
    """Chuẩn hóa giá trị để so khớp chính xác giữa OCR và file Excel."""
    if val is None:
        return ""
    val_str = str(val).strip()
    if not val_str or val_str.lower() in ("none", "nan", "null"):
        return ""

    if field_key == "stt":
        val_clean = re.sub(r"\.0+$", "", val_str)
        m_num = re.search(r"\d+", val_clean)
        if m_num:
            return str(int(m_num.group(0)))
        return val_clean.lower()

    if field_key == "ngay_sinh":
        norm_date = format_date_to_text(val_str)
        return norm_date or val_str

    if field_key == "ho_ten":
        val_norm = unicodedata.normalize("NFC", val_str)
        return re.sub(r"\s+", " ", val_norm).strip().lower()

    if field_key in ("so_hieu_bang", "so_vao_so"):
        return re.sub(r"\s+", "", val_str).lower()

    if field_key == "ten_truong":
        val_norm = unicodedata.normalize("NFC", val_str)
        return re.sub(r"\s+", " ", val_norm).strip().lower()

    return re.sub(r"\s+", " ", val_str).strip().lower()


class OCRColumnMappingManager:
    """
    Quản lý cấu hình ánh xạ cột từ tên cột trên ảnh sổ sang 28 cột chuẩn.
    Lưu trữ bền vững trong file JSON (ocr_column_mappings.json).
    Cho phép người dùng tạo Profile riêng cho từng loại sổ / trường học và tái sử dụng.
    """
    def __init__(self, config_path: Path | str | None = None):
        if config_path:
            self.config_path = Path(config_path)
        else:
            self.config_path = get_persistent_data_dir() / "ocr_column_mappings.json"
        self._cache: dict[str, Any] | None = None

    def get_default_mapping(self) -> dict[str, int]:
        return {
            "STT": 1,
            "Số thứ tự": 1,
            "Số hiệu bằng": 2,
            "Số hiệu": 2,
            "Số bằng": 2,
            "Số vào sổ gốc": 3,
            "Số vào sổ": 3,
            "Vào sổ": 3,
            "Tên văn bằng": 4,
            "Họ, chữ đệm và tên": 5,
            "Họ và tên": 5,
            "Họ tên": 5,
            "Tên học sinh": 5,
            "Mã người học": 6,
            "Số định danh cá nhân": 7,
            "CCCD": 7,
            "Hộ chiếu": 8,
            "Ngày, tháng, năm sinh": 9,
            "Ngày sinh": 9,
            "Năm sinh": 9,
            "Giới tính": 10,
            "GT": 10,
            "Dân tộc": 11,
            "Quốc tịch": 12,
            "Nơi sinh": 13,
            "Quê quán": 13,
            "Tên trường đã học": 14,
            "Tên trường": 14,
            "Trường": 14,
            "Niên khóa": 15,
            "Điểm các môn thi": 16,
            "Điểm thi": 16,
            "Xếp loại": 16,
            "Hội đồng thi": 17,
            "HĐT": 17,
            "Địa danh": 18,
            "Ngày tháng năm cấp bằng": 19,
            "Ngày cấp bằng": 19,
            "Tên cơ quan cấp bằng": 20,
            "Chức danh người ký": 21,
            "Người ký bằng": 22,
            "Tình trạng văn bằng": 23,
            "Ghi chú": 24,
            "Trạng thái số hóa": 25,
            "Mã đơn vị": 27,
            "Tên đơn vị": 28,
        }

    def load_data(self) -> dict[str, Any]:
        if self._cache is not None:
            return self._cache
        if self.config_path.is_file():
            try:
                data = json.loads(self.config_path.read_text(encoding="utf-8"))
                if isinstance(data, dict) and "profiles" in data:
                    self._cache = data
                    return data
            except Exception:
                pass

        default_map = self.get_default_mapping()
        data = {
            "active_profile": "Mặc định (Chuẩn 28 cột)",
            "profiles": {
                "Mặc định (Chuẩn 28 cột)": {
                    "name": "Mặc định (Chuẩn 28 cột)",
                    "mapping": default_map
                }
            }
        }
        self._cache = data
        self.save_data(data)
        return data

    def save_data(self, data: dict[str, Any]) -> None:
        self._cache = data
        try:
            self.config_path.parent.mkdir(parents=True, exist_ok=True)
            self.config_path.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
        except Exception:
            pass

    def list_profiles(self) -> list[str]:
        data = self.load_data()
        return list(data.get("profiles", {}).keys())

    def get_active_profile_name(self) -> str:
        data = self.load_data()
        return data.get("active_profile", "Mặc định (Chuẩn 28 cột)")

    def set_active_profile_name(self, name: str) -> None:
        data = self.load_data()
        if name in data.get("profiles", {}):
            data["active_profile"] = name
            self.save_data(data)

    def get_mapping(self, profile_name: str | None = None) -> dict[str, int]:
        data = self.load_data()
        p_name = profile_name or data.get("active_profile") or "Mặc định (Chuẩn 28 cột)"
        prof = data.get("profiles", {}).get(p_name)
        if prof and isinstance(prof.get("mapping"), dict):
            return prof["mapping"]
        return self.get_default_mapping()

    def save_profile(self, profile_name: str, mapping: dict[str, int]) -> None:
        data = self.load_data()
        if "profiles" not in data:
            data["profiles"] = {}
        data["profiles"][profile_name] = {
            "name": profile_name,
            "mapping": mapping
        }
        data["active_profile"] = profile_name
        self.save_data(data)

    def delete_profile(self, profile_name: str) -> bool:
        data = self.load_data()
        profiles = data.get("profiles", {})
        if profile_name in profiles and len(profiles) > 1:
            del profiles[profile_name]
            if data.get("active_profile") == profile_name:
                data["active_profile"] = list(profiles.keys())[0]
            self.save_data(data)
            return True
        return False


def resolve_field_to_col_idx(field_name: str, custom_mapping: dict[str, int] | None = None) -> int | None:
    """
    Nhận diện tên trường dữ liệu từ ảnh sang số cột chuẩn 1..28.
    Ưu tiên 100% theo bảng cấu hình ánh xạ do người dùng chỉ định.
    """
    if not field_name:
        return None
    raw_key = str(field_name).strip()
    norm = normalize_text_for_match(raw_key)

    # 1. Tra cứu trực tiếp trong cấu hình người dùng (nếu có)
    if custom_mapping:
        for k, col in custom_mapping.items():
            if str(k).strip().lower() == raw_key.lower():
                return int(col)
            if normalize_text_for_match(str(k)) == norm:
                return int(col)

    # 2. Tra cứu trong bảng chuẩn mặc định
    lower_k = raw_key.lower()
    if lower_k in FIELD_TO_COL_IDX:
        return FIELD_TO_COL_IDX[lower_k]

    if norm in FIELD_SYNONYMS:
        return FIELD_SYNONYMS[norm]

    for syn_text, c_idx in FIELD_SYNONYMS.items():
        if syn_text in norm or norm in syn_text:
            return c_idx

    return None


def match_student_row(
    candidate: dict[str, Any],
    sheet_rows: list[dict[str, Any]],
    selected_keys: list[str],
) -> tuple[int | None, list[int], str]:
    """
    Tìm dòng học sinh trong danh sách Excel khớp với thông tin trích xuất từ OCR.
    
    Hỗ trợ đối chiếu thông minh:
    - Nếu trường đang so sánh (ví dụ Họ tên) bị sửa đổi, cho phép khớp với cả giá trị cũ
      hoặc giá trị mới.
    - Trả về (matched_row_idx, all_candidate_indices, note)
    """
    if not selected_keys:
        return None, [], "Chưa chọn tiêu chí so khớp nào"

    # Trích xuất danh sách giá trị cũ nếu có trong cac_truong_sua
    old_vals: dict[str, str] = {}
    for mod in candidate.get("cac_truong_sua", []):
        t_key = mod.get("ten_truong", "")
        old_v = mod.get("gia_tri_cu")
        if t_key and old_v:
            old_vals[t_key] = str(old_v).strip()

    matches: list[int] = []

    for r_idx, row_data in enumerate(sheet_rows):
        is_all_match = True

        for k in selected_keys:
            cand_val = candidate.get(k)
            excel_val = row_data.get(k)

            norm_excel = normalize_match_val(excel_val, k)
            norm_cand_new = normalize_match_val(cand_val, k)
            norm_cand_old = normalize_match_val(old_vals.get(k), k) if k in old_vals else ""

            # Nếu phía OCR hoàn toàn không có dữ liệu cho tiêu chí này -> coi như không khớp
            if not norm_cand_new and not norm_cand_old:
                is_all_match = False
                break

            # Khớp nếu giá trị trong Excel trùng với giá trị mới HOẶC giá trị cũ của OCR
            matched_k = False
            if norm_cand_new and norm_excel == norm_cand_new:
                matched_k = True
            elif norm_cand_old and norm_excel == norm_cand_old:
                matched_k = True

            if not matched_k:
                is_all_match = False
                break

        if is_all_match:
            matches.append(r_idx)

    if len(matches) == 1:
        return matches[0], matches, "Khớp duy nhất"
    elif len(matches) > 1:
        return None, matches, f"Trùng {len(matches)} dòng (Cần kiểm tra lại)"

    # Cơ chế dự phòng thông minh: Nếu chọn cả STT + Họ tên mà không khớp cả hai
    # (do Họ tên bị sửa hoặc OCR đọc lệch dấu), ưu tiên định vị theo STT hoặc Họ tên
    if "stt" in selected_keys and len(selected_keys) > 1:
        stt_cand = normalize_match_val(candidate.get("stt"), "stt")
        if stt_cand:
            stt_matches = [
                i for i, r_d in enumerate(sheet_rows)
                if normalize_match_val(r_d.get("stt"), "stt") == stt_cand
            ]
            if len(stt_matches) == 1:
                return stt_matches[0], stt_matches, "Khớp theo STT"

    if "ho_ten" in selected_keys and len(selected_keys) > 1:
        name_new = normalize_match_val(candidate.get("ho_ten"), "ho_ten")
        name_old = normalize_match_val(old_vals.get("ho_ten"), "ho_ten") if "ho_ten" in old_vals else ""
        if name_new or name_old:
            name_matches = [
                i for i, r_d in enumerate(sheet_rows)
                if (name_new and normalize_match_val(r_d.get("ho_ten"), "ho_ten") == name_new)
                or (name_old and normalize_match_val(r_d.get("ho_ten"), "ho_ten") == name_old)
            ]
            if len(name_matches) == 1:
                return name_matches[0], name_matches, "Khớp theo Họ và tên"

    return None, [], "Không tìm thấy trong danh sách"


class OCRUpdaterProcessor:
    """Bộ xử lý cập nhật danh sách 28 cột từ ảnh chụp sửa đổi/bổ sung."""

    def __init__(self, config_mgr: OCRConfigManager | None = None):
        self.config_mgr = config_mgr or OCRConfigManager()

    def update_excel_from_images(
        self,
        excel_path: Path | str,
        image_paths: list[Path | str],
        selected_match_keys: list[str] | None = None,
        column_mapping: dict[str, int] | None = None,
        auto_append_new: bool = False,
        output_path: Path | str | None = None,
        api_key: str | None = None,
        base_url: str | None = None,
        model_name: str | None = None,
        prompt_text: str | None = None,
        progress_callback: Callable[[int, int, str], None] | None = None,
        file_status_callback: Callable[[Path, str], None] | None = None,
        log_callback: Callable[[str, str], None] | None = None,
        check_cancel: Callable[[], bool] | None = None,
        manual_review_callback: Callable[[dict[str, Any]], dict[str, Any]] | None = None,
        background_scan: bool = True,
        max_workers: int = 3,
    ) -> dict[str, Any]:
        """
        Thực thi OCR từng ảnh sửa đổi/bổ sung, định vị học sinh trong file Excel 28 cột
        và áp dụng các thay đổi kèm tô màu highlight và tạo sheet lịch sử.
        Hỗ trợ:
        - Cấu hình ánh xạ cột tùy biến 100% chuẩn xác.
        - Hộp thoại duyệt thủ công.
        - Cơ chế Producer-Consumer (`background_scan`): Luồng quét OCR vẫn tiếp tục chạy ngầm
          trên các ảnh tiếp theo trong khi người dùng đang xem và xác nhận hộp thoại chỉnh sửa.
        """
        def log(msg: str, tag: str = "info"):
            if log_callback:
                log_callback(msg, tag)

        def update_file_st(img_p: Path, status_str: str):
            if file_status_callback:
                try:
                    file_status_callback(img_p, status_str)
                except Exception:
                    pass

        excel_p = Path(excel_path)
        if not excel_p.is_file():
            raise FileNotFoundError(f"Không tìm thấy tệp Excel danh sách: {excel_p}")

        imgs = [Path(p) for p in image_paths if Path(p).is_file()]
        if not imgs:
            return {"status": "error", "message": "Không có tệp ảnh hợp lệ để xử lý."}

        keys_to_match = selected_match_keys or ["stt"]

        # Cấu hình Ánh xạ cột (100% chuẩn xác theo người dùng thiết lập)
        active_mapping = column_mapping
        if active_mapping is None:
            active_mapping = OCRColumnMappingManager().get_mapping()
        log(f"  🎯 Áp dụng bảng ánh xạ cột ({len(active_mapping)} quy tắc xác định 100%).", "info")

        # Cấu hình API OCR
        cfg = self.config_mgr.load_config()
        active_api_key = (api_key or cfg.get("api_key") or "").strip()
        active_base_url = (base_url or cfg.get("base_url") or DEFAULT_BASE_URL).strip()
        active_model_name = (model_name or cfg.get("model_name") or DEFAULT_MODEL_NAME).strip()
        active_prompt = (prompt_text or DEFAULT_PATCH_PROMPT).strip()

        if not active_api_key:
            return {
                "status": "error",
                "message": "Chưa có API Key. Vui lòng nhập API Key trên giao diện hoặc tệp cấu hình.",
            }

        log(f"📖 Đang nạp tệp Excel: {excel_p.name}...", "info")
        wb = openpyxl.load_workbook(excel_p, data_only=False)

        # Chọn Sheet dữ liệu chuẩn
        ws = None
        for s in wb.worksheets:
            if s.sheet_state == "visible" and s.title != "Lich_Su_Thay_Doi":
                ws = s
                break
        if ws is None:
            ws = wb.active

        header_info = detect_header_and_columns(ws)
        header_row_idx = header_info.get("header_row_idx", 1)
        raw_headers = header_info.get("raw_headers", [])
        log(f"  🔍 Dòng tiêu đề tại dòng {header_row_idx} với {len(raw_headers)} cột. Sheet: '{ws.title}'.", "info")

        # Nạp dữ liệu các dòng hiện có trong Excel vào bộ nhớ
        sheet_rows: list[dict[str, Any]] = []
        excel_row_indices: list[int] = []

        max_existing_stt = 0
        last_non_empty_row = header_row_idx

        for r_num in range(header_row_idx + 1, ws.max_row + 1):
            row_dict: dict[str, Any] = {}
            has_data = False

            # Đọc các trường nhận diện chính + 5 trường sửa đổi mục tiêu
            stt_val = ws.cell(row=r_num, column=1).value
            so_hieu_val = ws.cell(row=r_num, column=2).value
            so_vao_so_val = ws.cell(row=r_num, column=3).value
            ho_ten_val = ws.cell(row=r_num, column=5).value
            ngay_sinh_val = ws.cell(row=r_num, column=9).value
            gioi_tinh_val = ws.cell(row=r_num, column=10).value
            dan_toc_val = ws.cell(row=r_num, column=11).value
            noi_sinh_val = ws.cell(row=r_num, column=13).value
            ten_truong_val = ws.cell(row=r_num, column=14).value

            row_dict["excel_row"] = r_num
            row_dict["stt"] = stt_val
            row_dict["so_hieu_bang"] = so_hieu_val
            row_dict["so_vao_so"] = so_vao_so_val
            row_dict["ho_ten"] = ho_ten_val
            row_dict["ngay_sinh"] = ngay_sinh_val
            row_dict["gioi_tinh"] = gioi_tinh_val
            row_dict["dan_toc"] = dan_toc_val
            row_dict["noi_sinh"] = noi_sinh_val
            row_dict["ten_truong"] = ten_truong_val

            # Kiểm tra dòng có dữ liệu không
            if ho_ten_val or stt_val or so_hieu_val:
                has_data = True
                last_non_empty_row = r_num

            if stt_val is not None:
                try:
                    s_int = int(float(str(stt_val).strip()))
                    if s_int > max_existing_stt:
                        max_existing_stt = s_int
                except Exception:
                    pass

            sheet_rows.append(row_dict)
            excel_row_indices.append(r_num)

        log(f"  📊 Đã lập chỉ mục {len(sheet_rows):,} học sinh trong danh sách gốc (STT cao nhất: {max_existing_stt}).", "info")

        # Định dạng màu cho các ô thay đổi / bổ sung
        FILL_MODIFIED = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")  # Vàng nhạt
        FILL_ADDED = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")     # Xanh lá nhạt
        FONT_MODIFIED = Font(name="Times New Roman", size=11, bold=True, color="B25900")
        FONT_ADDED = Font(name="Times New Roman", size=11, bold=True, color="276A3C")
        BORDER_THIN = Border(
            left=Side(style="thin", color="D9D9D9"),
            right=Side(style="thin", color="D9D9D9"),
            top=Side(style="thin", color="D9D9D9"),
            bottom=Side(style="thin", color="D9D9D9"),
        )

        audit_entries: list[dict[str, Any]] = []
        total_images = len(imgs)
        success_images = 0
        modified_students_count = 0
        added_students_count = 0
        unmatched_count = 0

        # Lấy một dòng mẫu để kế thừa thông tin chung (Hội đồng thi, Ngày cấp, Tên trường...) khi bổ sung học sinh mới
        sample_context_row: dict[int, Any] = {}
        if last_non_empty_row > header_row_idx:
            for c in range(1, 29):
                sample_context_row[c] = ws.cell(row=last_non_empty_row, column=c).value

        log(f"🚀 BẮT ĐẦU CHẠY OCR & CẬP NHẬT TRÊN {total_images} ẢNH CHỤP...", "info")

        # Biến trạng thái duyệt thủ công
        review_active = bool(manual_review_callback)
        stop_process = False

        # Hàm quét 1 ảnh (có thể chạy trên luồng Producer hoặc chạy tuần tự)
        def _scan_single_image(idx: int, img_p: Path) -> tuple[int, Path, list[dict[str, Any]] | None, str | None]:
            prog_pct = int(((idx - 1) / total_images) * 100)
            if progress_callback:
                progress_callback(prog_pct, 100, f"Đang quét OCR [{idx}/{total_images}]: {img_p.name}")
            update_file_st(img_p, "⏳ Đang quét OCR...")
            log(f"📷 [{idx}/{total_images}] Đang đọc ảnh: {img_p.name}...", "info")

            try:
                raw_response = extract_table_from_image(
                    image_path=img_p,
                    prompt_text=active_prompt,
                    api_key=active_api_key,
                    base_url=active_base_url,
                    model_name=active_model_name,
                    log_callback=log_callback,
                )
            except Exception as e:
                err_s = f"Lỗi kết nối ({e})"
                update_file_st(img_p, f"❌ {err_s}")
                log(f"  ❌ Lỗi khi gửi ảnh {img_p.name}: {e}", "error")
                return idx, img_p, None, err_s

            records, err_msg = parse_ocr_json_response(raw_response)
            if records is None:
                update_file_st(img_p, f"⚠️ {err_msg}")
                log(f"  ⚠️ {img_p.name}: Không trích xuất được JSON ({err_msg}).", "warning")
                return idx, img_p, None, err_msg

            if not records:
                update_file_st(img_p, "ℹ️ Không có sửa/bổ sung")
                log(f"  ℹ️ {img_p.name}: Không phát hiện dòng nào có sửa đổi hoặc bổ sung.", "info")
                return idx, img_p, [], None

            update_file_st(img_p, f"✨ Phát hiện {len(records)} HS (Chờ duyệt)")
            log(f"  ✨ {img_p.name}: Phát hiện {len(records)} học sinh có đánh dấu sửa đổi/bổ sung.", "success")
            return idx, img_p, records, None

        # Khởi tạo hàng đợi kết quả OCR
        scan_queue: queue.Queue[tuple[int, Path, list[dict[str, Any]] | None, str | None] | None] = queue.Queue()

        def _producer_worker():
            workers = max(1, min(int(max_workers or 3), total_images))
            if workers <= 1:
                for idx, img_p in enumerate(imgs, start=1):
                    if stop_process or (check_cancel and check_cancel()):
                        log(f"🛑 Đã dừng luồng quét OCR tại ảnh {img_p.name}.", "warning")
                        break
                    res_item = _scan_single_image(idx, img_p)
                    scan_queue.put(res_item)
                scan_queue.put(None)
                return

            # Quét song song nhiều ảnh cùng lúc để tăng tốc gấp 2-3 lần nhưng vẫn đẩy vào queue theo đúng thứ tự trang
            with ThreadPoolExecutor(max_workers=workers) as executor:
                futures_map = {}
                next_submit_idx = 0
                next_yield_idx = 0

                # Nộp trước tối đa `workers * 2` ảnh vào pool
                while next_submit_idx < total_images and len(futures_map) < workers * 2:
                    if stop_process or (check_cancel and check_cancel()):
                        break
                    i_one = next_submit_idx + 1
                    p_img = imgs[next_submit_idx]
                    futures_map[next_submit_idx] = executor.submit(_scan_single_image, i_one, p_img)
                    next_submit_idx += 1

                while next_yield_idx in futures_map:
                    if stop_process or (check_cancel and check_cancel()):
                        log("🛑 Đã dừng luồng quét OCR song song.", "warning")
                        break
                    fut = futures_map.pop(next_yield_idx)
                    try:
                        res_item = fut.result()
                        scan_queue.put(res_item)
                    except Exception as e:
                        p_err = imgs[next_yield_idx]
                        scan_queue.put((next_yield_idx + 1, p_err, None, str(e)))
                    next_yield_idx += 1

                    # Nộp tiếp ảnh mới vào pool
                    if next_submit_idx < total_images and not (stop_process or (check_cancel and check_cancel())):
                        i_one = next_submit_idx + 1
                        p_img = imgs[next_submit_idx]
                        futures_map[next_submit_idx] = executor.submit(_scan_single_image, i_one, p_img)
                        next_submit_idx += 1

            scan_queue.put(None)  # Sentinel kết thúc quét

        if background_scan and total_images > 1:
            # Chạy quét OCR ở luồng riêng để không bị chặn khi người dùng mở hộp thoại duyệt
            producer_thread = threading.Thread(target=_producer_worker, daemon=True)
            producer_thread.start()
        else:
            # Nếu chỉ có 1 ảnh hoặc tắt background_scan, đẩy tuần tự hoặc chạy trực tiếp
            producer_thread = None

        img_cursor = 0
        while True:
            if stop_process or (check_cancel and check_cancel()):
                break

            if producer_thread is not None:
                q_item = scan_queue.get()
                if q_item is None:
                    break
                idx, img_p, records, err_msg = q_item
            else:
                if img_cursor >= total_images:
                    break
                img_p = imgs[img_cursor]
                img_cursor += 1
                idx = img_cursor
                if check_cancel and check_cancel():
                    log(f"🛑 Đã dừng theo yêu cầu tại ảnh {img_p.name}.", "warning")
                    break
                idx, img_p, records, err_msg = _scan_single_image(idx, img_p)

            if records is None:
                continue

            success_images += 1
            if not records:
                continue

            img_mod_cnt = 0
            img_add_cnt = 0

            for rec in records:
                if stop_process or (check_cancel and check_cancel()):
                    break

                rec_type = str(rec.get("loai") or "").strip().upper()
                ho_ten = str(rec.get("ho_ten") or "").strip()
                stt_val = rec.get("stt")
                ngay_sinh_ocr = str(rec.get("ngay_sinh") or "").strip()
                if not ngay_sinh_ocr:
                    for m_it in (rec.get("cac_truong_sua") or []):
                        if resolve_field_to_col_idx(str(m_it.get("ten_truong", "")), custom_mapping=active_mapping) == 9:
                            ngay_sinh_ocr = str(m_it.get("gia_tri_moi") or m_it.get("gia_tri_cu") or "").strip()
                            break

                # Nếu không bật auto_append_new, bỏ qua ngay các bản ghi loại BO_SUNG
                if rec_type == "BO_SUNG" and not auto_append_new:
                    log(f"    ⏭️ Bỏ qua dòng viết thêm (bổ sung): {ho_ten}.", "info")
                    continue

                # 1. Thử so khớp học sinh với các dòng hiện có
                matched_idx, all_cand, match_note = match_student_row(rec, sheet_rows, keys_to_match)

                # Nếu là bản ghi SUA_DOI, cho phép cập nhật khi đã khớp tự động HOẶC khi bật duyệt thủ công (để người dùng chọn lại STT)
                if matched_idx is not None or (rec_type != "BO_SUNG" and review_active and manual_review_callback):
                    excel_r = excel_row_indices[matched_idx] if matched_idx is not None else None
                    mod_fields = list(rec.get("cac_truong_sua") or [])

                    # Nếu AI không liệt kê cac_truong_sua cụ thể, so sánh trực tiếp 5 trường sửa đổi mục tiêu
                    if not mod_fields and excel_r is not None:
                        for field_k in ("ho_ten", "ngay_sinh", "gioi_tinh", "dan_toc", "noi_sinh"):
                            new_v = rec.get(field_k)
                            if new_v is not None and str(new_v).strip():
                                c_idx = resolve_field_to_col_idx(field_k, custom_mapping=active_mapping)
                                if c_idx:
                                    old_cell_v = ws.cell(row=excel_r, column=c_idx).value
                                    if normalize_match_val(old_cell_v, field_k) != normalize_match_val(new_v, field_k):
                                        mod_fields.append({
                                            "ten_truong": field_k,
                                            "gia_tri_cu": str(old_cell_v or ""),
                                            "gia_tri_moi": str(new_v).strip()
                                        })

                    # Chuẩn bị danh sách trường thay đổi
                    candidate_changes: list[dict[str, Any]] = []
                    for m_item in mod_fields:
                        f_name = m_item.get("ten_truong", "")
                        col_idx = resolve_field_to_col_idx(f_name, custom_mapping=active_mapping)
                        if not col_idx:
                            continue

                        val_new = m_item.get("gia_tri_moi", "")
                        val_old = m_item.get("gia_tri_cu", "")
                        curr_excel_val = ws.cell(row=excel_r, column=col_idx).value if excel_r is not None else None

                        col_label = TARGET_COLUMNS[col_idx - 1]["name"] if (1 <= col_idx <= len(TARGET_COLUMNS)) else f_name
                        candidate_changes.append({
                            "field_name": f_name,
                            "col_idx": col_idx,
                            "col_name": col_label,
                            "old_val": str(curr_excel_val if curr_excel_val is not None else (val_old or "")),
                            "new_val": str(val_new or ""),
                        })

                    if not candidate_changes:
                        continue

                    # NẾU BẬT DUYỆT THỦ CÔNG: GỌI CALLBACK HIỆN HỘP THOẠI KÈM STT & HỌ TÊN NHẬN DIỆN TỪ ẢNH
                    if review_active and manual_review_callback:
                        excel_stt_val = str(sheet_rows[matched_idx].get("stt", "") or "") if matched_idx is not None else str(stt_val or "")
                        excel_stu_name = str(sheet_rows[matched_idx].get("ho_ten", "") or "") if matched_idx is not None else ""
                        review_req = {
                            "type": "SUA_DOI",
                            "ocr_stt": stt_val if stt_val is not None else "",
                            "ocr_student_name": ho_ten,
                            "ocr_ngay_sinh": ngay_sinh_ocr,
                            "matched_idx": matched_idx,
                            "excel_stt": excel_stt_val,
                            "student_name": excel_stu_name or ho_ten,
                            "row_num": excel_r,
                            "match_note": match_note,
                            "image_name": img_p.name,
                            "reason": rec.get("ly_do_nghi_ngo"),
                            "changes": candidate_changes,
                            "sheet_rows": sheet_rows,
                        }
                        user_decision = manual_review_callback(review_req)
                        act = user_decision.get("action", "apply")
                        if act == "skip":
                            log(f"    ⏭️ Người dùng bỏ qua sửa học sinh '{ho_ten}' (STT ảnh: {stt_val}).", "warning")
                            audit_entries.append({
                                "row_num": excel_r or "---",
                                "type": "Bỏ qua",
                                "student_name": ho_ten,
                                "field_name": "---",
                                "old_val": "---",
                                "new_val": "---",
                                "image_name": img_p.name,
                                "status": "Người dùng bỏ qua",
                                "note": "Bỏ qua khi duyệt thủ công"
                            })
                            continue
                        elif act == "apply_all":
                            review_active = False
                            log("    ⚡ Chế độ: Tự động áp dụng tất cả các học sinh còn lại.", "info")
                        elif act == "stop":
                            log("🛑 Người dùng yêu cầu dừng tiến trình duyệt.", "warning")
                            stop_process = True
                            break

                        # Kiểm tra xem người dùng có chọn/khớp lại STT thực trong file Excel không
                        override_idx = user_decision.get("matched_idx", matched_idx)
                        if override_idx is not None and 0 <= override_idx < len(excel_row_indices):
                            matched_idx = override_idx
                            excel_r = excel_row_indices[matched_idx]

                        # Nhận các giá trị người dùng đã sửa đổi thủ công
                        candidate_changes = user_decision.get("changes", candidate_changes)

                    # Nếu sau khi duyệt vẫn không xác định được dòng Excel để sửa
                    if matched_idx is None or excel_r is None:
                        unmatched_count += 1
                        audit_entries.append({
                            "row_num": "---",
                            "type": "Không khớp",
                            "student_name": ho_ten,
                            "field_name": "---",
                            "old_val": "---",
                            "new_val": "---",
                            "image_name": img_p.name,
                            "status": "Không tìm thấy",
                            "note": match_note
                        })
                        log(f"    ⚠️ Không tìm thấy vị trí của học sinh '{ho_ten}' (STT ảnh: {stt_val} - {match_note}).", "warning")
                        continue

                    modified_students_count += 1
                    img_mod_cnt += 1
                    updated_cols_log = []
                    excel_current_name = str(sheet_rows[matched_idx].get("ho_ten") or ho_ten)
                    for c_info in candidate_changes:
                        col_idx = c_info["col_idx"]
                        f_name = c_info["field_name"]
                        val_new = c_info["new_val"]

                        # Chuẩn hóa giá trị theo nghiệp vụ cột
                        if col_idx == 5:
                            val_new = vietnamese_title_case(str(val_new))
                        elif col_idx == 9:
                            val_new = format_date_to_text(val_new)
                        elif col_idx in (2, 3):
                            val_new = re.sub(r"\s+", "", str(val_new))

                        cell = ws.cell(row=excel_r, column=col_idx)
                        curr_excel_val = cell.value
                        cell.value = val_new
                        cell.fill = FILL_MODIFIED
                        cell.font = FONT_MODIFIED

                        # Cập nhật lại chỉ mục trong bộ nhớ
                        for k, v_c in FIELD_TO_COL_IDX.items():
                            if v_c == col_idx:
                                sheet_rows[matched_idx][k] = val_new

                        updated_cols_log.append(f_name)
                        audit_entries.append({
                            "row_num": excel_r,
                            "type": "Sửa đổi",
                            "student_name": str(sheet_rows[matched_idx].get("ho_ten") or excel_current_name),
                            "field_name": c_info.get("col_name", f_name),
                            "old_val": str(curr_excel_val if curr_excel_val is not None else c_info.get("old_val", "")),
                            "new_val": val_new,
                            "image_name": img_p.name,
                            "status": "Thành công",
                            "note": rec.get("ly_do_nghi_ngo") or f"Đã cập nhật (STT: {sheet_rows[matched_idx].get('stt')})"
                        })

                    log(
                        f"    ✏️ STT {sheet_rows[matched_idx].get('stt')} (Dòng {excel_r} - {sheet_rows[matched_idx].get('ho_ten')}): "
                        f"Đã sửa {len(candidate_changes)} ô ({', '.join(updated_cols_log)}).",
                        "info"
                    )

                elif auto_append_new:
                    # BỔ SUNG HỌC SINH MỚI VÀO CUỐI DANH SÁCH
                    max_existing_stt += 1
                    new_r = ws.max_row + 1

                    # Tạo dữ liệu 28 cột cho học sinh mới
                    new_row_vals: list[Any] = [None] * 28

                    # 1: STT
                    new_row_vals[0] = max_existing_stt
                    # 2: Số hiệu bằng
                    new_row_vals[1] = re.sub(r"\s+", "", str(rec.get("so_hieu_bang") or ""))
                    # 3: Số vào sổ
                    new_row_vals[2] = re.sub(r"\s+", "", str(rec.get("so_vao_so") or ""))
                    # 4: Tên văn bằng
                    new_row_vals[3] = sample_context_row.get(4) or "Bằng tốt nghiệp trung học cơ sở"
                    # 5: Họ và tên
                    new_row_vals[4] = vietnamese_title_case(str(rec.get("ho_ten") or ""))
                    # 6: Mã người học
                    new_row_vals[5] = ""
                    # 7: CCCD
                    new_row_vals[6] = ""
                    # 8: Hộ chiếu
                    new_row_vals[7] = ""
                    # 9: Ngày sinh
                    new_row_vals[8] = format_date_to_text(rec.get("ngay_sinh") or "")
                    # 10: Giới tính
                    new_row_vals[9] = clean_whitespace(str(rec.get("gioi_tinh") or ""))
                    # 11: Dân tộc
                    new_row_vals[10] = clean_whitespace(str(rec.get("dan_toc") or sample_context_row.get(11) or "Kinh"))
                    # 12: Quốc tịch
                    new_row_vals[11] = sample_context_row.get(12) or "Việt Nam"
                    # 13: Nơi sinh
                    new_row_vals[12] = clean_whitespace(str(rec.get("noi_sinh") or ""))
                    # 14: Tên trường
                    new_row_vals[13] = clean_whitespace(str(rec.get("ten_truong") or sample_context_row.get(14) or ""))
                    # 15: Niên khóa
                    new_row_vals[14] = sample_context_row.get(15) or ""
                    # 16: Điểm/Xếp loại
                    new_row_vals[15] = clean_whitespace(str(rec.get("xep_loai") or ""))
                    # 17: Hội đồng thi
                    new_row_vals[16] = sample_context_row.get(17) or ""
                    # 18: Địa danh
                    new_row_vals[17] = sample_context_row.get(18) or ""
                    # 19: Ngày cấp bằng
                    new_row_vals[18] = sample_context_row.get(19) or ""
                    # 20: Tên cơ quan cấp bằng
                    new_row_vals[19] = sample_context_row.get(20) or ""
                    # 21: Chức danh người ký
                    new_row_vals[20] = sample_context_row.get(21) or ""
                    # 22: Họ tên người ký
                    new_row_vals[21] = sample_context_row.get(22) or ""
                    # 23: Tình trạng văn bằng
                    new_row_vals[22] = sample_context_row.get(23) or ""
                    # 24: Ghi chú
                    new_row_vals[23] = clean_whitespace(str(rec.get("ghi_chu") or "Bổ sung từ ảnh sửa đổi"))
                    # 25: Trạng thái số hóa
                    new_row_vals[24] = "Đã bổ sung từ ảnh sửa đổi"
                    # 26: Lý do lỗi
                    new_row_vals[25] = ""
                    # 27: Mã đơn vị
                    new_row_vals[26] = sample_context_row.get(27) or ""
                    # 28: Tên đơn vị
                    new_row_vals[27] = sample_context_row.get(28) or ""

                    # Áp dụng các trường tùy biến khác trên ảnh theo bảng ánh xạ người dùng
                    for k_rec, v_rec in rec.items():
                        if k_rec in ("loai", "stt", "can_kiem_tra", "ly_do_nghi_ngo", "cac_truong_sua"):
                            continue
                        if v_rec is None or not str(v_rec).strip():
                            continue
                        c_i = resolve_field_to_col_idx(k_rec, custom_mapping=active_mapping)
                        if c_i and 1 <= c_i <= 28:
                            val_str = str(v_rec).strip()
                            if c_i == 5:
                                val_str = vietnamese_title_case(val_str)
                            elif c_i == 9:
                                val_str = format_date_to_text(val_str)
                            elif c_i in (2, 3):
                                val_str = re.sub(r"\s+", "", val_str)
                            new_row_vals[c_i - 1] = val_str

                    # NẾU BẬT DUYỆT THỦ CÔNG KHI THÊM MỚI
                    if review_active and manual_review_callback:
                        new_student_fields = [
                            {"field_name": "ho_ten", "col_idx": 5, "col_name": "Họ và tên", "val": str(new_row_vals[4] or "")},
                            {"field_name": "ngay_sinh", "col_idx": 9, "col_name": "Ngày, tháng, năm sinh", "val": str(new_row_vals[8] or "")},
                            {"field_name": "gioi_tinh", "col_idx": 10, "col_name": "Giới tính", "val": str(new_row_vals[9] or "")},
                            {"field_name": "noi_sinh", "col_idx": 13, "col_name": "Nơi sinh", "val": str(new_row_vals[12] or "")},
                            {"field_name": "so_hieu_bang", "col_idx": 2, "col_name": "Số hiệu bằng", "val": str(new_row_vals[1] or "")},
                            {"field_name": "so_vao_so", "col_idx": 3, "col_name": "Số vào sổ", "val": str(new_row_vals[2] or "")},
                            {"field_name": "xep_loai", "col_idx": 16, "col_name": "Điểm / Xếp loại", "val": str(new_row_vals[15] or "")},
                            {"field_name": "ghi_chu", "col_idx": 24, "col_name": "Ghi chú", "val": str(new_row_vals[23] or "")},
                        ]
                        review_req = {
                            "type": "BO_SUNG",
                            "student_name": ho_ten,
                            "row_num": new_r,
                            "image_name": img_p.name,
                            "reason": rec.get("ly_do_nghi_ngo"),
                            "fields": new_student_fields,
                        }
                        user_decision = manual_review_callback(review_req)
                        act = user_decision.get("action", "apply")
                        if act == "skip":
                            log(f"    ⏭️ Người dùng bỏ qua thêm học sinh mới '{ho_ten}'.", "warning")
                            max_existing_stt -= 1
                            continue
                        elif act == "apply_all":
                            review_active = False
                            log("    ⚡ Chế độ: Tự động áp dụng tất cả các học sinh còn lại.", "info")
                        elif act == "stop":
                            log("🛑 Người dùng yêu cầu dừng tiến trình duyệt.", "warning")
                            max_existing_stt -= 1
                            stop_process = True
                            break

                        for ef in user_decision.get("fields", new_student_fields):
                            c_i = ef.get("col_idx")
                            v_ed = ef.get("val", "")
                            if c_i and 1 <= c_i <= 28:
                                new_row_vals[c_i - 1] = v_ed

                    added_students_count += 1
                    img_add_cnt += 1

                    # Ghi vào sheet
                    for c_i, v_val in enumerate(new_row_vals, start=1):
                        c_cell = ws.cell(row=new_r, column=c_i, value=v_val)
                        c_cell.fill = FILL_ADDED
                        c_cell.font = FONT_ADDED
                        c_cell.border = BORDER_THIN

                    # Cập nhật chỉ mục bộ nhớ
                    sheet_rows.append({
                        "stt": max_existing_stt,
                        "so_hieu_bang": new_row_vals[1],
                        "so_vao_so": new_row_vals[2],
                        "ho_ten": new_row_vals[4],
                        "ngay_sinh": new_row_vals[8],
                        "ten_truong": new_row_vals[13],
                    })
                    excel_row_indices.append(new_r)

                    audit_entries.append({
                        "row_num": new_r,
                        "type": "Bổ sung mới",
                        "student_name": ho_ten,
                        "field_name": "Toàn bộ dòng",
                        "old_val": "---",
                        "new_val": f"STT {max_existing_stt} - {ho_ten}",
                        "image_name": img_p.name,
                        "status": "Đã thêm mới",
                        "note": rec.get("ly_do_nghi_ngo") or "Học sinh mới ghi thêm trên sổ"
                    })
                    log(f"    ➕ Thêm mới dòng {new_r} (STT {max_existing_stt}): {ho_ten}.", "success")

                else:
                    # KHÔNG TÌM THẤY VÀ KHÔNG BẬT TỰ ĐỘNG THÊM MỚI
                    unmatched_count += 1
                    audit_entries.append({
                        "row_num": "---",
                        "type": "Không khớp",
                        "student_name": ho_ten,
                        "field_name": "---",
                        "old_val": "---",
                        "new_val": "---",
                        "image_name": img_p.name,
                        "status": "Không tìm thấy",
                        "note": match_note
                    })
                    log(f"    ⚠️ Không tìm thấy vị trí của học sinh '{ho_ten}' ({match_note}).", "warning")

            update_file_st(img_p, f"✅ Đã xử lý (Sửa {img_mod_cnt}, Thêm {img_add_cnt})")
            if stop_process:
                break

        # 4. TẠO HOẶC CẬP NHẬT SHEET 'Lich_Su_Thay_Doi'
        if "Lich_Su_Thay_Doi" in wb.sheetnames:
            del wb["Lich_Su_Thay_Doi"]

        ws_audit = wb.create_sheet(title="Lich_Su_Thay_Doi")
        audit_headers = [
            ("stt", "STT", 8),
            ("row_num", "Dòng Excel", 12),
            ("type", "Loại thao tác", 16),
            ("student_name", "Họ và tên", 26),
            ("field_name", "Trường thay đổi", 18),
            ("old_val", "Giá trị cũ trong file", 24),
            ("new_val", "Giá trị mới (từ ảnh)", 26),
            ("image_name", "Tệp ảnh nguồn", 22),
            ("status", "Trạng thái", 18),
            ("note", "Ghi chú đối soát", 32),
        ]

        # Tiêu đề bảng lịch sử
        HDR_FILL = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
        HDR_FONT = Font(name="Times New Roman", size=11, bold=True, color="FFFFFF")

        for col_i, (h_key, h_label, h_w) in enumerate(audit_headers, start=1):
            c = ws_audit.cell(row=1, column=col_i, value=h_label)
            c.fill = HDR_FILL
            c.font = HDR_FONT
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            c.border = BORDER_THIN
            col_letter = get_column_letter(col_i)
            ws_audit.column_dimensions[col_letter].width = h_w

        # Ghi các dòng lịch sử
        for a_idx, entry in enumerate(audit_entries, start=1):
            r_a = a_idx + 1
            ws_audit.cell(row=r_a, column=1, value=a_idx).alignment = Alignment(horizontal="center")
            ws_audit.cell(row=r_a, column=2, value=entry["row_num"]).alignment = Alignment(horizontal="center")
            ws_audit.cell(row=r_a, column=3, value=entry["type"]).alignment = Alignment(horizontal="center")
            ws_audit.cell(row=r_a, column=4, value=entry["student_name"])
            ws_audit.cell(row=r_a, column=5, value=entry["field_name"])
            ws_audit.cell(row=r_a, column=6, value=entry["old_val"])
            ws_audit.cell(row=r_a, column=7, value=entry["new_val"])
            ws_audit.cell(row=r_a, column=8, value=entry["image_name"])
            ws_audit.cell(row=r_a, column=9, value=entry["status"]).alignment = Alignment(horizontal="center")
            ws_audit.cell(row=r_a, column=10, value=entry["note"])

            for c_i in range(1, 11):
                cell_item = ws_audit.cell(row=r_a, column=c_i)
                cell_item.border = BORDER_THIN
                cell_item.font = Font(name="Times New Roman", size=11)
                if entry["type"] == "Sửa đổi":
                    cell_item.fill = PatternFill(start_color="FFF9E6", end_color="FFF9E6", fill_type="solid")
                elif entry["type"] == "Bổ sung mới":
                    cell_item.fill = PatternFill(start_color="F2F9F2", end_color="F2F9F2", fill_type="solid")
                elif entry["type"] == "Không khớp":
                    cell_item.fill = PatternFill(start_color="FDE8E8", end_color="FDE8E8", fill_type="solid")

        # Lưu tệp kết quả
        if output_path:
            final_out_path = Path(output_path)
        else:
            final_out_path = excel_p.parent / f"{excel_p.stem}_DaCapNhat_{time.strftime('%Y%m%d_%H%M%S')}.xlsx"

        final_out_path.parent.mkdir(parents=True, exist_ok=True)
        wb.save(final_out_path)

        if progress_callback:
            progress_callback(100, 100, "Hoàn tất cập nhật.")

        log(
            f"🎉 HOÀN TẤT CẬP NHẬT: Đã sửa {modified_students_count} học sinh, thêm mới {added_students_count} học sinh "
            f"từ {success_images}/{total_images} ảnh. ➔ Đã lưu: {final_out_path.name}",
            "success"
        )

        return {
            "status": "success",
            "total_images": total_images,
            "success_images": success_images,
            "modified_students": modified_students_count,
            "added_students": added_students_count,
            "unmatched_records": unmatched_count,
            "total_changes": len(audit_entries),
            "output_file": final_out_path,
        }
