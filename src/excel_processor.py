from __future__ import annotations

import os
import sys
import io
import re
import copy
import unicodedata
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from src.lookup_manager import LookupManager, get_data_dir, find_data_file
from src.normalizer import (
    clean_whitespace,
    is_cell_empty,
    vietnamese_title_case,
    normalize_birthplace,
    normalize_ethnic,
    normalize_nationality,
    format_date_to_text,
    format_grad_date_to_text,
    check_dob_validity,
    parse_nien_khoa,
    infer_cohort_from_birth_year,
    find_mode_birth_year,
    fuzzy_match_school,
    auto_fill_graduation_dates_by_school_and_cohort,
    determine_status_and_reason
)
from src.exporter import DataExporter

SUMMARY_KEYWORDS = [
    "danh sách này có", "danh sách có", "danh sách gồm",
    "tổng số", "tổng cộng", "số học sinh",
    "công nhận tốt nghiệp", "người lập", "người ký", "hiệu trưởng", "thủ trưởng"
]

def is_valid_stt(val):
    """Kiểm tra ô có phải là Số thứ tự hợp lệ của học sinh (1, 2, 3...)."""
    if val is None or isinstance(val, bool):
        return False
    if isinstance(val, (int, float)):
        import math
        if math.isnan(val):
            return False
        return val > 0 and float(val).is_integer()
    s = str(val).strip()
    s_clean = re.sub(r'\.0+$', '', s)
    return s_clean.isdigit() and int(s_clean) > 0

def is_summary_or_footer_row(row_cells):
    """Kiểm tra xem dòng có phải là dòng tổng kết, chân trang ký tên hay không."""
    for cell in row_cells[:15]:
        val_raw = cell.value if hasattr(cell, 'value') else cell
        val = str(val_raw or '').strip().lower()
        if any(kw in val for kw in SUMMARY_KEYWORDS):
            return True
    return False

def get_cached_row_fonts(ws, reference_row_idx, max_col=28):
    """Trích xuất và lưu sẵn định dạng Font từ dòng mẫu."""
    cached = []
    for col in range(1, max_col + 1):
        cell = ws.cell(row=reference_row_idx, column=col)
        if cell.has_style and cell.font:
            cached.append(Font(name=cell.font.name, size=cell.font.size))
        else:
            cached.append(None)
    return cached

def apply_cached_fonts(ws, target_row_idx, cached_fonts):
    """Gán font từ bộ nhớ cache vào dòng mục tiêu."""
    for col_idx, font in enumerate(cached_fonts, start=1):
        if font:
            ws.cell(row=target_row_idx, column=col_idx).font = font

def get_file_format_info(ws):
    """
    Nhận diện định dạng sheet:
    - Quét tối đa 30 dòng đầu tìm dòng tiêu đề chứa 'STT' hoặc 'Họ và tên'.
    - Phân loại: 'AE' (31 cột -> loại bỏ Q, R, S) hoặc 'AB' (28 cột chuẩn).
    - Trả về: (format_type, headers, header_row_idx)
    """
    for r_idx, row in enumerate(ws.iter_rows(min_row=1, max_row=30, values_only=True), start=1):
        if not row:
            continue
        row_str = [clean_whitespace(str(cell)) if cell is not None else "" for cell in row]
        row_lower = [s.lower() for s in row_str]

        has_stt = any(s in ["stt", "số tt", "số thứ tự", "stt.", "tt", "tt."] for s in row_lower[:5])
        has_name = any(("họ" in s and "tên" in s) for s in row_lower)

        if has_stt or has_name:
            headers = row_str
            col_idx = -1
            for idx, h in enumerate(row_lower):
                if "tên đơn vị" in h or h == "đơn vị":
                    col_idx = idx
                    break

            if col_idx >= 29:
                return "AE", headers, r_idx
            elif 25 <= col_idx <= 28:
                return "AB", headers, r_idx
            elif col_idx == -1 and has_name:
                non_empty_cols = sum(1 for c in headers if c)
                if non_empty_cols >= 30:
                    return "AE", headers, r_idx
                elif non_empty_cols >= 15:
                    return "AB", headers, r_idx
                return "INVALID", headers, r_idx
            else:
                return "INVALID", headers, r_idx

    return "INVALID", [], None

def extract_cohort_from_text(text: str) -> str:
    """Trích xuất chuỗi niên khóa dạng 2018-2022, 2018_2022, 18-22, 2020 từ chuỗi text (file name / sheet name)."""
    if not text:
        return ""
    # Mẫu 4 số - 4 số: 2018-2022, 2018_2022, 2018 - 2022
    m = re.search(r'\b(19\d{2}|20\d{2})\s*[-_–]\s*(19\d{2}|20\d{2})\b', text)
    if m:
        return f"{m.group(1)}-{m.group(2)}"
    # Mẫu 2 số - 2 số: 18-22
    m2 = re.search(r'\b(\d{2})\s*[-_–]\s*(\d{2})\b', text)
    if m2:
        y1 = int(m2.group(1))
        y2 = int(m2.group(2))
        full_y1 = (1900 + y1) if y1 >= 80 else (2000 + y1)
        full_y2 = (1900 + y2) if y2 >= 80 else (2000 + y2)
        if full_y2 > full_y1 and (full_y2 - full_y1) in (3, 4, 5):
            return f"{full_y1}-{full_y2}"
    # Mẫu 1 năm duy nhất: vd 2022
    m3 = re.search(r'\b(19\d{2}|20\d{2})\b', text)
    if m3:
        y_val = int(m3.group(1))
        return f"{y_val - 4}-{y_val}"
    return ""

def match_school_from_text(text: str, allowed_schools: list[str]) -> str:
    """Tìm trường khớp trong chuỗi văn bản (tên file, tên thư mục, tên sheet). Hỗ trợ chuẩn xác tiếng Việt có dấu, ký tự & và Unicode NFD/NFC."""
    if not text or not allowed_schools:
        return ""
    text_nfc = unicodedata.normalize('NFC', str(text))
    text_clean = re.sub(r'[_\-\.]', ' ', text_nfc).lower()
    for sch in allowed_schools:
        sch_nfc = unicodedata.normalize('NFC', str(sch)).lower()
        if sch_nfc in text_clean:
            return sch
    # Fuzzy match nếu gần khớp
    close = fuzzy_match_school(text_clean, allowed_schools, n=1, cutoff=0.75)
    if close:
        return close[0]
    return ""


class ExcelProcessor:
    """Bộ xử lý đọc, chuẩn hóa và xuất file Excel theo đúng quy chuẩn 28 cột của Bộ GD&ĐT."""
    def __init__(self, lookup_mgr: LookupManager, template_path: Path | str | None = None):
        self.lookup_mgr = lookup_mgr
        if template_path is None:
            self.template_path = find_data_file("Mau_DS_Cap_Bang_THCS.xlsx")
        else:
            self.template_path = Path(template_path)

        if not self.template_path.exists():
            raise FileNotFoundError(f"Không tìm thấy file mẫu template: {self.template_path}")

        # Nạp template bytes và font cache
        with open(self.template_path, "rb") as f:
            self.template_bytes = f.read()

        wb_tpl = openpyxl.load_workbook(io.BytesIO(self.template_bytes))
        ws_tpl = wb_tpl.active
        self.cached_fonts = get_cached_row_fonts(ws_tpl, ws_tpl.max_row, max_col=28)
        wb_tpl.close()

        self.yellow_fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
        self.exporter = DataExporter(self.template_bytes, self.cached_fonts, self.yellow_fill)

    def extract_students_from_file(
        self,
        file_path: Path,
        allowed_schools: list[str],
        log_callback=None,
        check_cancel=None
    ) -> list[dict]:
        """
        Quét và trích xuất dữ liệu thô của học sinh từ 1 tệp Excel:
        - Tự động nhận diện định dạng 28 hoặc 31 cột (AE/AB).
        - Trích xuất thông tin ngữ cảnh (thư mục, tên file, tên sheet).
        - Đóng file ngay sau khi đọc để giải phóng RAM tối đa.
        """
        def log(msg, level="info"):
            if log_callback:
                log_callback(msg, level)
            else:
                print(msg)

        try:
            wb_in = openpyxl.load_workbook(file_path, data_only=True)
        except Exception as e:
            log(f"❌ Không thể mở file {file_path.name}: {e}", "error")
            return []

        file_students = []

        # Heuristic 1 & 2: Nhận diện trường và niên khóa từ đường dẫn và tên tệp
        fallback_sch_parent = match_school_from_text(file_path.parent.name, allowed_schools)
        fallback_sch_file = match_school_from_text(file_path.stem, allowed_schools)
        file_level_school = fallback_sch_parent or fallback_sch_file
        file_level_cohort = extract_cohort_from_text(file_path.stem)

        try:
            for ws_in in wb_in.worksheets:
                if check_cancel and check_cancel():
                    log(f"  ⏹ Dừng trích xuất các sheet còn lại trong {file_path.name} theo yêu cầu.", "warning")
                    break

                if ws_in.sheet_state != 'visible':
                    continue

                fmt_type, headers, header_r_idx = get_file_format_info(ws_in)
                if fmt_type == "INVALID":
                    log(f"  -> Bỏ qua Sheet '{ws_in.title}' (Cấu trúc không khớp 28 hoặc 31 cột)", "warning")
                    continue

                remove_qrs = (fmt_type == "AE")
                log(f"  -> Quét Sheet '{ws_in.title}': {'Dạng AE 31 cột (cắt Q, R, S)' if remove_qrs else 'Dạng AB 28 cột chuẩn'}", "info")

                # Heuristic 3: Nhận diện trường hoặc niên khóa từ tên Sheet
                sheet_level_school = match_school_from_text(ws_in.title, allowed_schools)
                sheet_level_cohort = extract_cohort_from_text(ws_in.title)

                effective_context_school = sheet_level_school or file_level_school
                effective_context_cohort = sheet_level_cohort or file_level_cohort

                start_extracting = False
                empty_consecutive = 0
                sheet_students = []

                for row in ws_in.iter_rows():
                    r_idx = row[0].row if hasattr(row[0], 'row') else 0
                    if header_r_idx is not None and r_idx <= header_r_idx:
                        continue

                    cell_a_val = row[0].value
                    cell_a_str = str(cell_a_val).strip() if cell_a_val is not None else ""

                    name_cell_val = row[4].value if len(row) > 4 else None
                    name_cell_str = clean_whitespace(str(name_cell_val)) if name_cell_val is not None else ""
                    has_letters = bool(re.search(r'[a-zA-ZÀ-ỹ]', name_cell_str))
                    is_row_empty = not any(c.value is not None and str(c.value).strip() != "" for c in row[:15])

                    if not start_extracting:
                        is_stt_one = bool(re.match(r'^0*1(\.0*)?$', cell_a_str))
                        if (is_stt_one and has_letters) or (has_letters and name_cell_str.lower() not in ["họ và tên", "họ tên", "họ, chữ đệm và tên"]):
                            start_extracting = True

                    if start_extracting:
                        if is_row_empty:
                            empty_consecutive += 1
                            if empty_consecutive >= 20:
                                break
                            continue
                        else:
                            empty_consecutive = 0

                        if is_summary_or_footer_row(row):
                            continue

                        # Bắt buộc: STT hợp lệ và Họ tên có chữ cái
                        if not is_valid_stt(cell_a_val) or is_cell_empty(name_cell_val) or not has_letters:
                            continue

                        row_vals = [cell.value for cell in row]
                        req_len = 31 if remove_qrs else 28
                        while len(row_vals) < req_len:
                            row_vals.append(None)

                        if remove_qrs:
                            # Cắt bỏ 3 cột Q, R, S (index 16, 17, 18)
                            new_row = (row_vals[:16] + row_vals[19:])[:28]
                        else:
                            new_row = row_vals[:28]

                        while len(new_row) < 28:
                            new_row.append(None)

                        clean_stt = int(cell_a_val) if isinstance(cell_a_val, (int, float)) and float(cell_a_val).is_integer() else str(cell_a_val or '')

                        raw_dob = new_row[8]
                        raw_school = new_row[13]
                        raw_cohort = new_row[14]
                        raw_grad = new_row[18]

                        init_school = clean_whitespace(str(raw_school)) if not is_cell_empty(raw_school) else ""
                        init_cohort = clean_whitespace(str(raw_cohort)) if not is_cell_empty(raw_cohort) else ""

                        student_dict = {
                            "file_path": file_path,
                            "folder_name": file_path.parent.name,
                            "file_name": file_path.name,
                            "sheet_name": ws_in.title,
                            "sheet_st_idx": len(sheet_students),
                            "row_idx": r_idx,
                            "clean_stt": clean_stt,
                            "raw_name": name_cell_str,
                            "name": vietnamese_title_case(name_cell_str),
                            "raw_dob": raw_dob,
                            "dob": format_date_to_text(raw_dob),
                            "raw_school": raw_school,
                            "school": init_school,
                            "raw_cohort": raw_cohort,
                            "cohort": init_cohort,
                            "raw_grad": raw_grad,
                            "grad_date": format_grad_date_to_text(raw_grad) if not is_cell_empty(raw_grad) else "",
                            "row_values": new_row,
                            "effective_context_school": effective_context_school,
                            "effective_context_cohort": effective_context_cohort,
                        }
                        sheet_students.append(student_dict)

                total_sheet = len(sheet_students)
                if total_sheet > 0:
                    for st in sheet_students:
                        st["sheet_total"] = total_sheet
                    file_students.extend(sheet_students)
                    log(f"  -> Trích xuất được {total_sheet} học sinh từ Sheet '{ws_in.title}'.", "info")
        finally:
            wb_in.close()

        return file_students

    def process_file(
        self,
        file_path: Path,
        unit_code: str,
        unit_name: str,
        allowed_schools: list[str],
        output_dir: Path | None = None,
        dialog_callback=None,
        log_callback=None,
        cache_state=None,
        save_individual_file: bool = False,
        check_cancel=None
    ) -> tuple[list[dict], dict]:
        """
        Trích xuất và chuẩn hóa học sinh từ 1 tệp Excel riêng lẻ (tuân thủ pipeline toàn cục).
        - Giữ nguyên giao diện hàm để tương thích hoàn toàn.
        - Trả về: (file_students, stats)
        """
        unit_info = {
            "code": unit_code,
            "name": unit_name,
            "schools": allowed_schools or []
        }
        all_students, overall_stats = self.process_all_files(
            files=[file_path],
            unit_info=unit_info,
            output_dir=output_dir,
            dialog_callback=dialog_callback,
            log_callback=log_callback,
            check_cancel=check_cancel,
            cache_state=cache_state,
            save_individual_file=save_individual_file
        )
        file_stats = {
            "file_name": file_path.name,
            "total_students": overall_stats["total_students"],
            "da_so_hoa": overall_stats["da_so_hoa"],
            "chua_day_du": overall_stats["chua_day_du"],
            "chua_so_hoa": overall_stats["chua_so_hoa"],
            "output_file": overall_stats.get("output_file"),
            "error_file": overall_stats.get("error_file")
        }
        return all_students, file_stats

    def process_all_files(
        self,
        files: list[Path],
        unit_info: dict | str,
        unit_name: str = "",
        allowed_schools: list[str] | None = None,
        output_dir: Path | None = None,
        export_options: dict | None = None,
        dialog_callback=None,
        log_callback=None,
        progress_callback=None,
        check_cancel=None,
        cache_state=None,
        save_individual_file: bool = False
    ) -> tuple[list[dict], dict]:
        """
        Quy trình chuẩn hóa toàn cục xuyên suốt tất cả các File & Sheet theo 4 bước tuần tự:
        - Giai đoạn 0: Quét & trích xuất thô toàn bộ học sinh từ tất cả các file và sheet vào RAM.
        - Giai đoạn 1: Chuẩn hóa toàn bộ Ngày sinh (Cột I) cho toàn bộ học sinh.
        - Giai đoạn 2: Chuẩn hóa toàn bộ Tên trường (Cột N) cho toàn bộ học sinh theo danh mục đơn vị.
        - Giai đoạn 3: Chuẩn hóa toàn bộ Niên khóa (Cột O) dựa trên Năm sinh đa số đã làm sạch ở Giai đoạn 1.
        - Giai đoạn 4: Chuẩn hóa toàn bộ Ngày tháng năm cấp bằng (Cột S) theo cặp (Trường, Niên khóa).
        - Giai đoạn 5: Hoàn tất dữ liệu, tính Trạng thái/Lý do và Xuất các tệp kết quả.
        """
        def log(msg, level="info"):
            if log_callback:
                log_callback(msg, level)
            else:
                print(msg)

        if isinstance(unit_info, dict):
            eff_unit_info = unit_info
            u_code = unit_info.get("code", "")
            u_name = unit_info.get("name", "")
            u_schools = unit_info.get("schools", [])
        else:
            u_code = str(unit_info)
            u_name = unit_name
            u_schools = allowed_schools or []
            eff_unit_info = self.lookup_mgr.get_unit_info(u_name) or {
                "code": u_code,
                "name": u_name,
                "schools": u_schools,
                "school_map": {}
            }

        if cache_state is None:
            cache_state = {
                "school_cache": {},
                "cohort_cache": {},
                "skip_all_dob": False
            }

        school_cache = cache_state.setdefault("school_cache", {})
        cohort_cache = cache_state.setdefault("cohort_cache", {})

        is_single_file = (len(files) == 1)

        overall_stats = {
            "total_files": len(files),
            "total_students": 0,
            "da_so_hoa": 0,
            "chua_day_du": 0,
            "chua_so_hoa": 0,
            "output_file": None,
            "error_file": None,
            "export_results": {}
        }

        def update_progress(curr, total, stage_msg=None):
            if progress_callback:
                try:
                    progress_callback(curr, total, overall_stats, stage_msg)
                except TypeError:
                    progress_callback(curr, total, overall_stats)

        # ==============================================================
        # GIAI ĐOẠN 0: QUÉT & TRÍCH XUẤT THÔ TỪ TẤT CẢ CÁC TỆP EXCEL
        # ==============================================================
        log("\n" + "=" * 80, "info")
        log("🔍 [GIAI ĐOẠN 0] BẮT ĐẦU QUÉT & TRÍCH XUẤT DỮ LIỆU TỪ TẤT CẢ CÁC TỆP...", "info")
        log("=" * 80, "info")

        all_unit_students = []
        sheet_map: dict[tuple[str, str], list[dict]] = {}

        for idx, f_path in enumerate(files, start=1):
            if check_cancel and check_cancel():
                log("⏹ Quá trình xử lý đã bị dừng theo yêu cầu của người dùng.", "warning")
                return [], overall_stats

            log(f"\n[{idx}/{len(files)}] ĐANG TRÍCH XUẤT: {f_path.name}...", "info")
            f_students = self.extract_students_from_file(
                file_path=f_path,
                allowed_schools=u_schools,
                log_callback=log_callback,
                check_cancel=check_cancel
            )

            for st in f_students:
                sheet_map.setdefault((st["file_name"], st["sheet_name"]), []).append(st)

            all_unit_students.extend(f_students)
            overall_stats["total_students"] = len(all_unit_students)

            update_progress(idx, len(files), f"Giai đoạn 0: Đang đọc tệp [{idx}/{len(files)}] {f_path.name}...")

        if not all_unit_students:
            log("⚠️ Không có học sinh nào được trích xuất từ các file nguồn.", "warning")
            return [], overall_stats

        total_extracted = len(all_unit_students)
        log(f"\n🎉 Hoàn thành Giai đoạn 0: Đã trích xuất tổng cộng {total_extracted} học sinh từ {len(files)} tệp.", "success")

        if check_cancel and check_cancel():
            log("⏹ Quá trình xử lý đã bị dừng theo yêu cầu của người dùng.", "warning")
            return all_unit_students, overall_stats

        # ==============================================================
        # GIAI ĐOẠN 1/4: CHUẨN HÓA TOÀN BỘ NGÀY SINH (CỘT I)
        # ==============================================================
        log("\n" + "=" * 80, "info")
        log("📅 [GIAI ĐOẠN 1/4] BẮT ĐẦU CHUẨN HÓA TOÀN BỘ NGÀY SINH (CỘT I)...", "info")
        log("=" * 80, "info")
        update_progress(1, 4, f"Giai đoạn 1/4: Đang chuẩn hóa Ngày sinh toàn cục ({total_extracted} học sinh)...")

        for idx_st, st_item in enumerate(all_unit_students):
            if check_cancel and check_cancel():
                log("⏹ Đã dừng chuẩn hóa ngày sinh theo yêu cầu.", "warning")
                return all_unit_students, overall_stats

            norm_dob = st_item["dob"]
            is_v_dob, err_t, err_d = check_dob_validity(norm_dob)
            st_item["is_valid_dob"] = is_v_dob
            st_item["dob_err_desc"] = err_d

            if not is_v_dob and not cache_state.get("skip_all_dob", False):
                if not is_cell_empty(norm_dob):
                    # Cửa sổ trượt động theo ngữ cảnh cùng sheet:
                    sheet_key = (st_item["file_name"], st_item["sheet_name"])
                    sheet_peers = sheet_map.get(sheet_key, [st_item])
                    total_sheet = len(sheet_peers)
                    peer_idx = st_item.get("sheet_st_idx", 0)

                    if peer_idx < 5:
                        ctx_start = 0
                        ctx_end = min(total_sheet, peer_idx + 11)
                    elif peer_idx >= total_sheet - 5:
                        ctx_start = max(0, peer_idx - 10)
                        ctx_end = total_sheet
                    else:
                        ctx_start = max(0, peer_idx - 5)
                        ctx_end = min(total_sheet, peer_idx + 6)

                    context_rows = []
                    for c_idx in range(ctx_start, ctx_end):
                        item_c = sheet_peers[c_idx]
                        context_rows.append({
                            "stt": item_c["clean_stt"],
                            "name": item_c["name"],
                            "dob": item_c["dob"],
                            "cohort": item_c["cohort"],
                            "grad": item_c["grad_date"],
                            "is_current": (c_idx == peer_idx)
                        })

                    sugg_dob = ""
                    cohort_text = str(st_item.get("cohort") or "") or st_item.get("effective_context_cohort") or ""
                    m_nk = re.search(r'(\d{4})$', cohort_text)
                    if m_nk:
                        y_expected = int(m_nk.group(1)) - 15
                    else:
                        y_expected = find_mode_birth_year(sheet_peers) or find_mode_birth_year(all_unit_students)

                    if y_expected:
                        m_dm = re.search(r'\b(\d{1,2})[/-](\d{1,2})\b', str(st_item["raw_dob"] or norm_dob))
                        if m_dm:
                            d_v, m_v = int(m_dm.group(1)), int(m_dm.group(2))
                            if 1 <= d_v <= 31 and 1 <= m_v <= 12:
                                cand = f"{d_v:02d}/{m_v:02d}/{y_expected}"
                                c_v, _, _ = check_dob_validity(cand)
                                if c_v:
                                    sugg_dob = cand

                    if dialog_callback:
                        student_info = {
                            "file_name": st_item["file_name"],
                            "sheet_name": st_item["sheet_name"],
                            "stt": st_item["clean_stt"],
                            "name": st_item["name"],
                            "raw_val": st_item["raw_dob"],
                            "curr_norm": norm_dob,
                            "err_desc": err_d,
                            "context_rows": context_rows
                        }
                        res_dob, skip_all_flag = dialog_callback("date_fix", student_info, "Ngày sinh", sugg_dob)
                        if skip_all_flag:
                            cache_state["skip_all_dob"] = True
                        if res_dob:
                            st_item["dob"] = res_dob
                            is_v_after, _, err_d_after = check_dob_validity(res_dob)
                            st_item["is_valid_dob"] = is_v_after
                            st_item["dob_err_desc"] = err_d_after

        log(f"✅ Hoàn thành Giai đoạn 1: Chuẩn hóa Ngày sinh cho toàn bộ {total_extracted} học sinh.", "success")

        if check_cancel and check_cancel():
            log("⏹ Quá trình xử lý đã bị dừng theo yêu cầu của người dùng.", "warning")
            return all_unit_students, overall_stats

        # ==============================================================
        # GIAI ĐOẠN 2/4: CHUẨN HÓA TOÀN BỘ TÊN TRƯỜNG (CỘT N)
        # ==============================================================
        log("\n" + "=" * 80, "info")
        log("🏫 [GIAI ĐOẠN 2/4] BẮT ĐẦU CHUẨN HÓA TOÀN BỘ TÊN TRƯỜNG (CỘT N)...", "info")
        log("=" * 80, "info")
        update_progress(2, 4, f"Giai đoạn 2/4: Đang chuẩn hóa Tên trường toàn cục...")

        def _norm_sch_key(s: str) -> str:
            return re.sub(r'\s+', ' ', str(s or "")).strip().lower()

        allowed_schools_clean = [clean_whitespace(s) for s in (u_schools or []) if clean_whitespace(s)]
        school_lookup = {_norm_sch_key(s): s for s in allowed_schools_clean}

        # Gom nhóm học sinh theo (file_path, sheet_name, raw_school)
        distinct_school_blocks: dict[tuple, list[dict]] = {}
        for st in all_unit_students:
            raw_sch = clean_whitespace(str(st.get("raw_school") or "")) if not is_cell_empty(st.get("raw_school")) else ""
            key = (st["file_path"], st["sheet_name"], raw_sch)
            distinct_school_blocks.setdefault(key, []).append(st)

        for (f_path, s_name, raw_sch), grp in distinct_school_blocks.items():
            if check_cancel and check_cancel():
                log("⏹ Đã dừng chuẩn hóa tên trường theo yêu cầu.", "warning")
                return all_unit_students, overall_stats

            raw_sch_lower = _norm_sch_key(raw_sch)

            # 1. Trùng khớp với danh mục trường trực thuộc DS_Don_Vi
            if raw_sch_lower in school_lookup:
                canonical_sch = school_lookup[raw_sch_lower]
                for st in grp:
                    st["school"] = canonical_sch
                log(f"   ✅ Tên trường '{raw_sch}' khớp chuẩn DS_Don_Vi: '{canonical_sch}' ({len(grp)} HS - {s_name})", "info")
                continue

            # Gợi ý ngữ cảnh từ Thư mục -> Tên File -> Tên Sheet
            sugg_from_folder = match_school_from_text(f_path.parent.name, allowed_schools_clean)
            sugg_from_file = match_school_from_text(f_path.stem, allowed_schools_clean)
            sugg_from_sheet = match_school_from_text(s_name, allowed_schools_clean)
            ctx_school_suggestions = []
            for s_opt in [sugg_from_folder, sugg_from_file, sugg_from_sheet]:
                if s_opt and s_opt not in ctx_school_suggestions:
                    ctx_school_suggestions.append(s_opt)
            default_ctx_school = sugg_from_folder or sugg_from_file or sugg_from_sheet or (allowed_schools_clean[0] if allowed_schools_clean else "")

            # 2. Bị trống -> Hiển thị dialog kèm Thư mục, File, Sheet để người dùng tự chọn
            if not raw_sch:
                empty_file_key = f"__EMPTY_SCH__{f_path}"
                if empty_file_key in school_cache:
                    applied_sch = school_cache[empty_file_key]
                    for st in grp:
                        st["school"] = applied_sch
                    log(f"   ⚡ Tự động áp dụng Tên trường đã nhớ cho {len(grp)} HS bị trống trường ({f_path.name} | {s_name}): '{applied_sch}'", "info")
                    continue

                block_info = {
                    "folder_name": f_path.parent.name,
                    "file_name": f_path.name,
                    "sheet_name": s_name,
                    "start_stt": grp[0]["clean_stt"],
                    "end_stt": grp[-1]["clean_stt"],
                    "count": len(grp),
                    "raw_val": "(Trống)",
                    "fuzzy_suggestions": ctx_school_suggestions,
                    "context_rows": [{
                        "stt": st["clean_stt"],
                        "name": st["name"],
                        "dob": st["dob"],
                        "cohort": st["cohort"],
                        "grad": st["grad_date"],
                        "is_in_block": True
                    } for st in grp[:15]]
                }
                if dialog_callback:
                    res_sch, rem = dialog_callback("school_block", block_info, allowed_schools_clean, default_ctx_school)
                    applied_sch = res_sch if res_sch is not None else ""
                    if rem and applied_sch:
                        school_cache[empty_file_key] = applied_sch
                else:
                    applied_sch = default_ctx_school
                    if applied_sch:
                        school_cache[empty_file_key] = applied_sch

                for st in grp:
                    st["school"] = applied_sch
                if applied_sch:
                    log(f"   📍 Đã áp dụng Tên trường '{applied_sch}' cho {len(grp)} HS bị trống trường ({f_path.name} | {s_name})", "info")
                continue

            # 3. Không trùng với danh mục -> Kiểm tra cache hoặc mở dialog gợi ý gần giống
            if raw_sch in school_cache:
                applied_sch = school_cache[raw_sch]
                for st in grp:
                    st["school"] = applied_sch
                log(f"   ⚡ Tự động áp dụng Tên trường đã nhớ cho {len(grp)} HS: '{applied_sch}' ({f_path.name} | {s_name})", "info")
                continue

            fuzzy_suggs = fuzzy_match_school(raw_sch, allowed_schools_clean, n=3) if allowed_schools_clean else []
            default_sugg = fuzzy_suggs[0] if fuzzy_suggs else default_ctx_school
            block_info = {
                "folder_name": f_path.parent.name,
                "file_name": f_path.name,
                "sheet_name": s_name,
                "start_stt": grp[0]["clean_stt"],
                "end_stt": grp[-1]["clean_stt"],
                "count": len(grp),
                "raw_val": raw_sch,
                "fuzzy_suggestions": fuzzy_suggs,
                "context_rows": [{
                    "stt": st["clean_stt"],
                    "name": st["name"],
                    "dob": st["dob"],
                    "cohort": st["cohort"],
                    "grad": st["grad_date"],
                    "is_in_block": True
                } for st in grp[:15]]
            }

            if dialog_callback:
                res_sch, rem = dialog_callback("school_block", block_info, allowed_schools_clean, default_sugg)
                applied_sch = res_sch if res_sch is not None else ""
                if rem and applied_sch:
                    school_cache[raw_sch] = applied_sch
            else:
                applied_sch = default_sugg
                if applied_sch:
                    school_cache[raw_sch] = applied_sch

            for st in grp:
                st["school"] = applied_sch

        log(f"✅ Hoàn thành Giai đoạn 2: Chuẩn hóa Tên trường cho toàn bộ {total_extracted} học sinh.", "success")

        if check_cancel and check_cancel():
            log("⏹ Quá trình xử lý đã bị dừng theo yêu cầu của người dùng.", "warning")
            return all_unit_students, overall_stats

        # ==============================================================
        # GIAI ĐOẠN 3/4: CHUẨN HÓA TOÀN BỘ NIÊN KHÓA (CỘT O)
        # ==============================================================
        log("\n" + "=" * 80, "info")
        log("🎓 [GIAI ĐOẠN 3/4] BẮT ĐẦU CHUẨN HÓA TOÀN BỘ NIÊN KHÓA (CỘT O)...", "info")
        log("=" * 80, "info")
        update_progress(3, 4, f"Giai đoạn 3/4: Đang chuẩn hóa Niên khóa toàn cục...")

        distinct_cohort_blocks: dict[tuple, list[dict]] = {}
        for st in all_unit_students:
            raw_nk = clean_whitespace(str(st.get("raw_cohort") or "")) if not is_cell_empty(st.get("raw_cohort")) else ""
            key = (st["file_path"], st["sheet_name"], raw_nk)
            distinct_cohort_blocks.setdefault(key, []).append(st)

        for (f_path, s_name, raw_nk), grp in distinct_cohort_blocks.items():
            if check_cancel and check_cancel():
                log("⏹ Đã dừng chuẩn hóa niên khóa theo yêu cầu.", "warning")
                return all_unit_students, overall_stats

            # Năm sinh đa số: tìm từ chính nhóm này (ngày sinh đã được chuẩn hóa sạch ở Giai đoạn 1)
            mode_yr = find_mode_birth_year(grp)
            if not mode_yr:
                sheet_key = (f_path.name, s_name)
                mode_yr = find_mode_birth_year(sheet_map.get(sheet_key, [])) or find_mode_birth_year(all_unit_students)

            is_valid_diff_4, y1, y2, fmt_nk = parse_nien_khoa(raw_nk)
            has_range_sep = bool(re.search(r'\d+\s*[-/–—._~]\s*\d+', raw_nk))
            m_single = re.match(r'^\s*(19\d{2}|20\d{2})\s*$', raw_nk)
            is_empty_nk = is_cell_empty(raw_nk)

            # Trường hợp 1: Y_ra - Y_vao = 4
            if has_range_sep and is_valid_diff_4:
                if (mode_yr is None) or (y2 == mode_yr + 15):
                    for st in grp:
                        st["cohort"] = fmt_nk
                    log(f"   ✅ Niên khóa '{fmt_nk}' ({y1}-{y2}) chuẩn 4 năm, khớp năm sinh {mode_yr or 'N/A'} ({len(grp)} HS - {s_name})", "info")
                    continue
                else:
                    default_y_ra = y2
                    reason = f"Năm sinh đa số là {mode_yr} (dự kiến tốt nghiệp {mode_yr + 15}) nhưng niên khóa ghi {fmt_nk}"

            # Trường hợp 2: Có 2 năm nhưng Y_ra - Y_vao != 4
            elif has_range_sep and (not is_valid_diff_4):
                default_y_ra = y2 if y2 else (mode_yr + 15 if mode_yr else "")
                reason = f"Hiệu số năm ({y1}-{y2}) khác 4 năm chuẩn THCS"

            # Trường hợp 3: Chỉ có năm tốt nghiệp (ví dụ 2022)
            elif m_single:
                default_y_ra = int(m_single.group(1))
                reason = f"Dữ liệu gốc chỉ có năm tốt nghiệp ({default_y_ra}), cần xác nhận niên khóa 4 năm"

            # Trường hợp 4: Niên khóa bị trống
            elif is_empty_nk:
                if mode_yr:
                    default_y_ra = mode_yr + 15
                    reason = f"Niên khóa bị trống (Gợi ý tốt nghiệp năm {default_y_ra} từ năm sinh đa số {mode_yr})"
                else:
                    cohort_ctx = extract_cohort_from_text(f_path.stem) or extract_cohort_from_text(s_name)
                    if cohort_ctx:
                        _, _, y_ctx_ra, _ = parse_nien_khoa(cohort_ctx)
                        default_y_ra = y_ctx_ra or ""
                    else:
                        default_y_ra = ""
                    reason = "Niên khóa bị trống"

            # Trường hợp 5: Dạng khác
            else:
                default_y_ra = mode_yr + 15 if mode_yr else ""
                reason = f"Định dạng niên khóa không hợp lệ: '{raw_nk}'"

            # Kiểm tra cache trước khi mở dialog (kèm cơ chế bảo vệ an toàn 2 lớp theo năm sinh)
            cohort_cache_key = raw_nk if raw_nk else "__EMPTY_COHORT__"
            if cohort_cache_key in cohort_cache:
                applied_nk = cohort_cache[cohort_cache_key]
                _, _, y_applied_ra, _ = parse_nien_khoa(applied_nk)

                # Chỉ tự động áp dụng khi Năm ra trường đã nhớ khớp với Năm sinh đa số + 15 của nhóm này
                if mode_yr and y_applied_ra and (y_applied_ra == mode_yr + 15):
                    for st in grp:
                        st["cohort"] = applied_nk
                    log(f"   ⚡ Tự động áp dụng Niên khóa đã nhớ cho {len(grp)} HS: '{applied_nk}' (khớp năm sinh đa số {mode_yr} + 15) ({f_path.name} | {s_name})", "info")
                    continue
                elif mode_yr and y_applied_ra:
                    default_y_ra = mode_yr + 15
                    reason = f"Đã nhớ '{applied_nk}' cho nhóm tương tự, nhưng nhóm này có năm sinh đa số là {mode_yr} (dự kiến tốt nghiệp {mode_yr + 15} != {y_applied_ra})"
                    log(f"   ⚠️ Phát hiện lệch tuổi: Nhóm '{raw_nk or '(Trống)'}' có năm sinh {mode_yr} (dự kiến {mode_yr + 15}), khác niên khóa đã nhớ '{applied_nk}'. Mở dialog xác nhận thủ công.", "warning")
                elif not mode_yr:
                    reason = f"Đã nhớ '{applied_nk}' cho nhóm tương tự, nhưng nhóm này không có thông tin năm sinh để đối chiếu an toàn"
                    log(f"   ⚠️ Nhóm '{raw_nk or '(Trống)'}' không có thông tin năm sinh để đối chiếu an toàn với niên khóa đã nhớ '{applied_nk}'. Mở dialog xác nhận thủ công.", "warning")

            # Mở dialog đối chiếu niên khóa (hiển thị tối đa 15 HS)
            block_info = {
                "folder_name": f_path.parent.name,
                "file_name": f_path.name,
                "sheet_name": s_name,
                "start_stt": grp[0]["clean_stt"],
                "end_stt": grp[-1]["clean_stt"],
                "count": len(grp),
                "raw_val": raw_nk if raw_nk else "(Trống)",
                "reason": reason,
                "context_rows": [{
                    "stt": st["clean_stt"],
                    "name": st["name"],
                    "dob": st["dob"],
                    "cohort": st["cohort"],
                    "grad": st["grad_date"],
                    "is_in_block": True
                } for st in grp[:15]]
            }

            if dialog_callback:
                res_nk, rem = dialog_callback("cohort_block", block_info, default_y_ra, reason)
                if res_nk:
                    chosen_nk = res_nk
                else:
                    if default_y_ra and str(default_y_ra).isdigit():
                        chosen_nk = f"{int(default_y_ra) - 4}-{default_y_ra}"
                    else:
                        chosen_nk = fmt_nk if fmt_nk else raw_nk

                if rem:
                    cohort_cache[cohort_cache_key] = chosen_nk
            else:
                if default_y_ra and str(default_y_ra).isdigit():
                    chosen_nk = f"{int(default_y_ra) - 4}-{default_y_ra}"
                else:
                    chosen_nk = fmt_nk if fmt_nk else raw_nk
                cohort_cache[cohort_cache_key] = chosen_nk

            for st in grp:
                st["cohort"] = chosen_nk

        log(f"✅ Hoàn thành Giai đoạn 3: Chuẩn hóa Niên khóa cho toàn bộ {total_extracted} học sinh.", "success")

        if check_cancel and check_cancel():
            log("⏹ Quá trình xử lý đã bị dừng theo yêu cầu của người dùng.", "warning")
            return all_unit_students, overall_stats

        # ==============================================================
        # GIAI ĐOẠN 4/4: CHUẨN HÓA TOÀN BỘ NGÀY CẤP BẰNG (CỘT S)
        # ==============================================================
        log("\n" + "=" * 80, "info")
        log("📜 [GIAI ĐOẠN 4/4] BẮT ĐẦU CHUẨN HÓA NGÀY CẤP BẰNG TOÀN CỤC (CỘT S)...", "info")
        log("=" * 80, "info")
        update_progress(4, 4, f"Giai đoạn 4/4: Đang chuẩn hóa Ngày cấp bằng toàn cục...")
        auto_fill_graduation_dates_by_school_and_cohort(all_unit_students)
        log(f"✅ Hoàn thành Giai đoạn 4: Đã tự động đối khớp và chuẩn hóa Ngày cấp bằng cho toàn bộ {total_extracted} học sinh.", "success")

        if check_cancel and check_cancel():
            log("⏹ Quá trình xử lý đã bị dừng theo yêu cầu của người dùng.", "warning")
            return all_unit_students, overall_stats

        # ==============================================================
        # GIAI ĐOẠN 5: HOÀN TẤT DỮ LIỆU & XUẤT CÁC TỆP KẾT QUẢ
        # ==============================================================
        log("\n" + "=" * 80, "info")
        log("📋 [GIAI ĐOẠN 5] HOÀN TẤT DỮ LIỆU & BẮT ĐẦU XUẤT CÁC TỆP KẾT QUẢ...", "info")
        log("=" * 80, "info")
        update_progress(4, 4, "Giai đoạn 5: Đang hoàn tất dữ liệu & xuất các tệp kết quả...")

        overall_stats["total_students"] = len(all_unit_students)

        for st_item in all_unit_students:
            row_vals = list(st_item["row_values"])

            # Cột D (Index 3): Tên văn bằng
            row_vals[3] = "Bằng tốt nghiệp trung học cơ sở"
            # Cột E (Index 4): Họ và tên Title Case
            row_vals[4] = st_item["name"]
            # Cột I (Index 8): Ngày sinh chuẩn dd/MM/yyyy
            row_vals[8] = st_item["dob"]
            # Cột J (Index 9): Giới tính
            row_vals[9] = self.lookup_mgr.map_gender(row_vals[9])
            # Cột K (Index 10): Dân tộc (Title Case)
            row_vals[10] = normalize_ethnic(row_vals[10], self.lookup_mgr)
            # Cột L (Index 11): Quốc tịch (Title Case)
            row_vals[11] = normalize_nationality(row_vals[11], self.lookup_mgr)
            # Cột M (Index 12): Nơi sinh (Title Case + Thành phố/Thị xã/Thị trấn)
            row_vals[12] = normalize_birthplace(row_vals[12], self.lookup_mgr)
            # Cột N (Index 13): Tên trường đã học đã chuẩn hóa
            row_vals[13] = st_item["school"]
            # Cột O (Index 14): Niên khóa đã chuẩn hóa
            row_vals[14] = st_item["cohort"]
            # Cột S (Index 18): Ngày tháng năm cấp bằng
            row_vals[18] = st_item["grad_date"]
            # Cột V (Index 21): Người ký bằng Title Case
            if not is_cell_empty(row_vals[21]):
                row_vals[21] = vietnamese_title_case(str(row_vals[21]))
            else:
                row_vals[21] = None

            # Cột X (Index 23): Ghi chú -> Giữ nguyên ghi chú gốc nếu có, nếu rỗng thì để trống
            raw_note_val = row_vals[23]
            orig_note = clean_whitespace(str(raw_note_val)) if not is_cell_empty(raw_note_val) else ""
            row_vals[23] = orig_note if orig_note else None

            # Cột Y (Index 24) & Cột Z (Index 25): Trạng thái số hóa & Nguyên nhân
            status_y, reason_z = determine_status_and_reason(
                st_item["raw_dob"],
                st_item["dob"],
                st_item["is_valid_dob"],
                st_item["dob_err_desc"]
            )
            row_vals[24] = status_y
            row_vals[25] = reason_z if reason_z else None

            # Cột AA (Index 26): Mã đơn vị
            row_vals[26] = u_code
            # Cột AB (Index 27): Tên đơn vị
            row_vals[27] = u_name

            st_item["final_row"] = row_vals

            if status_y == "Đã số hóa đầy đủ, chính xác":
                overall_stats["da_so_hoa"] += 1
            elif status_y == "Số hoá chưa đầy đủ, chính xác":
                overall_stats["chua_day_du"] += 1
            else:
                overall_stats["chua_so_hoa"] += 1

        # Lưu file riêng lẻ nếu là chế độ đơn tệp
        if (is_single_file or save_individual_file) and output_dir and all_unit_students:
            output_dir.mkdir(parents=True, exist_ok=True)
            out_wb, out_ws = self.exporter._create_base_workbook("Data")
            self.exporter._write_student_rows(out_ws, all_unit_students, renumber=True)

            out_file_path = output_dir / files[0].name
            try:
                out_wb.save(out_file_path)
                overall_stats["output_file"] = out_file_path
                log(f"  ✅ Đã lưu file chuẩn hóa đơn lẻ: Ket_Qua_So_Hoa/{out_file_path.name}", "success")
            except PermissionError:
                log(f"❌ Không thể ghi file '{out_file_path.name}' do đang mở trong Excel.", "error")
            finally:
                out_wb.close()

        # Xuất các file gộp, tách, file DS_Loi duy nhất, file Dữ liệu gốc và Tài liệu số hóa
        if output_dir:
            log("\n📦 BẮT ĐẦU XUẤT CÁC TỆP DỮ LIỆU THEO TÙY CHỌN...", "info")
            export_res = self.exporter.export_all(
                all_students=all_unit_students,
                unit_info=eff_unit_info,
                output_dir=output_dir,
                options=export_options or {},
                log_callback=log_callback
            )
            overall_stats["export_results"] = export_res

        update_progress(len(files), len(files), "Hoàn tất chuẩn hóa & xuất tệp thành công!")

        return all_unit_students, overall_stats
