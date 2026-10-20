from __future__ import annotations

import io
import re
import copy
import unicodedata
from datetime import datetime
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from src.lookup_manager import get_data_dir
from src.normalizer import (
    clean_whitespace,
    is_cell_empty,
    parse_nien_khoa,
    extract_birth_year_from_dob
)

def sanitize_filename(name: str) -> str:
    r"""Loại bỏ các ký tự cấm của Windows/macOS (\ / : * ? " < > |) khỏi tên file/thư mục. Giữ nguyên tiếng Việt có dấu và ký tự & hoàn toàn hợp lệ."""
    if not name:
        return "Chua_Xac_Dinh"
    s = unicodedata.normalize('NFC', str(name)).strip()
    s = re.sub(r'[\\/*?:"<>|]', '_', s)
    s = re.sub(r'\s+', ' ', s).strip()
    s = re.sub(r'_+', '_', s).strip('._ ')
    return s if s else "Chua_Xac_Dinh"

def get_graduation_year(student: dict) -> int | str:
    """Xác định năm ra trường (tốt nghiệp) của học sinh để gom nhóm theo năm."""
    cohort = student.get("cohort", "")
    _, _, y2, _ = parse_nien_khoa(cohort)
    if y2:
        return y2
    
    # Thử lấy từ ngày cấp bằng
    grad = student.get("grad_date", "")
    m = re.search(r'\b(19\d{2}|20\d{2})\b', str(grad))
    if m:
        return int(m.group(1))
        
    # Thử lấy từ năm sinh + 15
    dob = student.get("dob", "")
    y_birth = extract_birth_year_from_dob(dob)
    if y_birth:
        return y_birth + 15
        
    return "Chua_Xac_Dinh"

def sort_key_by_school(st: dict):
    """Khóa sắp xếp: Tên trường -> Năm ra trường -> STT gốc."""
    sch = (st.get("school") or "").lower()
    yr = get_graduation_year(st)
    yr_num = yr if isinstance(yr, int) else 9999
    stt = st.get("clean_stt")
    stt_num = stt if isinstance(stt, (int, float)) else 999999
    return (sch, yr_num, stt_num)

def sort_key_by_year(st: dict):
    """Khóa sắp xếp: Năm ra trường -> Tên trường -> STT gốc."""
    yr = get_graduation_year(st)
    yr_num = yr if isinstance(yr, int) else 9999
    sch = (st.get("school") or "").lower()
    stt = st.get("clean_stt")
    stt_num = stt if isinstance(stt, (int, float)) else 999999
    return (yr_num, sch, stt_num)


class DataExporter:
    """Chuyên trách việc xuất các file Gộp, file Tách, file DS_Lỗi duy nhất, Dữ liệu gốc và Danh sách tài liệu số hoá."""
    def __init__(self, template_bytes: bytes, cached_fonts: list, yellow_fill: PatternFill, doc_template_path: Path | None = None):
        self.template_bytes = template_bytes
        self.cached_fonts = cached_fonts
        self.yellow_fill = yellow_fill

        # Nạp template cho mô-đun 2: Danh sách tài liệu số hóa
        if doc_template_path is None:
            self.doc_template_path = get_data_dir() / "Mau_DS_Tai_Lieu_So_Hoa.xlsx"
        else:
            self.doc_template_path = Path(doc_template_path)

        if self.doc_template_path.exists():
            with open(self.doc_template_path, "rb") as f:
                self.doc_template_bytes = f.read()
            wb_doc = openpyxl.load_workbook(io.BytesIO(self.doc_template_bytes))
            ws_doc = wb_doc.active
            self.cached_doc_fonts = []
            for col in range(1, 12):
                cell = ws_doc.cell(row=2, column=col)
                self.cached_doc_fonts.append(Font(name=cell.font.name, size=cell.font.size) if (cell.has_style and cell.font) else None)
            wb_doc.close()
        else:
            self.doc_template_bytes = None
            self.cached_doc_fonts = []

    def _create_base_workbook(self, sheet_title="Data") -> tuple[openpyxl.Workbook, openpyxl.worksheet.worksheet.Worksheet]:
        """Tạo workbook chuẩn từ template 28 cột, xóa sạch các dòng trống thừa sau dòng header 2."""
        wb = openpyxl.load_workbook(io.BytesIO(self.template_bytes))
        ws = wb.active
        ws.title = sheet_title
        if ws.max_row > 2:
            ws.delete_rows(3, ws.max_row - 2)
        return wb, ws

    def _create_doc_workbook(self, sheet_title="Data") -> tuple[openpyxl.Workbook, openpyxl.worksheet.worksheet.Worksheet]:
        """Tạo workbook chuẩn từ template Mau_DS_Tai_Lieu_So_Hoa.xlsx, xóa sạch các dòng thừa sau dòng 2."""
        if not self.doc_template_bytes:
            raise FileNotFoundError(f"Không tìm thấy template Danh sách tài liệu số hóa tại: {self.doc_template_path}")
        wb = openpyxl.load_workbook(io.BytesIO(self.doc_template_bytes))
        ws = wb.active
        ws.title = sheet_title
        if ws.max_row > 2:
            ws.delete_rows(3, ws.max_row - 2)
        return wb, ws

    def _write_student_rows(self, ws, students: list[dict], renumber: bool = True) -> int:
        """
        Ghi danh sách học sinh vào worksheet:
        - Đánh lại STT liên tục.
        - ÉP TOÀN BỘ CỘT NGÀY THÁNG (Cột I, S) VÀ CỘT VĂN BẢN VỀ KIỂU TEXT (@) TUYỆT ĐỐI.
        - Tô vàng các ô bị thiếu từ cột B (2) đến AB (28).
        """
        yellow_count = 0
        for idx, st in enumerate(students, start=1):
            row_vals = list(st["final_row"])
            if renumber:
                row_vals[0] = idx
                
            ws.append(row_vals)
            curr_row = ws.max_row
            
            # Gán font và ép định dạng TEXT
            for col_idx, font in enumerate(self.cached_fonts, start=1):
                cell = ws.cell(row=curr_row, column=col_idx)
                if font:
                    cell.font = font
                    
                # Ép định dạng TEXT cho các cột:
                # Cột 9 (I: Ngày sinh), Cột 19 (S: Ngày cấp bằng), Cột 2 (B: Số hiệu), Cột 3 (C: Số vào sổ),
                # Cột 6 (F: Mã người học), Cột 7 (G: CCCD), Cột 8 (H: Hộ chiếu), Cột 24 (X: Ghi chú), Cột 27 (AA: Mã đơn vị)
                if col_idx in (2, 3, 6, 7, 8, 9, 19, 24, 27):
                    cell.number_format = '@'
                    if cell.value is not None:
                        cell.data_type = 's'
                        cell.value = str(cell.value)
                    
            # Tô vàng các ô bị thiếu từ cột B (2) đến AB (28)
            for col in range(1, min(28, len(row_vals))):
                if is_cell_empty(row_vals[col]):
                    ws.cell(row=curr_row, column=col + 1).fill = self.yellow_fill
                    yellow_count += 1
                    
        return yellow_count

    def export_all(
        self,
        all_students: list[dict],
        unit_info: dict,
        output_dir: Path,
        options: dict,
        log_callback=None
    ) -> dict:
        """
        Thực hiện toàn bộ các tùy chọn xuất file của Mô-đun 1 và Mô-đun 2:
        MÔ-ĐUN 1 (Chuẩn hóa văn bằng THCS):
        - Gộp theo trường
        - Gộp theo năm
        - Tách theo trường
        - Tách theo năm
        - Tách theo trường và năm
        - 1 File DS_Loi duy nhất kèm 4 cột truy vết nguồn (Cột AC..AF)
        - 1 File Du_Lieu_Goc tổng hợp
        
        MÔ-ĐUN 2 (Danh sách tài liệu số hoá):
        - Xuất tài liệu số hóa theo trường
        - Xuất tài liệu số hóa theo năm
        - Xuất tài liệu số hóa theo trường và năm
        """
        def log(msg, level="info"):
            if log_callback:
                log_callback(msg, level)
            else:
                print(msg)

        results = {
            "merged_school_file": None,
            "merged_year_file": None,
            "split_school_files": [],
            "split_year_files": [],
            "split_school_year_files": [],
            "error_file": None,
            "raw_file": None,
            "doc_school_file": None,
            "doc_year_file": None,
            "doc_school_year_file": None
        }

        if not all_students:
            return results

        unit_code = unit_info.get("code", "")
        unit_name = unit_info.get("name", "")
        unit_safe = sanitize_filename(unit_name)
        output_dir.mkdir(parents=True, exist_ok=True)

        # ----------------------------------------------------------------------
        # 1. XUẤT FILE DỮ LIỆU GỐC TỔNG HỢP: {Ten_Don_Vi}_Du_Lieu_Goc.xlsx
        # ----------------------------------------------------------------------
        try:
            raw_wb, raw_ws = self._create_base_workbook("Du_Lieu_Goc")
            raw_ws.cell(row=2, column=29, value="Tệp nguồn")
            raw_ws.cell(row=2, column=30, value="Sheet nguồn")
            raw_ws.cell(row=2, column=31, value="Dòng gốc")
            
            for idx, st in enumerate(all_students, start=1):
                raw_row = list(st["row_values"])
                raw_row[0] = idx
                raw_row.extend([st.get("file_name", ""), st.get("sheet_name", ""), st.get("row_idx", "")])
                raw_ws.append(raw_row)
                c_row = raw_ws.max_row
                for col_idx, font in enumerate(self.cached_fonts, start=1):
                    if font:
                        raw_ws.cell(row=c_row, column=col_idx).font = font
                # Định dạng TEXT cho các cột ngày tháng
                for c_t in (9, 19):
                    c_cell = raw_ws.cell(row=c_row, column=c_t)
                    c_cell.number_format = '@'
                    if c_cell.value is not None:
                        c_cell.data_type = 's'
                        c_cell.value = str(c_cell.value)

            raw_file_path = output_dir / f"{unit_safe}_Du_Lieu_Goc.xlsx"
            raw_wb.save(raw_file_path)
            raw_wb.close()
            results["raw_file"] = raw_file_path
            log(f"  📄 Đã lưu file Dữ liệu gốc tổng hợp: {raw_file_path.name}", "info")
        except Exception as e:
            log(f"⚠️ Lỗi khi xuất file Dữ liệu gốc tổng hợp: {e}", "warning")

        # ----------------------------------------------------------------------
        # 2. XUẤT FILE GỘP THEO TRƯỜNG: {Ten_Don_Vi}_Đa_Gop_Theo_Truong.xlsx
        # ----------------------------------------------------------------------
        students_by_school = sorted(all_students, key=sort_key_by_school)
        if options.get("gop_theo_truong", True):
            try:
                m_wb, m_ws = self._create_base_workbook("Data")
                self._write_student_rows(m_ws, students_by_school, renumber=True)
                out_p = output_dir / f"{unit_safe}_Đa_Gop_Theo_Truong.xlsx"
                m_wb.save(out_p)
                m_wb.close()
                results["merged_school_file"] = out_p
                log(f"  ✅ Đã xuất file Gộp theo trường: {out_p.name} ({len(students_by_school)} học sinh)", "success")
            except Exception as e:
                log(f"❌ Lỗi khi xuất file Gộp theo trường: {e}", "error")

        # ----------------------------------------------------------------------
        # 3. XUẤT FILE GỘP THEO NĂM: {Ten_Don_Vi}_Đa_Gop_Theo_Nam.xlsx
        # ----------------------------------------------------------------------
        students_by_year = sorted(all_students, key=sort_key_by_year)
        if options.get("gop_theo_nam", False):
            try:
                m_wb, m_ws = self._create_base_workbook("Data")
                self._write_student_rows(m_ws, students_by_year, renumber=True)
                out_p = output_dir / f"{unit_safe}_Đa_Gop_Theo_Nam.xlsx"
                m_wb.save(out_p)
                m_wb.close()
                results["merged_year_file"] = out_p
                log(f"  ✅ Đã xuất file Gộp theo năm: {out_p.name} ({len(students_by_year)} học sinh)", "success")
            except Exception as e:
                log(f"❌ Lỗi khi xuất file Gộp theo năm: {e}", "error")

        # ----------------------------------------------------------------------
        # 4. XUẤT 1 FILE DS_LOI DUY NHẤT: {Ten_Don_Vi}_DS_Loi.xlsx
        # ----------------------------------------------------------------------
        error_students = [
            st for st in students_by_school
            if len(st.get("final_row", [])) > 24 and st["final_row"][24] != "Đã số hóa đầy đủ, chính xác"
        ]
        if error_students:
            try:
                err_wb, err_ws = self._create_base_workbook("DS_Loi")
                # Thêm 4 cột chỉ đường truy vết nguồn gốc (Cột AC đến AF)
                err_ws.cell(row=2, column=29, value="Tệp nguồn gốc")
                err_ws.cell(row=2, column=30, value="Sheet nguồn")
                err_ws.cell(row=2, column=31, value="Dòng Excel gốc")
                err_ws.cell(row=2, column=32, value="STT ban đầu")

                for col_c in range(29, 33):
                    c_ref = err_ws.cell(row=2, column=col_c)
                    if self.cached_fonts and self.cached_fonts[0]:
                        c_ref.font = copy.copy(self.cached_fonts[0])

                for err_idx, st in enumerate(error_students, start=1):
                    row_vals = list(st["final_row"])
                    row_vals[0] = err_idx
                    row_vals.extend([
                        st.get("file_name", ""),
                        st.get("sheet_name", ""),
                        f"Dòng {st.get('row_idx', '')}",
                        st.get("clean_stt", "")
                    ])
                    err_ws.append(row_vals)
                    c_row = err_ws.max_row
                    for col_idx, font in enumerate(self.cached_fonts, start=1):
                        cell_c = err_ws.cell(row=c_row, column=col_idx)
                        if font:
                            cell_c.font = font
                        if col_idx in (2, 3, 6, 7, 8, 9, 19, 24, 27):
                            cell_c.number_format = '@'
                            if cell_c.value is not None:
                                cell_c.data_type = 's'
                                cell_c.value = str(cell_c.value)

                    for col in range(1, min(28, len(row_vals))):
                        if is_cell_empty(row_vals[col]):
                            err_ws.cell(row=c_row, column=col + 1).fill = self.yellow_fill

                err_file_path = output_dir / f"{unit_safe}_DS_Loi.xlsx"
                err_wb.save(err_file_path)
                err_wb.close()
                results["error_file"] = err_file_path
                log(f"  ⚠️ Đã xuất 1 file Danh sách lỗi duy nhất: {err_file_path.name} ({len(error_students)} học sinh lỗi kèm 4 cột truy vết nguồn)", "warning")
            except Exception as e:
                log(f"⚠️ Lỗi khi xuất file Danh sách lỗi duy nhất: {e}", "warning")

        # ----------------------------------------------------------------------
        # 5. TÁCH THEO TRƯỜNG: Da_Tach_Theo_Truong/{Ten_Truong}_Tu_{Min}_Den_{Max}.xlsx
        # ----------------------------------------------------------------------
        if options.get("tach_theo_truong", False):
            split_sch_dir = output_dir / "Da_Tach_Theo_Truong"
            split_sch_dir.mkdir(parents=True, exist_ok=True)
            
            school_groups = {}
            for st in students_by_school:
                sch_name = st.get("school") or "Chua_Xac_Dinh_Truong"
                school_groups.setdefault(sch_name, []).append(st)

            for sch_name, s_list in school_groups.items():
                sch_safe = sanitize_filename(sch_name)
                years = [get_graduation_year(st) for st in s_list if isinstance(get_graduation_year(st), int)]
                if years:
                    y_min, y_max = min(years), max(years)
                    fname = f"{sch_safe}_Tu_{y_min}_Den_{y_max}.xlsx"
                else:
                    fname = f"{sch_safe}.xlsx"

                try:
                    wb_s, ws_s = self._create_base_workbook("Data")
                    self._write_student_rows(ws_s, s_list, renumber=True)
                    fpath = split_sch_dir / fname
                    wb_s.save(fpath)
                    wb_s.close()
                    results["split_school_files"].append(fpath)
                except Exception as e:
                    log(f"⚠️ Lỗi khi tách trường {sch_name}: {e}", "warning")

            log(f"  📂 Đã tách thành {len(school_groups)} file theo trường tại: Da_Tach_Theo_Truong/", "success")

        # ----------------------------------------------------------------------
        # 6. TÁCH THEO NĂM: Da_Tach_Theo_Nam/{Ten_Don_Vi}_{Nam}.xlsx
        # ----------------------------------------------------------------------
        if options.get("tach_theo_nam", False):
            split_yr_dir = output_dir / "Da_Tach_Theo_Nam"
            split_yr_dir.mkdir(parents=True, exist_ok=True)
            
            year_groups = {}
            for st in students_by_year:
                yr = get_graduation_year(st)
                year_groups.setdefault(yr, []).append(st)

            for yr_val, s_list in year_groups.items():
                fname = f"{unit_safe}_{yr_val}.xlsx"
                try:
                    wb_y, ws_y = self._create_base_workbook("Data")
                    self._write_student_rows(ws_y, s_list, renumber=True)
                    fpath = split_yr_dir / fname
                    wb_y.save(fpath)
                    wb_y.close()
                    results["split_year_files"].append(fpath)
                except Exception as e:
                    log(f"⚠️ Lỗi khi tách theo năm {yr_val}: {e}", "warning")

            log(f"  📂 Đã tách thành {len(year_groups)} file theo năm tại: Da_Tach_Theo_Nam/", "success")

        # ----------------------------------------------------------------------
        # 7. TÁCH THEO TRƯỜNG VÀ NĂM: Tach_Theo_Truong_Va_Nam/{Ten_Truong}/{Ten_Truong}_{Nam}.xlsx
        # ----------------------------------------------------------------------
        if options.get("tach_theo_truong_va_nam", False):
            split_sch_yr_dir = output_dir / "Tach_Theo_Truong_Va_Nam"
            split_sch_yr_dir.mkdir(parents=True, exist_ok=True)

            matrix = {}
            for st in students_by_school:
                sch_name = st.get("school") or "Chua_Xac_Dinh_Truong"
                yr = get_graduation_year(st)
                matrix.setdefault(sch_name, {}).setdefault(yr, []).append(st)

            total_split_files = 0
            for sch_name, yr_dict in matrix.items():
                sch_safe = sanitize_filename(sch_name)
                sch_folder = split_sch_yr_dir / sch_safe
                sch_folder.mkdir(parents=True, exist_ok=True)

                for yr_val, s_list in yr_dict.items():
                    fname = f"{sch_safe}_{yr_val}.xlsx"
                    try:
                        wb_sy, ws_sy = self._create_base_workbook("Data")
                        self._write_student_rows(ws_sy, s_list, renumber=True)
                        fpath = sch_folder / fname
                        wb_sy.save(fpath)
                        wb_sy.close()
                        results["split_school_year_files"].append(fpath)
                        total_split_files += 1
                    except Exception as e:
                        log(f"⚠️ Lỗi khi tách {sch_name} năm {yr_val}: {e}", "warning")

            log(f"  📂 Đã tách thành {total_split_files} file theo Trường & Năm tại: Tach_Theo_Truong_Va_Nam/", "success")

        # ======================================================================
        # MÔ-ĐUN 2: XUẤT DANH SÁCH TÀI LIỆU SỐ HOÁ (3 TÙY CHỌN)
        # Mẫu: data/Mau_DS_Tai_Lieu_So_Hoa.xlsx (11 cột A..K)
        # ======================================================================
        self._export_digitization_documents(
            all_students=all_students,
            unit_info=unit_info,
            output_dir=output_dir,
            options=options,
            results=results,
            log=log
        )

        return results

    def _export_digitization_documents(
        self,
        all_students: list[dict],
        unit_info: dict,
        output_dir: Path,
        options: dict,
        results: dict,
        log
    ):
        """Xuất các file Danh sách tài liệu số hoá theo 3 tùy chọn: Theo Trường / Theo Năm / Theo Trường & Năm."""
        if not self.doc_template_bytes:
            log("⚠️ Không tìm thấy file mẫu Mau_DS_Tai_Lieu_So_Hoa.xlsx để xuất tài liệu số hóa.", "warning")
            return

        unit_code = unit_info.get("code", "")
        unit_name = unit_info.get("name", "")
        unit_safe = sanitize_filename(unit_name)
        school_map = unit_info.get("school_map", {})
        curr_date_str = datetime.now().strftime("%d/%m/%Y")

        def lookup_sch_code(s_name: str) -> str:
            if not s_name:
                return ""
            s_clean = s_name.strip().lower()
            if s_clean in school_map:
                return school_map[s_clean]
            # Khớp mềm
            s_alpha = re.sub(r'\W+', '', s_clean)
            for k, cd in school_map.items():
                if re.sub(r'\W+', '', k) == s_alpha:
                    return cd
            for k, cd in school_map.items():
                if k in s_clean or s_clean in k:
                    return cd
            return ""

        def write_doc_rows(ws, doc_rows):
            """Ghi các dòng tài liệu số hóa và định dạng font, kiểu TEXT @."""
            for r_idx, r_vals in enumerate(doc_rows, start=1):
                r_vals[0] = r_idx  # STT 1..N
                ws.append(r_vals)
                cur_r = ws.max_row
                for c_idx in range(1, 12):
                    cell = ws.cell(row=cur_r, column=c_idx)
                    if self.cached_doc_fonts and c_idx - 1 < len(self.cached_doc_fonts) and self.cached_doc_fonts[c_idx - 1]:
                        cell.font = self.cached_doc_fonts[c_idx - 1]
                    # Ép định dạng TEXT @ cho các cột text: Cột B (2), Cột D (4), Cột F (6), Cột G (7), Cột H (8 - Ngày lập sổ)
                    if c_idx in (2, 4, 6, 7, 8):
                        cell.number_format = '@'
                        if cell.value is not None:
                            cell.data_type = 's'
                            cell.value = str(cell.value)

        # ----------------------------------------------------------------------
        # TÙY CHỌN 1: Xuất theo Trường ({Ten_Don_Vi}_DS_Tai_Lieu_So_Hoa_Theo_Truong.xlsx)
        # ----------------------------------------------------------------------
        if options.get("doc_theo_truong", True):
            try:
                school_groups = {}
                for st in all_students:
                    sch_name = st.get("school") or "Chưa xác định trường"
                    school_groups.setdefault(sch_name, []).append(st)

                doc_rows = []
                for sch_name, s_list in school_groups.items():
                    sch_code = lookup_sch_code(sch_name)
                    so_goc = f"{unit_code}_{sch_code}" if sch_code else f"{unit_code}_{sanitize_filename(sch_name)}"

                    years = [get_graduation_year(st) for st in s_list if isinstance(get_graduation_year(st), int)]
                    if years:
                        y_min, y_max = min(years), max(years)
                        if y_min == y_max:
                            doc_title = f"Danh sách cấp bằng THCS trường {sch_name} năm {y_min}"
                        else:
                            doc_title = f"Danh sách cấp bằng THCS trường {sch_name} từ {y_min} đến {y_max}"
                    else:
                        y_min, y_max = "", ""
                        doc_title = f"Danh sách cấp bằng THCS trường {sch_name}"

                    # 11 cột: A(STT), B(Mã đơn vị), C(Tên đơn vị), D(Mã sổ gốc), E(Tên tài liệu),
                    #         F(Từ năm), G(Đến năm), H(Ngày lập sổ), I(Số trang), J(Tình trạng), K(Thông tư)
                    doc_rows.append([
                        0, unit_code, unit_name, so_goc, doc_title,
                        y_min, y_max, curr_date_str, None, "Nguyên vẹn", None
                    ])

                wb_doc, ws_doc = self._create_doc_workbook("Data")
                write_doc_rows(ws_doc, doc_rows)
                out_p = output_dir / f"{unit_safe}_DS_Tai_Lieu_So_Hoa_Theo_Truong.xlsx"
                wb_doc.save(out_p)
                wb_doc.close()
                results["doc_school_file"] = out_p
                log(f"  📑 Đã xuất Danh sách tài liệu số hoá theo Trường: {out_p.name} ({len(doc_rows)} trường)", "success")
            except Exception as e:
                log(f"❌ Lỗi khi xuất Danh sách tài liệu số hoá theo trường: {e}", "error")

        # ----------------------------------------------------------------------
        # TÙY CHỌN 2: Xuất theo Năm ({Ten_Don_Vi}_DS_Tai_Lieu_So_Hoa_Theo_Nam.xlsx)
        # ----------------------------------------------------------------------
        if options.get("doc_theo_nam", False):
            try:
                year_groups = {}
                for st in all_students:
                    yr = get_graduation_year(st)
                    year_groups.setdefault(yr, []).append(st)

                doc_rows = []
                # Sắp xếp theo năm tăng dần
                sorted_years = sorted(year_groups.keys(), key=lambda y: y if isinstance(y, int) else 9999)
                for yr_val in sorted_years:
                    so_goc = f"{unit_code}_{yr_val}"
                    doc_title = f"Danh sách cấp bằng THCS đơn vị {unit_name} năm {yr_val}"
                    doc_rows.append([
                        0, unit_code, unit_name, so_goc, doc_title,
                        yr_val, yr_val, curr_date_str, None, "Nguyên vẹn", None
                    ])

                wb_doc, ws_doc = self._create_doc_workbook("Data")
                write_doc_rows(ws_doc, doc_rows)
                out_p = output_dir / f"{unit_safe}_DS_Tai_Lieu_So_Hoa_Theo_Nam.xlsx"
                wb_doc.save(out_p)
                wb_doc.close()
                results["doc_year_file"] = out_p
                log(f"  📑 Đã xuất Danh sách tài liệu số hoá theo Năm: {out_p.name} ({len(doc_rows)} năm)", "success")
            except Exception as e:
                log(f"❌ Lỗi khi xuất Danh sách tài liệu số hoá theo năm: {e}", "error")

        # ----------------------------------------------------------------------
        # TÙY CHỌN 3: Xuất theo Trường và Năm ({Ten_Don_Vi}_DS_Tai_Lieu_So_Hoa_Theo_Truong_Va_Nam.xlsx)
        # ----------------------------------------------------------------------
        if options.get("doc_theo_truong_va_nam", False):
            try:
                matrix = {}
                for st in all_students:
                    sch_name = st.get("school") or "Chưa xác định trường"
                    yr = get_graduation_year(st)
                    matrix.setdefault(sch_name, {}).setdefault(yr, []).append(st)

                doc_rows = []
                for sch_name in sorted(matrix.keys()):
                    yr_dict = matrix[sch_name]
                    sch_code = lookup_sch_code(sch_name)
                    sorted_yrs = sorted(yr_dict.keys(), key=lambda y: y if isinstance(y, int) else 9999)

                    for yr_val in sorted_yrs:
                        so_goc = f"{unit_code}_{sch_code}_{yr_val}" if sch_code else f"{unit_code}_{sanitize_filename(sch_name)}_{yr_val}"
                        doc_title = f"Danh sách cấp bằng THCS trường {sch_name} năm {yr_val}"
                        doc_rows.append([
                            0, unit_code, unit_name, so_goc, doc_title,
                            yr_val, yr_val, curr_date_str, None, "Nguyên vẹn", None
                        ])

                wb_doc, ws_doc = self._create_doc_workbook("Data")
                write_doc_rows(ws_doc, doc_rows)
                out_p = output_dir / f"{unit_safe}_DS_Tai_Lieu_So_Hoa_Theo_Truong_Va_Nam.xlsx"
                wb_doc.save(out_p)
                wb_doc.close()
                results["doc_school_year_file"] = out_p
                log(f"  📑 Đã xuất Danh sách tài liệu số hoá theo Trường & Năm: {out_p.name} ({len(doc_rows)} mục)", "success")
            except Exception as e:
                log(f"❌ Lỗi khi xuất Danh sách tài liệu số hoá theo trường và năm: {e}", "error")
