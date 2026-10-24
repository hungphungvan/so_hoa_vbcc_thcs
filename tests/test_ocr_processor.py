import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

import openpyxl

from src.ocr_processor import (
    OCRConfigManager,
    OCRProcessor,
    natural_sort_key,
    parse_ocr_json_response,
)


class TestOCRProcessor(unittest.TestCase):
    def setUp(self):
        self.tmp_dir = Path(tempfile.mkdtemp())

    def test_natural_sort_key(self):
        files = ["anh_10.jpg", "anh_2.jpg", "anh_1.jpg", "anh_01.jpg", "anh_9.jpg"]
        sorted_files = sorted(files, key=natural_sort_key)
        self.assertEqual(sorted_files, ["anh_01.jpg", "anh_1.jpg", "anh_2.jpg", "anh_9.jpg", "anh_10.jpg"])

    def test_parse_ocr_json_response(self):
        raw = """```json
        [
          {
            "stt": 1,
            "sbd": "001",
            "ho_ten": "NGUYỄN VĂN AN",
            "gioi_tinh": "Nam",
            "ngay_sinh": "12/05/2008",
            "noi_sinh": "Thái Bình",
            "hoi_dong_thi": "THCS Yên Dương",
            "xep_loai": "Giỏi",
            "ghi_chu": "456/2023",
            "can_kiem_tra": null,
            "ly_do_nghi_ngo": null
          }
        ]
        ```"""
        records, err = parse_ocr_json_response(raw)
        self.assertIsNone(err)
        self.assertEqual(len(records), 1)
        self.assertEqual(records[0]["ho_ten"], "NGUYỄN VĂN AN")

        # Trường hợp không có mảng JSON
        records2, err2 = parse_ocr_json_response("Xin lỗi, ảnh bị mờ không đọc được")
        self.assertIsNone(records2)
        self.assertIn("mảng JSON", err2)

    def test_prompt_3_levels_hierarchy(self):
        cfg_file = self.tmp_dir / "ocr_config.json"
        mgr = OCRConfigManager(config_path=cfg_file)

        year_dir = self.tmp_dir / "Nam_2023"
        school_dir = year_dir / "Truong_01"
        school_dir.mkdir(parents=True)

        # Cấp 3: Mặc định hệ thống
        txt3, src3 = mgr.resolve_prompt_for_folder(school_dir)
        self.assertIn("DANH SÁCH HỌC SINH", txt3)
        self.assertIn("MẶC ĐỊNH", src3)

        # Cấp 2: Prompt chung của năm học
        (year_dir / "prompt.txt").write_text("Prompt nam hoc 2023", encoding="utf-8")
        txt2, src2 = mgr.resolve_prompt_for_folder(school_dir)
        self.assertEqual(txt2, "Prompt nam hoc 2023")
        self.assertIn("NĂM HỌC", src2)

        # Cấp 1: Prompt ưu tiên của trường
        (school_dir / "prompt.txt").write_text("Prompt rieng truong 01", encoding="utf-8")
        txt1, src1 = mgr.resolve_prompt_for_folder(school_dir)
        self.assertEqual(txt1, "Prompt rieng truong 01")
        self.assertIn("TRƯỜNG", src1)

    def test_export_to_original_excel(self):
        proc = OCRProcessor(config_mgr=OCRConfigManager(self.tmp_dir / "cfg.json"))
        sample_records = [
            {
                "stt": 1,
                "sbd": "012",
                "ho_ten": "nguyễn thị bình",
                "gioi_tinh": "Nữ",
                "ngay_sinh": "05/08/2008",
                "noi_sinh": "Nam Định",
                "hoi_dong_thi": "THCS Yên Dương",
                "xep_loai": "Khá",
                "ghi_chu": "Số 123",
                "can_kiem_tra": False,
                "ly_do_nghi_ngo": None,
            },
            {
                "stt": 2,
                "sbd": "013",
                "ho_ten": "trần văn cường",
                "gioi_tinh": "Nam",
                "ngay_sinh": "19/11/2008",
                "noi_sinh": "Hà Nội",
                "hoi_dong_thi": "THCS Yên Dương",
                "xep_loai": "TB",
                "ghi_chu": "",
                "can_kiem_tra": True,
                "ly_do_nghi_ngo": "Năm sinh sửa bằng mực đỏ hơi mờ",
            },
        ]

        orig_out = self.tmp_dir / "out_original.xlsx"
        proc.export_to_original_excel(sample_records, orig_out)
        self.assertTrue(orig_out.exists())

        wb1 = openpyxl.load_workbook(orig_out)
        ws1 = wb1.active
        self.assertEqual(ws1.cell(1, 3).value, "ho_ten")
        self.assertEqual(ws1.cell(2, 3).value, "nguyễn thị bình")
        self.assertEqual(ws1.cell(3, 11).value, "Năm sinh sửa bằng mực đỏ hơi mờ")
        wb1.close()

        # Kiểm tra với custom prompt chỉ có 3 cột hoàn toàn khác
        custom_records = [
            {"ma_hs": "HS01", "ten_day_du": "Lê Văn Tám", "diem_tb": 8.5},
            {"ma_hs": "HS02", "ten_day_du": "Võ Thị Sáu", "diem_tb": 9.0, "dan_toc": "Kinh"},
        ]
        custom_out = self.tmp_dir / "out_custom.xlsx"
        proc.export_to_original_excel(custom_records, custom_out)
        wb2 = openpyxl.load_workbook(custom_out)
        ws2 = wb2.active
        headers = [ws2.cell(1, c).value for c in range(1, ws2.max_column + 1)]
        self.assertEqual(headers, ["ma_hs", "ten_day_du", "diem_tb", "dan_toc"])
        self.assertEqual(ws2.cell(2, 2).value, "Lê Văn Tám")
        self.assertEqual(ws2.cell(3, 4).value, "Kinh")
        wb2.close()

    @patch("src.ocr_processor.extract_table_from_image")
    def test_process_ocr_batch_end_to_end(self, mock_extract):
        img1 = self.tmp_dir / "page_1.jpg"
        img2 = self.tmp_dir / "page_2.jpg"
        img1.write_bytes(b"fake_jpg_1")
        img2.write_bytes(b"fake_jpg_2")

        mock_extract.side_effect = [
            json.dumps([{"stt": 1, "ho_ten": "Lê Văn A", "ngay_sinh": "01/01/2008", "gioi_tinh": "Nam"}]),
            json.dumps([{"stt": 2, "ho_ten": "Phạm Thị B", "ngay_sinh": "02/02/2008", "gioi_tinh": "Nữ", "can_kiem_tra": True, "ly_do_nghi_ngo": "Chữ mờ"}]),
        ]

        proc = OCRProcessor(config_mgr=OCRConfigManager(self.tmp_dir / "cfg.json"))
        out_excel = self.tmp_dir / "truong_01_ketqua.xlsx"

        res = proc.process_ocr_batch(
            image_paths=[img2, img1],  # truyền ngược thứ tự để kiểm tra natural sort
            output_excel_path=out_excel,
            prompt_text="Test prompt",
            api_key="sk-test-key",
        )

        self.assertEqual(res["status"], "success")
        self.assertEqual(res["total_images"], 2)
        self.assertEqual(res["success_images"], 2)
        self.assertEqual(res["total_students"], 2)
        self.assertEqual(res["suspect_rows_count"], 1)
        self.assertTrue(out_excel.exists())
        self.assertFalse(out_excel.with_suffix(".json").exists())


if __name__ == "__main__":
    unittest.main()
