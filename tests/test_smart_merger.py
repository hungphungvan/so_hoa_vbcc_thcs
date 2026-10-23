from __future__ import annotations

import unittest
import tempfile
import shutil
from pathlib import Path
import openpyxl

from src.smart_merger import (
    TARGET_COLUMNS,
    compute_fuzzy_similarity,
    compute_headers_signature,
    auto_suggest_mapping,
    parse_column_expression,
    MappingProfileManager,
    SmartMerger
)


class TestSmartMerger(unittest.TestCase):
    def setUp(self):
        self.tmp_dir = Path(tempfile.mkdtemp())
        self.profile_file = self.tmp_dir / "test_profiles.json"
        self.profile_mgr = MappingProfileManager(self.profile_file)

    def tearDown(self):
        shutil.rmtree(self.tmp_dir, ignore_errors=True)

    def test_target_columns_definition(self):
        """Kiểm tra danh sách 28 cột chuẩn."""
        self.assertEqual(len(TARGET_COLUMNS), 28)
        self.assertEqual(TARGET_COLUMNS[0]["name"], "STT")
        self.assertEqual(TARGET_COLUMNS[4]["name"], "Họ, chữ đệm và tên")
        self.assertEqual(TARGET_COLUMNS[6]["name"], "Số định danh cá nhân")
        self.assertEqual(TARGET_COLUMNS[27]["name"], "Tên đơn vị")

    def test_fuzzy_similarity(self):
        """Kiểm tra khả năng so khớp mờ trực tiếp giữa tiêu đề chuẩn 28 cột và cột nguồn."""
        score_so_hieu = compute_fuzzy_similarity("Số hiệu bằng", "Số hiệu")
        self.assertGreaterEqual(score_so_hieu, 0.70)

        score_ngay_sinh = compute_fuzzy_similarity("Ngày, tháng, năm sinh", "Ngày sinh")
        self.assertGreaterEqual(score_ngay_sinh, 0.65)

        score_truong = compute_fuzzy_similarity("Tên trường đã học", "Tên trường")
        self.assertGreaterEqual(score_truong, 0.70)

        score_noi_sinh = compute_fuzzy_similarity("Nơi sinh", "Nơi sinh (Huyện - Tỉnh)")
        self.assertGreaterEqual(score_noi_sinh, 0.80)

        score_gioi_tinh = compute_fuzzy_similarity("Giới tính", "Giới tính")
        self.assertEqual(score_gioi_tinh, 1.0)

    def test_auto_suggest_mapping(self):
        """Kiểm tra thuật toán tự động gợi ý ánh xạ."""
        src_headers = ["STT", "Họ và tên", "Ngày sinh", "Nơi sinh", "Số định danh", "Niên khóa", "Số hiệu", "Số vào sổ", None]
        mapping = auto_suggest_mapping(src_headers)

        self.assertEqual(mapping[1], 0)  # STT
        self.assertEqual(mapping[5], 1)  # Họ, chữ đệm và tên <- Họ và tên
        self.assertEqual(mapping[9], 2)  # Ngày sinh
        self.assertEqual(mapping[13], 3) # Nơi sinh
        self.assertEqual(mapping[7], 4)  # Số định danh cá nhân <- Số định danh
        self.assertEqual(mapping[15], 5) # Niên khóa
        self.assertEqual(mapping[2], 6)  # Số hiệu bằng <- Số hiệu
        self.assertEqual(mapping[3], 7)  # Số vào sổ gốc <- Số vào sổ

        # Cột null ở vị trí 8 và các cột không có trong nguồn phải là None (Để trống)
        self.assertIsNone(mapping[4])   # Tên văn bằng
        self.assertIsNone(mapping[8])   # Hộ chiếu
        self.assertIsNone(mapping[12])  # Quốc tịch
        self.assertIsNone(mapping[16])  # Điểm thi

    def test_profile_manager(self):
        """Kiểm tra học từ mới, lấy danh sách từ đã học và xóa từ khỏi kho tri thức tập trung."""
        headers = ["Mã STT", "Họ tên mới", "Ngày sinh mới"]

        # Học từ mới vào kho tri thức tập trung
        self.profile_mgr.learn_mapping(
            mapping={1: 0, 5: 1, 9: 2},
            raw_headers=headers
        )
        terms_5 = self.profile_mgr.get_learned_terms(5)
        self.assertIn("Họ tên mới", terms_5)

        # Xóa từ đã học
        deleted = self.profile_mgr.delete_learned_term(5, len(terms_5) - 1)
        self.assertTrue(deleted)
        self.assertNotIn("Họ tên mới", self.profile_mgr.get_learned_terms(5))

    def test_smart_merger_end_to_end(self):
        """Kiểm tra quy trình Ghép & Chuẩn hóa thực tế từ file Excel nguồn với tiêu đề mới."""
        # Tạo file nguồn 1 với tiêu đề chưa có trong từ điển mẫu để kiểm tra kích hoạt callback học hỏi
        file1 = self.tmp_dir / "src_file_1.xlsx"
        wb1 = openpyxl.Workbook()
        ws1 = wb1.active
        ws1.title = "Lop9A"
        ws1.append(["DANH SÁCH TỐT NGHIỆP TRƯỜNG A"])
        ws1.append(["STT_Mới", "Tên_Đầy_Đủ_Mới", "Sinh_Nhật_Mới", "Phái_Mới", "Mã_Định_Danh_Mới", "Quê_Hương_Mới", "Mã_Số_Hiệu_Mới", "Số_Sổ_Gốc_Mới"])
        ws1.append([1, "nguyễn văn an", "15/05/2005", "Nam", "001205001234", "Hà Nội", "THCS001", "01/2020"])
        ws1.append([2, "TRẦN THỊ BÌNH", "20/10/2005", "Nữ", "001205005678", "Hải Phòng", "THCS002", "02/2020"])
        ws1.append(["Tổng số học sinh: 2", None, None, None, None, None, None, None])
        ws1.append(["Người lập biểu", None, None, None, None, None, None, "Hiệu trưởng"])
        wb1.save(file1)
        wb1.close()

        # Tạo file nguồn 2 (Cùng cấu trúc cột với file 1 để kiểm tra tự động nhận diện Profile không hỏi lại)
        file2 = self.tmp_dir / "src_file_2.xlsx"
        wb2 = openpyxl.Workbook()
        ws2 = wb2.active
        ws2.title = "Lop9B"
        ws2.append(["DANH SÁCH TỐT NGHIỆP TRƯỜNG B"])
        ws2.append(["STT_Mới", "Tên_Đầy_Đủ_Mới", "Sinh_Nhật_Mới", "Phái_Mới", "Mã_Định_Danh_Mới", "Quê_Hương_Mới", "Mã_Số_Hiệu_Mới", "Số_Sổ_Gốc_Mới"])
        ws2.append([1, "LÊ VĂN CƯỜNG", "01/01/2005", "Nam", "001205009999", "Đà Nẵng", "THCS003", "03/2020"])
        wb2.save(file2)
        wb2.close()

        output_file = self.tmp_dir / "Tong_Hop_Chuan_Hoa_28_Cot.xlsx"

        dialog_calls = []

        def mock_mapping_callback(payload):
            dialog_calls.append(payload["file_name"])
            # Ánh xạ thủ công cho các cột mới lạ
            mapping = {1: 0, 5: 1, 9: 2, 10: 3, 7: 4, 13: 5, 2: 6, 3: 7}
            return mapping, "Tri thức mới", False

        merger = SmartMerger(profile_mgr=self.profile_mgr)
        res = merger.process_merge_and_standardize(
            source_files=[file1, file2],
            output_file_path=output_file,
            mapping_callback=mock_mapping_callback,
            renumber_stt=True,
            fill_defaults=True,
            merge_all_sheets=True
        )

        self.assertEqual(res["status"], "success")
        self.assertEqual(res["total_rows"], 3)  # 2 từ file 1 + 1 từ file 2 (bỏ dòng tổng kết & footer)

        # QUAN TRỌNG: Callback chỉ được gọi 1 lần cho file 1. File 2 tự động áp dụng tri thức đã học mà KHÔNG gọi callback!
        self.assertEqual(len(dialog_calls), 1)
        self.assertEqual(dialog_calls[0], file1.name)

        # Kiểm tra nội dung file đích đã ghép chuẩn hóa
        self.assertTrue(output_file.exists())
        wb_res = openpyxl.load_workbook(output_file, data_only=True)
        ws_res = wb_res.active

        self.assertEqual(ws_res.max_row, 5) # Dòng 1 banner, Dòng 2 tiêu đề, Dòng 3,4,5 là 3 học sinh

        # Kiểm tra học sinh 1 (nguyen van an -> Nguyễn Văn An)
        r3 = [ws_res.cell(3, c).value for c in range(1, 29)]
        self.assertEqual(r3[0], 1) # STT đánh lại
        self.assertEqual(r3[1], "THCS001") # Số hiệu
        self.assertEqual(r3[2], "01/2020") # Số vào sổ
        self.assertEqual(r3[3], "Bằng tốt nghiệp trung học cơ sở") # Điền mặc định vì để trống
        self.assertEqual(r3[4], "Nguyễn Văn An") # Title case
        self.assertEqual(r3[6], "001205001234") # CCCD
        self.assertEqual(r3[8], "15/05/2005") # Ngày sinh
        self.assertEqual(r3[9], "Nam") # Giới tính
        self.assertEqual(r3[11], "Việt Nam") # Quốc tịch điền mặc định
        self.assertEqual(r3[12], "Hà Nội") # Nơi sinh
        self.assertEqual(r3[24], "Đã số hoá đầy đủ, chính xác") # Cột 25 điền mặc định từ TARGET_COLUMNS

        # Kiểm tra học sinh 3 từ file 2 (LE VAN CUONG -> Lê Văn Cường, STT = 3)
        r5 = [ws_res.cell(5, c).value for c in range(1, 29)]
        self.assertEqual(r5[0], 3) # STT liên tục toàn bảng
        self.assertEqual(r5[4], "Lê Văn Cường") # Title case chuẩn
        self.assertEqual(r5[6], "001205009999") # CCCD bảo toàn số 0 đầu
        self.assertEqual(r5[3], "Bằng tốt nghiệp trung học cơ sở")
        self.assertEqual(r5[11], "Việt Nam")
        self.assertEqual(r5[24], "Đã số hoá đầy đủ, chính xác")

        wb_res.close()

    def test_smart_merger_without_stt_column(self):
        """Kiểm tra trích xuất học sinh từ sheet hoàn toàn KHÔNG CÓ cột STT (chỉ cần có Họ và tên)."""
        file_no_stt = self.tmp_dir / "src_no_stt.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "KhongCoSTT"
        ws.append(["DANH SÁCH KHÔNG CỘT STT"])
        # Tiêu đề không hề có cột STT
        ws.append(["Họ và tên", "Ngày sinh", "Giới tính", "Nơi sinh"])
        ws.append(["Phạm Nhật Minh", "12/03/2005", "Nam", "Hà Nội"])
        ws.append(["Hoàng Thu Trang", "25/11/2005", "Nữ", "Hải Dương"])
        ws.append(["Vũ Đức Đam", "09/08/2005", "Nam", "Quảng Ninh"])
        ws.append(["Tổng số: 3 học sinh", None, None, None]) # Dòng footer
        wb.save(file_no_stt)
        wb.close()

        output_file = self.tmp_dir / "Out_No_STT.xlsx"

        merger = SmartMerger(profile_mgr=self.profile_mgr)
        res = merger.process_merge_and_standardize(
            source_files=[file_no_stt],
            output_file_path=output_file,
            mapping_callback=None, # Tự động nhận diện
            renumber_stt=True,
            fill_defaults=True,
            merge_all_sheets=True
        )

        self.assertEqual(res["status"], "success")
        self.assertEqual(res["total_rows"], 3) # Phải trích xuất đủ cả 3 học sinh

        wb_res = openpyxl.load_workbook(output_file, data_only=True)
        ws_res = wb_res.active

        # Kiểm tra 3 học sinh đều có STT tự đánh 1, 2, 3
        stt_list = [ws_res.cell(r, 1).value for r in [3, 4, 5]]
        name_list = [ws_res.cell(r, 5).value for r in [3, 4, 5]]
        self.assertEqual(stt_list, [1, 2, 3])
        self.assertEqual(name_list, ["Phạm Nhật Minh", "Hoàng Thu Trang", "Vũ Đức Đam"])

        wb_res.close()

    def test_smart_merger_with_custom_default_values(self):
        """Kiểm tra tính năng nhập giá trị mặc định cho cột để trống khi người dùng cấu hình qua callback."""
        file1 = self.tmp_dir / "src_custom_defaults_1.xlsx"
        wb1 = openpyxl.Workbook()
        ws1 = wb1.active
        ws1.title = "Lop9"
        # Dùng tiêu đề mới lạ để kích hoạt callback
        ws1.append(["Họ và tên thí sinh mới", "Ngày sinh thí sinh mới", "Giới tính thí sinh mới"])
        ws1.append(["Đặng Thái Sơn", "02/07/2005", "Nam"])
        ws1.append(["Ngô Bảo Châu", "28/06/2005", "Nam"])
        wb1.save(file1)
        wb1.close()

        file2 = self.tmp_dir / "src_custom_defaults_2.xlsx"
        wb2 = openpyxl.Workbook()
        ws2 = wb2.active
        ws2.title = "Lop9"
        ws2.append(["Họ và tên thí sinh mới", "Ngày sinh thí sinh mới", "Giới tính thí sinh mới"])
        ws2.append(["Lưu Lệ Hằng", "04/10/2005", "Nữ"])
        wb2.save(file2)
        wb2.close()

        output_file = self.tmp_dir / "Out_Custom_Defaults.xlsx"

        call_count = 0
        def mock_callback(payload):
            nonlocal call_count
            call_count += 1
            mapping = {
                5: 0, # Họ và tên
                9: 1, # Ngày sinh
                10: 2 # Giới tính
            }
            # Người dùng tự nhập giá trị mặc định cho đợt xử lý này
            custom_defs = {
                14: "THCS Chu Văn An",
                15: "2019-2023",
                18: "Hà Nội",
                20: "Phòng GD&ĐT Tây Hồ"
            }
            return mapping, "Tri thức mới", custom_defs, False

        merger = SmartMerger(profile_mgr=self.profile_mgr)
        res = merger.process_merge_and_standardize(
            source_files=[file1, file2],
            output_file_path=output_file,
            mapping_callback=mock_callback,
            renumber_stt=True,
            fill_defaults=True,
            merge_all_sheets=True
        )

        self.assertEqual(res["status"], "success")
        self.assertEqual(res["total_rows"], 3)
        self.assertEqual(call_count, 1) # File 2 tự động áp dụng tri thức đã học không cần hỏi lại

        wb_res = openpyxl.load_workbook(output_file, data_only=True)
        ws_res = wb_res.active

        # Kiểm tra học sinh 1 (Đặng Thái Sơn) có giá trị mặc định của hệ thống
        self.assertEqual(ws_res.cell(3, 1).value, 1) # STT
        self.assertEqual(ws_res.cell(3, 4).value, "Bằng tốt nghiệp trung học cơ sở") # Mặc định hệ thống
        self.assertEqual(ws_res.cell(3, 5).value, "Đặng Thái Sơn") # Họ tên
        self.assertEqual(ws_res.cell(3, 12).value, "Việt Nam") # Mặc định hệ thống
        self.assertEqual(ws_res.cell(3, 25).value, "Đã số hoá đầy đủ, chính xác") # Mặc định hệ thống từ TARGET_COLUMNS

        wb_res.close()

    def test_parse_column_expression(self):
        """Kiểm tra phân tích các cú pháp biểu thức ghép cột (Phương án 3)."""
        formatted = ["Cột A: STT", "Cột B: Họ và chữ đệm", "Cột C: Tên", "Cột D: Ngày sinh", "Cột E: Quê quán"]
        
        # 1. Để trống
        self.assertIsNone(parse_column_expression("", formatted))
        self.assertIsNone(parse_column_expression("-- Để trống --", formatted))
        self.assertIsNone(parse_column_expression("--", formatted))

        # 2. Cột đơn khớp chính xác
        self.assertEqual(parse_column_expression("Cột B: Họ và chữ đệm", formatted), 1)
        self.assertEqual(parse_column_expression("Cột D: Ngày sinh", formatted), 3)

        # 3. Biểu thức ghép bằng chữ cái cột: B + C, b, c, B & C
        self.assertEqual(parse_column_expression("B + C", formatted), [1, 2])
        self.assertEqual(parse_column_expression("b, c", formatted), [1, 2])
        self.assertEqual(parse_column_expression("B & C", formatted), [1, 2])
        self.assertEqual(parse_column_expression("cột b + cột c", formatted), [1, 2])

        # 4. Biểu thức ghép bằng nhãn đầy đủ
        self.assertEqual(parse_column_expression("Cột B: Họ và chữ đệm + Cột C: Tên", formatted), [1, 2])

        # 5. Biểu thức ghép bằng số thứ tự 1-indexed: 2 + 3
        self.assertEqual(parse_column_expression("2 + 3", formatted), [1, 2])

        # 6. Biểu thức ghép bằng từ khóa tương đối
        self.assertEqual(parse_column_expression("Họ và chữ đệm + Tên", formatted), [1, 2])

    def test_auto_suggest_mapping_split_name(self):
        """Kiểm tra thuật toán tự động nhận diện tệp nguồn có cột Họ đệm và Tên tách rời."""
        # Trường hợp 1: Họ và chữ đệm + Tên
        headers1 = ["STT", "Họ và chữ đệm", "Tên", "Ngày sinh", "Giới tính"]
        m1 = auto_suggest_mapping(headers1)
        self.assertEqual(m1[1], 0)        # STT
        self.assertEqual(m1[5], [1, 2])   # Họ, chữ đệm và tên <- ghép cột 1 và 2
        self.assertEqual(m1[9], 3)        # Ngày sinh
        self.assertEqual(m1[10], 4)       # Giới tính

        # Trường hợp 2: Họ đệm + Tên
        headers2 = ["STT", "Họ đệm", "Tên", "Ngày sinh"]
        m2 = auto_suggest_mapping(headers2)
        self.assertEqual(m2[5], [1, 2])

        # Trường hợp 3: Họ + Tên (không có STT)
        headers3 = ["Họ", "Tên", "Ngày sinh"]
        m3 = auto_suggest_mapping(headers3)
        self.assertEqual(m3[5], [0, 1])

        # Trường hợp 4: Có cột Họ và tên gộp thì không tách
        headers4 = ["STT", "Họ và tên", "Ngày sinh"]
        m4 = auto_suggest_mapping(headers4)
        self.assertEqual(m4[5], 1)

    def test_smart_merger_with_split_name_columns(self):
        """Kiểm tra quy trình Ghép & Chuẩn hóa thực tế khi cột Họ và Tên bị tách thành 2 cột."""
        file_split = self.tmp_dir / "src_split_name.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Lop9_TachCot"
        ws.append(["DANH SÁCH TỐT NGHIỆP CÓ HỌ VÀ TÊN TÁCH RỜI"])
        ws.append(["STT", "Họ và chữ đệm", "Tên", "Ngày sinh", "Nơi sinh"])
        ws.append([1, "nguyễn văn", "an", "15/05/2005", "Hà Nội"])
        ws.append([2, "TRẦN THỊ BÍCH", "NGỌC", "20/10/2005", "Hải Phòng"])
        ws.append([3, "  LÊ   VĂN ", " CƯỜNG ", "01/01/2005", "Đà Nẵng"])
        ws.append(["Tổng số học sinh: 3", None, None, None, None]) # Dòng footer
        wb.save(file_split)
        wb.close()

        output_file = self.tmp_dir / "Out_Split_Name.xlsx"

        merger = SmartMerger(profile_mgr=self.profile_mgr)
        res = merger.process_merge_and_standardize(
            source_files=[file_split],
            output_file_path=output_file,
            mapping_callback=None, # Tự động phát hiện mapping[5] = [1, 2]
            renumber_stt=True,
            fill_defaults=True,
            merge_all_sheets=True
        )

        self.assertEqual(res["status"], "success")
        self.assertEqual(res["total_rows"], 3)

        wb_res = openpyxl.load_workbook(output_file, data_only=True)
        ws_res = wb_res.active

        # Kiểm tra học sinh 1: "nguyễn văn" + "an" -> "Nguyễn Văn An"
        self.assertEqual(ws_res.cell(3, 1).value, 1)
        self.assertEqual(ws_res.cell(3, 5).value, "Nguyễn Văn An")
        self.assertEqual(ws_res.cell(3, 9).value, "15/05/2005")
        self.assertEqual(ws_res.cell(3, 13).value, "Hà Nội")

        # Kiểm tra học sinh 2: "TRẦN THỊ BÍCH" + "NGỌC" -> "Trần Thị Bích Ngọc"
        self.assertEqual(ws_res.cell(4, 1).value, 2)
        self.assertEqual(ws_res.cell(4, 5).value, "Trần Thị Bích Ngọc")

        # Kiểm tra học sinh 3: khoảng trắng thừa được dọn sạch -> "Lê Văn Cường"
        self.assertEqual(ws_res.cell(5, 1).value, 3)
        self.assertEqual(ws_res.cell(5, 5).value, "Lê Văn Cường")

        wb_res.close()

    def test_profile_with_combined_columns(self):
        """Kiểm tra học và nạp lại cấu hình ghép nhiều cột trong kho tri thức tập trung."""
        headers = ["STT_A", "Họ và chữ đệm A", "Tên A", "Ngày sinh A"]

        self.profile_mgr.learn_mapping(
            mapping={1: 0, 5: [1, 2], 9: 3},
            raw_headers=headers
        )
        terms_5 = self.profile_mgr.get_learned_terms(5)
        self.assertIn(["Họ và chữ đệm A", "Tên A"], terms_5)

    def test_find_exact_matching_profile_reordered_columns(self):
        """Kiểm tra nhận diện 100% khớp tên cột dù các cột bị đảo thứ tự."""
        prof_headers = ["STT_X", "Họ và tên X", "Ngày sinh X", "Nơi sinh X", "Số hiệu bằng X"]
        self.profile_mgr.learn_mapping(
            mapping={1: 0, 5: 1, 9: 2, 13: 3, 2: 4},
            raw_headers=prof_headers
        )

        # File mới bị đảo lộn thứ tự các cột
        reordered_headers = ["Số hiệu bằng X", "Nơi sinh X", "STT_X", "Họ và tên X", "Ngày sinh X"]
        matched_prof, resolved_map = self.profile_mgr.find_exact_matching_profile(reordered_headers)

        self.assertIsNotNone(matched_prof)
        # Kiểm tra resolved_map lấy đúng chỉ số cột mới:
        self.assertEqual(resolved_map[2], 0)   # "Số hiệu bằng X" ở vị trí 0
        self.assertEqual(resolved_map[13], 1)  # "Nơi sinh X" ở vị trí 1
        self.assertEqual(resolved_map[1], 2)   # "STT_X" ở vị trí 2
        self.assertEqual(resolved_map[5], 3)   # "Họ và tên X" ở vị trí 3
        self.assertEqual(resolved_map[9], 4)   # "Ngày sinh X" ở vị trí 4

    def test_find_exact_profile_refuses_when_missing_mapped_column(self):
        """Kiểm tra từ chối tự động áp dụng 100% khi có cột lạ chưa từng học."""
        # File có cột lạ chưa từng biết "Cột Lạ Hoàn Toàn 999"
        unknown_headers = ["STT", "Họ và tên", "Ngày sinh", "Cột Lạ Hoàn Toàn 999"]
        matched_prof, resolved_map = self.profile_mgr.find_exact_matching_profile(unknown_headers)
        self.assertIsNone(matched_prof)
        self.assertIsNone(resolved_map)

    def test_two_tier_auto_suggest_mapping(self):
        """Kiểm tra thuật toán gợi ý 2 tầng: Tầng 1 khớp mờ kho tri thức, Tầng 2 theo 28 cột chuẩn."""
        from src.smart_merger import auto_suggest_mapping_with_details

        # Học các biến thể tiêu đề đặc thù vào kho tri thức
        prof_headers = ["Mã Số Học Sinh", "Họ Tên Thí Sinh", "Ngày Tháng Năm Đẻ", "Nơi Chôn Rau Cắt Rốn"]
        self.profile_mgr.learn_mapping(
            mapping={1: 0, 5: 1, 9: 2, 13: 3},
            raw_headers=prof_headers
        )

        # File mới có tiêu đề khớp mờ với tri thức đã học (Tầng 1) và có cột chuẩn 28 (Tầng 2: Số định danh cá nhân)
        new_headers = ["Họ Tên Thí Sinh", "Ngày Tháng Năm Đẻ", "Nơi Chôn Rau Cắt Rốn", "Số định danh cá nhân", ""]
        mapping, defaults, matched_pname = auto_suggest_mapping_with_details(new_headers, self.profile_mgr)

        # Cột 5 (Họ tên) -> index 0 (khớp Tầng 1)
        self.assertEqual(mapping.get(5), 0)
        # Cột 9 (Ngày sinh) -> index 1 (khớp Tầng 1)
        self.assertEqual(mapping.get(9), 1)
        # Cột 13 (Nơi sinh) -> index 2 (khớp Tầng 1)
        self.assertEqual(mapping.get(13), 2)
        # Cột 7 (Số định danh) -> index 3 (khớp Tầng 2 chuẩn 28 cột)
        self.assertEqual(mapping.get(7), 3)
        # Cột rỗng ở vị trí 4 không bao giờ được gán vào bất kỳ cột nào
        for t_idx, s_val in mapping.items():
            self.assertNotEqual(s_val, 4)

    def test_fallback_columns_when_no_dedicated_exam_board_column(self):
        """Kiểm tra Cột 17 (Hội đồng thi) tự động dự phòng lấy từ Cột 14 (Tên trường) khi không có cột riêng."""
        from src.smart_merger import auto_suggest_mapping_with_details

        headers = ["STT", "Họ và tên", "Ngày sinh", "Nơi học"]
        mapping, defaults, _ = auto_suggest_mapping_with_details(headers, self.profile_mgr)

        # Cột 14 nhận cột "Nơi học" (index 3)
        self.assertEqual(mapping.get(14), 3)
        # Cột 17 tự động nhận theo Cột 14 nhờ quy tắc fallback
        self.assertEqual(mapping.get(17), 3)

        # Kiểm tra learn_mapping không nhét "Nơi học" vào từ điển riêng của Cột 17
        self.profile_mgr.learn_mapping(mapping, headers)
        terms_17 = self.profile_mgr.get_learned_terms(17)
        self.assertNotIn("Nơi học", terms_17)

    def test_fallback_columns_when_dedicated_exam_board_column_exists(self):
        """Kiểm tra Cột 17 ưu tiên nhận cột Hội đồng thi riêng biệt khi file có cả hai cột."""
        from src.smart_merger import auto_suggest_mapping_with_details

        headers = ["STT", "Họ và tên", "Ngày sinh", "Tên trường", "Hội đồng coi thi"]
        mapping, defaults, _ = auto_suggest_mapping_with_details(headers, self.profile_mgr)

        # Cột 14 nhận cột "Tên trường" (index 3)
        self.assertEqual(mapping.get(14), 3)
    def test_inverted_profiles_set_matching(self):
        """Kiểm tra toàn diện cơ chế so khớp tập hợp: Trùng khớp 100%, Tập con, Tập mẹ."""
        # Lưu Profile gốc A: 5 cột
        prof_headers = ["STT", "Họ đệm", "Tên", "Ngày sinh", "Nơi sinh"]
        self.profile_mgr.save_profile(
            name="Mẫu Trường A",
            source_headers=prof_headers,
            mapping={
                "stt": 1,
                "ho dem": [5, 0],
                "ten": [5, 1],
                "ngay sinh": 9,
                "noi sinh": 13
            }
        )

        # 1. Trùng khớp 100% (Exact Match)
        match_type, p, res_map, _ = self.profile_mgr.find_matching_profile(
            ["Nơi sinh", "Ngày sinh", "Tên", "Họ đệm", "STT"]
        )
        self.assertEqual(match_type, "exact")
        self.assertEqual(p["name"], "Mẫu Trường A")
        self.assertEqual(res_map[1], 4)   # STT ở vị trí 4
        self.assertEqual(res_map[5], [3, 2]) # Họ đệm ở 3, Tên ở 2

        # 2. Tập con (Subset Match) - có đủ Họ tên
        sub_headers = ["STT", "Họ đệm", "Tên", "Ngày sinh"]
        match_type_sub, p_sub, res_map_sub, _ = self.profile_mgr.find_matching_profile(sub_headers)
        self.assertEqual(match_type_sub, "subset")
        self.assertEqual(p_sub["name"], "Mẫu Trường A")
        self.assertEqual(res_map_sub[1], 0)
        self.assertEqual(res_map_sub[5], [1, 2])
        self.assertIsNone(res_map_sub[13]) # Nơi sinh bị thiếu trong tập con -> None

        # 3. Tập con nhưng thiếu Họ tên -> Không được tự động áp dụng (trả về None)
        no_name_sub = ["STT", "Ngày sinh", "Nơi sinh"]
        match_type_no_name, _, _, _ = self.profile_mgr.find_matching_profile(no_name_sub)
        self.assertIsNone(match_type_no_name)

        # 4. Tập mẹ (Superset Match) - Cột thừa KHÔNG có dữ liệu -> Tự động bỏ qua
        sup_headers = ["STT", "Họ đệm", "Tên", "Ngày sinh", "Nơi sinh", "Cột Thừa 1", "Cột Thừa 2"]
        sample_rows_empty_extras = [
            [1, "Nguyễn Văn", "An", "01/01/2005", "Hà Nội", None, ""],
            [2, "Trần Thị", "Bình", "02/02/2005", "Hải Phòng", "", None]
        ]
        match_type_sup_auto, p_sup, res_map_sup, _ = self.profile_mgr.find_matching_profile(
            sup_headers, sample_rows=sample_rows_empty_extras
        )
        self.assertEqual(match_type_sup_auto, "superset_auto")
        self.assertEqual(p_sup["name"], "Mẫu Trường A")
        self.assertEqual(res_map_sup[1], 0)

        # 5. Tập mẹ (Superset Match) - Cột thừa CÓ dữ liệu -> Yêu cầu người dùng duyệt (superset_ask)
        sample_rows_with_extras = [
            [1, "Nguyễn Văn", "An", "01/01/2005", "Hà Nội", "Có dữ liệu điểm", "8.5"],
            [2, "Trần Thị", "Bình", "02/02/2005", "Hải Phòng", "Có dữ liệu điểm", "9.0"]
        ]
        match_type_sup_ask, p_ask, _, extra_cols = self.profile_mgr.find_matching_profile(
            sup_headers, sample_rows=sample_rows_with_extras
        )
        self.assertEqual(match_type_sup_ask, "superset_ask")
        self.assertEqual(p_ask["name"], "Mẫu Trường A")
        self.assertIn("Cột Thừa 1", extra_cols)
        self.assertIn("Cột Thừa 2", extra_cols)

    def test_batch_delete_profiles(self):
        """Kiểm tra xóa nhiều Profile cùng lúc và xóa tất cả Profile."""
        p1 = self.profile_mgr.save_profile("Prof 1", ["STT", "Họ tên"], {"stt": 1, "ho ten": 5})
        p2 = self.profile_mgr.save_profile("Prof 2", ["STT", "Tên"], {"stt": 1, "ten": 5})
        p3 = self.profile_mgr.save_profile("Prof 3", ["STT", "Họ"], {"stt": 1, "ho": 5})

        self.assertGreaterEqual(len(self.profile_mgr.get_all_profiles()), 3)

        # Xóa 2 profile cùng lúc
        deleted_cnt = self.profile_mgr.delete_profiles([p1["id"], p2["id"]])
        self.assertEqual(deleted_cnt, 2)
        self.assertIsNone(self.profile_mgr.get_profile_by_id(p1["id"]))
        self.assertIsNone(self.profile_mgr.get_profile_by_id(p2["id"]))
        self.assertIsNotNone(self.profile_mgr.get_profile_by_id(p3["id"]))

        # Xóa tất cả
        total_deleted = self.profile_mgr.delete_all_profiles()
        self.assertGreaterEqual(total_deleted, 1)
        self.assertEqual(len(self.profile_mgr.get_all_profiles()), 0)

    def test_immediate_cancel_all(self):
        """Kiểm tra nút 'Dừng toàn bộ' dừng ngay lập tức, không tiếp tục các file kế tiếp."""
        file1 = self.tmp_dir / "file_cancel_1.xlsx"
        wb1 = openpyxl.Workbook()
        ws1 = wb1.active
        ws1.title = "Sheet1"
        ws1.append(["STT_La", "HoTen_La"])
        ws1.append([1, "Nguyen Van A"])
        wb1.save(file1)
        wb1.close()

        file2 = self.tmp_dir / "file_cancel_2.xlsx"
        wb2 = openpyxl.Workbook()
        ws2 = wb2.active
        ws2.title = "Sheet1"
        ws2.append(["STT_La", "HoTen_La"])
        ws2.append([1, "Tran Van B"])
        wb2.save(file2)
        wb2.close()

        out_f = self.tmp_dir / "Out_Cancel_Test.xlsx"
        callback_history = []

        def cancel_callback(payload):
            callback_history.append(payload["file_name"])
            # Người dùng chọn "Dừng toàn bộ"
            return ("cancel", None, None)

        merger = SmartMerger(profile_mgr=self.profile_mgr)
        res = merger.process_merge_and_standardize(
            source_files=[file1, file2],
            output_file_path=out_f,
            mapping_callback=cancel_callback,
            merge_all_sheets=True
        )

        self.assertEqual(res["status"], "cancelled")
        # Phải dừng ngay tại file 1, KHÔNG được mở file 2
        self.assertEqual(len(callback_history), 1)
        self.assertEqual(callback_history[0], file1.name)

    def test_code_lookup_mapping_in_smart_merger(self):
        """Kiểm tra tự động tra cứu mã số (Giới tính, Dân tộc, Quốc tịch, Nơi sinh) trong Sub-tab 3."""
        file_codes = self.tmp_dir / "src_codes.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "DS_MaSo"
        ws.append(["STT", "Họ và tên", "Ngày sinh", "Giới tính", "Dân tộc", "Quốc tịch", "Nơi sinh"])
        # Học sinh 1: nhập mã số nguyên / chuỗi có số 0 ở đầu
        # Gioi_Tinh: 1 -> Nam; Dan_Toc: 29 -> Kinh; Quoc_Tich: 83 -> Việt Nam; Noi_Sinh: 3 -> Thành phố Hải Phòng
        ws.append([1, "nguyễn văn an", "15/05/2008", 1, 29, 83, 3])
        # Học sinh 2: nhập chuỗi có số 0 đứng đầu ("03" -> Nữ, "01" -> Mạ, "83" -> Việt Nam, "01" -> Tỉnh Bắc Ninh)
        ws.append([2, "trần thị bình", "20/10/2008", "03", "01", "83", "01"])
        wb.save(file_codes)
        wb.close()

        out_f = self.tmp_dir / "Out_Codes.xlsx"
        merger = SmartMerger(profile_mgr=self.profile_mgr)
        res = merger.process_merge_and_standardize(
            source_files=[file_codes],
            output_file_path=out_f,
            mapping_callback=None,
            renumber_stt=True,
            fill_defaults=True,
            merge_all_sheets=True
        )

        self.assertEqual(res["status"], "success")
        self.assertEqual(res["total_rows"], 2)

        wb_res = openpyxl.load_workbook(out_f, data_only=True)
        ws_res = wb_res.active

        # Kiểm tra học sinh 1 (Dòng 3)
        self.assertEqual(ws_res.cell(3, 5).value, "Nguyễn Văn An")
        self.assertEqual(ws_res.cell(3, 10).value, "Nam")
        self.assertEqual(ws_res.cell(3, 11).value, "Kinh")
        self.assertEqual(ws_res.cell(3, 12).value, "Việt Nam")
        self.assertEqual(ws_res.cell(3, 13).value, "Thành phố Hải Phòng")
        self.assertEqual(ws_res.cell(3, 25).value, "Đã số hoá đầy đủ, chính xác")

        # Kiểm tra học sinh 2 (Dòng 4)
        self.assertEqual(ws_res.cell(4, 5).value, "Trần Thị Bình")
        self.assertEqual(ws_res.cell(4, 10).value, "Nữ")
        self.assertEqual(ws_res.cell(4, 11).value, "Mạ")
        self.assertEqual(ws_res.cell(4, 12).value, "Việt Nam")
        self.assertEqual(ws_res.cell(4, 13).value, "Tỉnh Bắc Ninh")

        wb_res.close()


if __name__ == '__main__':
    unittest.main()

