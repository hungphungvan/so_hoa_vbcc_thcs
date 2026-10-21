import sys
import unittest
import tempfile
from pathlib import Path
import openpyxl

from src.custom_excel_manager import (
    is_valid_stt,
    is_valid_name,
    is_header_duplicate_row,
    detect_header_and_columns,
    ExcelSplitter,
    ExcelMerger
)

TEMP_DIR = Path(tempfile.gettempdir())

class TestModule3(unittest.TestCase):
    def test_stt_validation(self):
        self.assertTrue(is_valid_stt(1))
        self.assertTrue(is_valid_stt(2.0))
        self.assertTrue(is_valid_stt("3"))
        self.assertTrue(is_valid_stt("04"))
        self.assertTrue(is_valid_stt("15.0"))

        self.assertFalse(is_valid_stt(None))
        self.assertFalse(is_valid_stt(False))
        self.assertFalse(is_valid_stt(True))
        self.assertFalse(is_valid_stt("STT"))
        self.assertFalse(is_valid_stt("Số TT"))
        self.assertFalse(is_valid_stt(0))
        self.assertFalse(is_valid_stt(-5))
        self.assertFalse(is_valid_stt("abc"))
        self.assertFalse(is_valid_stt(""))

    def test_name_validation(self):
        self.assertTrue(is_valid_name("Nguyễn Văn An"))
        self.assertTrue(is_valid_name("Trần Thị Bích"))
        self.assertTrue(is_valid_name("Lê Cường"))
        self.assertTrue(is_valid_name("John Doe"))

        self.assertFalse(is_valid_name(None))
        self.assertFalse(is_valid_name(""))
        self.assertFalse(is_valid_name("   "))
        self.assertFalse(is_valid_name("Họ và tên"))
        self.assertFalse(is_valid_name("HỌ VÀ TÊN"))
        self.assertFalse(is_valid_name("Họ tên"))
        self.assertFalse(is_valid_name("123456"))
        self.assertFalse(is_valid_name("---"))

    def test_header_duplicate_detection(self):
        """Kiểm tra hàm loại trừ dòng tiêu đề lặp lại hoặc có chứa từ khóa tiêu đề."""
        header = ["STT", "Số hiệu bằng", "Họ và tên", "Ngày tháng năm sinh", "Trường học", "Đơn vị"]
        
        # 1. Dòng tiêu đề lặp lại nguyên bản
        self.assertTrue(is_header_duplicate_row(header, header_row_data=header, s_idx=0, n_idx=2))
        
        # 2. Dòng có STT nhưng cột tên lại ghi "Họ và tên" hoặc "Họ tên học sinh"
        row_fake_header1 = [1, "SH01", "Họ và tên", "Ngày sinh", "THCS Ba Đình", "Hà Nội"]
        self.assertTrue(is_header_duplicate_row(row_fake_header1, header_row_data=header, s_idx=0, n_idx=2))
        
        row_fake_header2 = [1, "SH02", "Họ tên học sinh", "2005", "Trường THCS", "Hà Nội"]
        self.assertTrue(is_header_duplicate_row(row_fake_header2, header_row_data=header, s_idx=0, n_idx=2))

        # 3. Dòng tiêu đề lặp lại có STT cột là 'STT'
        row_fake_header3 = ["STT", "Số hiệu", "Nguyễn Văn A", "10/02/2005", "THCS", "Hà Nội"]
        self.assertTrue(is_header_duplicate_row(row_fake_header3, header_row_data=header, s_idx=0, n_idx=2))

        # 4. Dòng học sinh thật sự -> Không được coi là header duplicate
        valid_student = [1, "SH03", "Nguyễn Văn An", "15/05/2005", "THCS Ba Đình", "Hà Nội"]
        self.assertFalse(is_header_duplicate_row(valid_student, header_row_data=header, s_idx=0, n_idx=2))

    def test_split_by_column_with_early_exit(self):
        """Kiểm tra tính năng tách file theo cột, lọc bỏ tiêu đề lặp lại và dừng khi có 15 dòng trống."""
        test_file = TEMP_DIR / "test_split_mixed_units.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Sheet1"

        # Dòng 1: Banner
        ws.append(["SỞ GIÁO DỤC VÀ ĐÀO TẠO HÀ NỘI"])
        # Dòng 2: Header (6 cột ví dụ: STT, Số hiệu, Họ và tên, Ngày sinh, Tên trường, Tên đơn vị)
        header_row = ["STT", "Số hiệu", "Họ và tên", "Ngày sinh", "Tên trường", "Tên đơn vị"]
        ws.append(header_row)

        # 3 học sinh Ba Đình
        ws.append([1, "SH01", "Nguyễn Văn An", "15/05/2005", "THCS Ba Đình", "Phòng GD&ĐT Ba Đình"])
        ws.append([2, "SH02", "Trần Thị Bích", "20/10/2005", "THCS Ba Đình", "Phòng GD&ĐT Ba Đình"])
        ws.append([3, "SH03", "Lê Văn Cường", "12/03/2005", "THCS Giảng Võ", "Phòng GD&ĐT Ba Đình"])

        # 1 dòng trống ngắt trang
        ws.append([])
        # 1 dòng tiêu đề lặp lại (Page break header)
        ws.append(header_row)

        # 2 học sinh Cầu Giấy
        ws.append([4, "SH04", "Phạm Văn Dũng", "10/01/2005", "THCS Nghĩa Tân", "Phòng GD&ĐT Cầu Giấy"])
        ws.append([5, "SH05", "Hoàng Thị Em", "05/08/2005", "THCS Dịch Vọng", "Phòng GD&ĐT Cầu Giấy"])

        # 1 dòng chân trang
        ws.append(["Tổng số học sinh: 5"])
        ws.append(["NGƯỜI LẬP BIỂU (Ký tên)"])

        # 20 dòng trống liên tiếp ở cuối file
        for _ in range(20):
            ws.append([])

        wb.save(test_file)
        wb.close()

        # 1. Kiểm tra preview
        preview_res = ExcelSplitter.preview_split(test_file, split_col_idx=5)
        self.assertEqual(preview_res["total_valid_rows"], 5)
        self.assertEqual(len(preview_res["groups"]), 2)
        self.assertEqual(preview_res["groups"]["Phòng GD&ĐT Ba Đình"], 3)
        self.assertEqual(preview_res["groups"]["Phòng GD&ĐT Cầu Giấy"], 2)

        # 2. Kiểm tra thực hiện tách file
        out_split_dir = TEMP_DIR / "test_out_split"
        split_res = ExcelSplitter.split_file(
            file_path=test_file,
            split_col_idx=5,  # Cột Tên đơn vị (index 5)
            output_dir=out_split_dir,
            renumber_stt=True
        )

        self.assertEqual(split_res["status"], "success")
        self.assertEqual(len(split_res["files"]), 2)

        # Kiểm tra file Ba Đình
        badinh_file = out_split_dir / "Phòng GD&ĐT Ba Đình.xlsx"
        self.assertTrue(badinh_file.exists())
        wb_bd = openpyxl.load_workbook(badinh_file)
        ws_bd = wb_bd.active
        # Dòng 1: Banner, Dòng 2: Header, Dòng 3, 4, 5: 3 học sinh
        self.assertEqual(ws_bd.max_row, 5)
        self.assertEqual(ws_bd.cell(3, 1).value, 1) # STT 1
        self.assertEqual(ws_bd.cell(4, 1).value, 2) # STT 2
        self.assertEqual(ws_bd.cell(5, 1).value, 3) # STT 3
        self.assertEqual(ws_bd.cell(3, 3).value, "Nguyễn Văn An")
        wb_bd.close()

        # Kiểm tra file Cầu Giấy
        caugiay_file = out_split_dir / "Phòng GD&ĐT Cầu Giấy.xlsx"
        self.assertTrue(caugiay_file.exists())
        wb_cg = openpyxl.load_workbook(caugiay_file)
        ws_cg = wb_cg.active
        self.assertEqual(ws_cg.max_row, 4)
        self.assertEqual(ws_cg.cell(3, 1).value, 1) # STT đánh lại từ 1
        self.assertEqual(ws_cg.cell(4, 1).value, 2) # STT đánh lại từ 2
        self.assertEqual(ws_cg.cell(3, 3).value, "Phạm Văn Dũng")
        wb_cg.close()

        print("--> Test ExcelSplitter PASSED successfully!")

    def test_merge_files(self):
        """Kiểm tra gộp 2 file đã tách lại thành 1 file tổng hợp."""
        # Chạy tách file trước để sinh ra các file con
        self.test_split_by_column_with_early_exit()

        badinh_file = TEMP_DIR / "test_out_split" / "Phòng GD&ĐT Ba Đình.xlsx"
        caugiay_file = TEMP_DIR / "test_out_split" / "Phòng GD&ĐT Cầu Giấy.xlsx"
        out_merged = TEMP_DIR / "test_out_merged" / "Tong_Hop_2_Don_Vi.xlsx"

        val_res = ExcelMerger.validate_headers([badinh_file, caugiay_file])
        self.assertTrue(val_res["compatible"])

        merge_res = ExcelMerger.merge_files(
            file_paths=[badinh_file, caugiay_file],
            output_file_path=out_merged,
            renumber_stt=True,
            add_source_col=True
        )

        self.assertEqual(merge_res["status"], "success")
        self.assertEqual(merge_res["total_rows"], 5)

        wb_m = openpyxl.load_workbook(out_merged)
        ws_m = wb_m.active
        # Dòng 1: Banner, Dòng 2: Header, Dòng 3..7: 5 học sinh
        self.assertEqual(ws_m.max_row, 7)
        for i in range(1, 6):
            self.assertEqual(ws_m.cell(2 + i, 1).value, i) # STT đánh lại 1..5

        # Cột thứ 7 (thêm mới): Tên file nguồn gốc
        self.assertEqual(ws_m.cell(3, 7).value, "Phòng GD&ĐT Ba Đình.xlsx")
        self.assertEqual(ws_m.cell(6, 7).value, "Phòng GD&ĐT Cầu Giấy.xlsx")
        wb_m.close()

        print("--> Test ExcelMerger PASSED successfully!")

if __name__ == "__main__":
    unittest.main()
