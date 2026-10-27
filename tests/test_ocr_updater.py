from __future__ import annotations

import tempfile
from pathlib import Path
from unittest.mock import patch

import openpyxl
import pytest

from src.ocr_updater import (
    AVAILABLE_MATCH_FIELDS,
    OCRUpdaterProcessor,
    match_student_row,
    normalize_match_val,
    resolve_field_to_col_idx,
)


def test_normalize_match_val():
    assert normalize_match_val("15.0", "stt") == "15"
    assert normalize_match_val(15, "stt") == "15"
    assert normalize_match_val("  Nguyễn   Văn   A  ", "ho_ten") == "nguyễn văn a"
    assert normalize_match_val("05/10/2008", "ngay_sinh") == "05/10/2008"
    assert normalize_match_val("2008-10-05", "ngay_sinh") == "05/10/2008"
    assert normalize_match_val("THCS Lê Quý Đôn", "ten_truong") == "thcs lê quý đôn"
    assert normalize_match_val(" A123 456 ", "so_hieu_bang") == "a123456"


def test_resolve_field_to_col_idx():
    assert resolve_field_to_col_idx("stt") == 1
    assert resolve_field_to_col_idx("ho_ten") == 5
    assert resolve_field_to_col_idx("Họ và tên") == 5
    assert resolve_field_to_col_idx("ngay_sinh") == 9
    assert resolve_field_to_col_idx("Ngày sinh") == 9
    assert resolve_field_to_col_idx("gioi_tinh") == 10
    assert resolve_field_to_col_idx("noi_sinh") == 13
    assert resolve_field_to_col_idx("so_hieu_bang") == 2
    assert resolve_field_to_col_idx("so_vao_so") == 3


def test_match_student_row():
    sheet_rows = [
        {"stt": "1", "ho_ten": "Nguyễn Văn An", "ngay_sinh": "15/03/2008", "so_hieu_bang": "A001", "ten_truong": "THCS A"},
        {"stt": "2", "ho_ten": "Trần Thị Bích", "ngay_sinh": "20/05/2008", "so_hieu_bang": "A002", "ten_truong": "THCS A"},
        {"stt": "3", "ho_ten": "Lê Văn Cường", "ngay_sinh": "10/11/2008", "so_hieu_bang": "A003", "ten_truong": "THCS B"},
    ]

    # Test khớp theo Họ tên + Ngày sinh
    cand1 = {"ho_ten": "Nguyễn Văn An", "ngay_sinh": "15/03/2008"}
    idx, all_cand, note = match_student_row(cand1, sheet_rows, ["ho_ten", "ngay_sinh"])
    assert idx == 0

    # Test khớp khi Họ tên bị sửa (có trong cac_truong_sua với gia_tri_cu)
    cand2 = {
        "ho_ten": "Trần Thị Ngọc Bích",
        "ngay_sinh": "20/05/2008",
        "cac_truong_sua": [
            {"ten_truong": "ho_ten", "gia_tri_cu": "Trần Thị Bích", "gia_tri_moi": "Trần Thị Ngọc Bích"}
        ]
    }
    idx2, all_cand2, note2 = match_student_row(cand2, sheet_rows, ["ho_ten", "ngay_sinh"])
    assert idx2 == 1

    # Test không tìm thấy
    cand3 = {"ho_ten": "Phạm Văn D", "ngay_sinh": "01/01/2008"}
    idx3, all_cand3, note3 = match_student_row(cand3, sheet_rows, ["ho_ten", "ngay_sinh"])
    assert idx3 is None
    assert len(all_cand3) == 0


def test_updater_end_to_end(tmp_path: Path):
    # Tạo 1 file Excel 28 cột mẫu
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Mau28Cot"

    # Header chuẩn 28 cột
    from src.smart_merger import TARGET_COLUMNS
    headers = [col["name"] for col in TARGET_COLUMNS]
    ws.append(headers)

    # Thêm 2 học sinh
    row1 = [None] * 28
    row1[0] = 1
    row1[1] = "A001"
    row1[2] = "01/2024"
    row1[3] = "Bằng tốt nghiệp trung học cơ sở"
    row1[4] = "Nguyễn Văn An"
    row1[8] = "15/03/2008"
    row1[9] = "Nam"
    row1[12] = "Hà Nội"
    row1[13] = "THCS Dịch Vọng"
    ws.append(row1)

    row2 = [None] * 28
    row2[0] = 2
    row2[1] = "A002"
    row2[2] = "02/2024"
    row2[3] = "Bằng tốt nghiệp trung học cơ sở"
    row2[4] = "Trần Thị Bích"
    row2[8] = "20/05/2008"
    row2[9] = "Nữ"
    row2[12] = "Hải Phòng"
    row2[13] = "THCS Dịch Vọng"
    ws.append(row2)

    excel_file = tmp_path / "ds_goc.xlsx"
    wb.save(excel_file)

    # Tạo tệp ảnh giả lập
    fake_img = tmp_path / "trang_01.jpg"
    fake_img.write_bytes(b"FAKE_IMAGE_BYTES")

    # Mock OCR response
    mock_json_response = """[
        {
            "loai": "SUA_DOI",
            "stt": 1,
            "ho_ten": "Nguyễn Văn An",
            "ngay_sinh": "16/03/2008",
            "cac_truong_sua": [
                {
                    "ten_truong": "ngay_sinh",
                    "gia_tri_cu": "15/03/2008",
                    "gia_tri_moi": "16/03/2008"
                }
            ],
            "can_kiem_tra": false
        },
        {
            "loai": "BO_SUNG",
            "stt": 3,
            "ho_ten": "Lê Văn Cường",
            "ngay_sinh": "10/11/2008",
            "gioi_tinh": "Nam",
            "noi_sinh": "Đà Nẵng",
            "ten_truong": "THCS Dịch Vọng",
            "ghi_chu": "Bổ sung danh sách",
            "can_kiem_tra": false
        }
    ]"""

    out_file = tmp_path / "ds_updated.xlsx"

    processor = OCRUpdaterProcessor()

    with patch("src.ocr_updater.extract_table_from_image", return_value=mock_json_response):
        res = processor.update_excel_from_images(
            excel_path=excel_file,
            image_paths=[fake_img],
            selected_match_keys=["ho_ten"],
            auto_append_new=True,
            output_path=out_file,
            api_key="test_api_key",
        )

    assert res["status"] == "success"
    assert res["modified_students"] == 1
    assert res["added_students"] == 1
    assert out_file.exists()

    # Kiểm tra nội dung sau khi cập nhật
    wb_res = openpyxl.load_workbook(out_file)
    ws_res = wb_res["Mau28Cot"]

    # Học sinh 1: Ngày sinh đổi từ 15/03/2008 thành 16/03/2008
    assert ws_res.cell(row=2, column=9).value == "16/03/2008"
    assert ws_res.cell(row=2, column=9).fill.start_color.rgb == "00FFF2CC" or ws_res.cell(row=2, column=9).fill.start_color.index == "FFF2CC"

    # Học sinh 3 được bổ sung ở dòng 4 (row 4)
    assert ws_res.cell(row=4, column=1).value == 3
    assert ws_res.cell(row=4, column=5).value == "Lê Văn Cường"
    assert ws_res.cell(row=4, column=9).value == "10/11/2008"

    # Kiểm tra sheet Lich_Su_Thay_Doi
    assert "Lich_Su_Thay_Doi" in wb_res.sheetnames
    ws_audit = wb_res["Lich_Su_Thay_Doi"]
    assert ws_audit.max_row >= 3  # Header + 2 entries


def test_updater_with_manual_review(tmp_path: Path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Mau28Cot"

    from src.smart_merger import TARGET_COLUMNS
    headers = [col["name"] for col in TARGET_COLUMNS]
    ws.append(headers)

    row1 = [None] * 28
    row1[0] = 1
    row1[4] = "Nguyễn Văn An"
    row1[8] = "15/03/2008"
    ws.append(row1)

    excel_file = tmp_path / "ds_review.xlsx"
    wb.save(excel_file)
    fake_img = tmp_path / "trang_01.jpg"
    fake_img.write_bytes(b"FAKE_BYTES")

    # AI trích xuất nhầm 18/03/2008
    mock_json = """[
        {
            "loai": "SUA_DOI",
            "stt": 1,
            "ho_ten": "Nguyễn Văn An",
            "ngay_sinh": "18/03/2008",
            "cac_truong_sua": [
                {"ten_truong": "ngay_sinh", "gia_tri_cu": "15/03/2008", "gia_tri_moi": "18/03/2008"}
            ]
        }
    ]"""

    out_file = tmp_path / "ds_reviewed.xlsx"
    processor = OCRUpdaterProcessor()

    # Giả lập người dùng can thiệp sửa tay thành 19/03/2008
    def fake_review_callback(req):
        assert req["student_name"] == "Nguyễn Văn An"
        changes = req["changes"]
        changes[0]["new_val"] = "19/03/2008"  # Người dùng sửa tay
        return {"action": "apply", "changes": changes}

    with patch("src.ocr_updater.extract_table_from_image", return_value=mock_json):
        res = processor.update_excel_from_images(
            excel_path=excel_file,
            image_paths=[fake_img],
            selected_match_keys=["ho_ten"],
            output_path=out_file,
            api_key="test_api_key",
            manual_review_callback=fake_review_callback,
        )

    assert res["status"] == "success"
    wb_res = openpyxl.load_workbook(out_file)
    ws_res = wb_res["Mau28Cot"]
    # Giá trị phải là 19/03/2008 do người dùng sửa tay
    assert ws_res.cell(row=2, column=9).value == "19/03/2008"


def test_column_mapping_manager(tmp_path: Path):
    from src.ocr_updater import OCRColumnMappingManager

    cfg_file = tmp_path / "custom_mappings.json"
    mgr = OCRColumnMappingManager(config_path=cfg_file)

    # 1. Mặc định có profile "Mặc định (Chuẩn 28 cột)"
    assert "Mặc định (Chuẩn 28 cột)" in mgr.list_profiles()

    # 2. Tạo profile mới cho ảnh có tên cột khác
    custom_map = {
        "Họ và tên học sinh": 5,
        "Ngày tháng sinh": 9,
        "Số hiệu VB": 2,
        "Vào sổ bộ": 3,
        "Xếp loại tốt nghiệp": 16,
    }
    mgr.save_profile("Sổ THCS Lạc Long Quân", custom_map)
    assert "Sổ THCS Lạc Long Quân" in mgr.list_profiles()
    assert mgr.get_active_profile_name() == "Sổ THCS Lạc Long Quân"

    # 3. Tra cứu mapping chính xác 100%
    loaded_map = mgr.get_mapping("Sổ THCS Lạc Long Quân")
    assert loaded_map["Họ và tên học sinh"] == 5
    assert resolve_field_to_col_idx("Họ và tên học sinh", custom_mapping=loaded_map) == 5
    assert resolve_field_to_col_idx("Vào sổ bộ", custom_mapping=loaded_map) == 3
    assert resolve_field_to_col_idx("Chưa ánh xạ", custom_mapping=loaded_map) is None

    # 4. Xóa profile
    assert mgr.delete_profile("Sổ THCS Lạc Long Quân") is True
    assert "Sổ THCS Lạc Long Quân" not in mgr.list_profiles()


def test_updater_with_custom_column_mapping(tmp_path: Path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Mau28Cot"

    from src.smart_merger import TARGET_COLUMNS
    headers = [col["name"] for col in TARGET_COLUMNS]
    ws.append(headers)

    row1 = [None] * 28
    row1[0] = 1
    row1[1] = "A001"
    row1[4] = "Lê Văn Cường"
    row1[8] = "10/11/2008"
    ws.append(row1)

    excel_file = tmp_path / "ds_custom_map.xlsx"
    wb.save(excel_file)
    fake_img = tmp_path / "trang_sua_la.jpg"
    fake_img.write_bytes(b"IMAGE_CONTENT")

    # Ảnh dùng tên cột khác lạ: "Họ và tên thí sinh" và "Số bằng TN"
    mock_json = """[
        {
            "loai": "SUA_DOI",
            "stt": 1,
            "ho_ten": "Lê Văn Cường",
            "cac_truong_sua": [
                {"ten_truong": "Họ và tên thí sinh", "gia_tri_cu": "Lê Văn Cường", "gia_tri_moi": "Lê Quốc Cường"},
                {"ten_truong": "Số bằng TN", "gia_tri_cu": "A001", "gia_tri_moi": "A999"}
            ]
        }
    ]"""

    out_file = tmp_path / "ds_custom_map_res.xlsx"
    processor = OCRUpdaterProcessor()

    custom_map = {
        "Họ và tên thí sinh": 5,
        "Số bằng TN": 2
    }

    with patch("src.ocr_updater.extract_table_from_image", return_value=mock_json):
        res = processor.update_excel_from_images(
            excel_path=excel_file,
            image_paths=[fake_img],
            selected_match_keys=["ho_ten"],
            column_mapping=custom_map,
            output_path=out_file,
            api_key="test_api_key",
        )

    assert res["status"] == "success"
    wb_res = openpyxl.load_workbook(out_file)
    ws_res = wb_res["Mau28Cot"]
    # Kiểm tra cột 5 và cột 2 đã được cập nhật chính xác 100% theo ánh xạ người dùng
    assert ws_res.cell(row=2, column=5).value == "Lê Quốc Cường"
    assert ws_res.cell(row=2, column=2).value == "A999"

