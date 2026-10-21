from __future__ import annotations

import os
import shutil
import tempfile
import unittest
from pathlib import Path
import openpyxl

from src.db_manager import DatabaseManager
from src.export_manager import (
    ExportManager,
    analyze_cohorts,
    parse_cohort_grad_year,
    sanitize_filename
)

class TestPhanHe3Export(unittest.TestCase):
    def setUp(self):
        self.temp_dir = Path(tempfile.mkdtemp())
        self.db_path = self.temp_dir / "test_export.db"
        self.db = DatabaseManager(self.db_path)
        self.out_dir = self.temp_dir / "output"
        self.out_dir.mkdir(parents=True, exist_ok=True)
        self.export_mgr = ExportManager()

        # Nạp dữ liệu mẫu đa dạng:
        # 2 Trường thuộc Huyện TP Hòa Bình:
        #   - Trường A (Độc Lập, Mã: 1460, Đơn vị: PT DTBT Độc Lập, Mã ĐV: 001, Xã: Phường Kỳ Sơn, Mã Xã: 04894, Huyện: TP Hòa Bình, Mã Huyện: H34)
        #   - Trường B (Lê Quý Đôn, Mã: 1462, Đơn vị: THCS Lê Quý Đôn, Mã ĐV: 002, Xã: Phường Kỳ Sơn, Mã Xã: 04894, Huyện: TP Hòa Bình, Mã Huyện: H34)
        # 3 Niên khóa: 2016-2020 (năm 2020), 2017-2021 (năm 2021), 2019-2023 (năm 2023)
        self.sample_students = [
            # Học sinh 1: Trường A, 2016-2020
            {
                "final_row": [
                    1, "'001234", "'0042", "Bằng tốt nghiệp trung học cơ sở", "Nguyễn Văn An",
                    None, "'036123456789", None, "15/05/2005", "Nam",
                    "Kinh", "Việt Nam", "Hòa Bình", "THCS Độc Lập", "2016-2020",
                    None, None, "TP Hòa Bình", "15/06/2020", "Phòng GD&ĐT TP Hòa Bình",
                    "Trưởng phòng", "Lê Văn Ký", "Cấp mới", None, "Đã số hóa đầy đủ, chính xác",
                    None, "001", "PT DTBT Độc Lập"
                ],
                "school_code": "1460",
                "commune_code": "04894",
                "commune": "Phường Kỳ Sơn",
                "district_code": "H34",
                "district_name": "TP Hòa Bình",
                "province_name": "Hòa Bình",
                "source_file": "2020/THCS_DocLap.xlsx",
                "source_sheet": "Lop9A",
                "row_idx": 4
            },
            # Học sinh 2: Trường A, 2017-2021
            {
                "final_row": [
                    2, "'001235", "'0043", "Bằng tốt nghiệp trung học cơ sở", "Trần Thị Bình",
                    None, "'036123456790", None, "20/08/2006", "Nữ",
                    "Mường", "Việt Nam", "Hòa Bình", "THCS Độc Lập", "2017-2021",
                    None, None, "TP Hòa Bình", "15/06/2021", "Phòng GD&ĐT TP Hòa Bình",
                    "Trưởng phòng", "Lê Văn Ký", "Cấp mới", None, "Đã số hóa đầy đủ, chính xác",
                    None, "001", "PT DTBT Độc Lập"
                ],
                "school_code": "1460",
                "commune_code": "04894",
                "commune": "Phường Kỳ Sơn",
                "district_code": "H34",
                "district_name": "TP Hòa Bình",
                "province_name": "Hòa Bình",
                "source_file": "2021/THCS_DocLap.xlsx",
                "source_sheet": "Lop9A",
                "row_idx": 5
            },
            # Học sinh 3: Trường A, 2019-2023 (Có lỗi để test DS_Loi)
            {
                "final_row": [
                    3, "'001236", "'0044", "Bằng tốt nghiệp trung học cơ sở", "Lê Văn Cường",
                    None, None, None, "2008", "Nam",
                    "Kinh", "Việt Nam", "Hòa Bình", "THCS Độc Lập", "2019-2023",
                    None, None, "TP Hòa Bình", "15/06/2023", "Phòng GD&ĐT TP Hòa Bình",
                    "Trưởng phòng", "Lê Văn Ký", "Cấp mới", None, "Số hoá chưa đầy đủ, chính xác",
                    "Chỉ có năm sinh", "001", "PT DTBT Độc Lập"
                ],
                "school_code": "1460",
                "commune_code": "04894",
                "commune": "Phường Kỳ Sơn",
                "district_code": "H34",
                "district_name": "TP Hòa Bình",
                "province_name": "Hòa Bình",
                "source_file": "2023/THCS_DocLap.xlsx",
                "source_sheet": "Lop9B",
                "row_idx": 6
            },
            # Học sinh 4: Trường B, 2016-2020
            {
                "final_row": [
                    4, "'001237", "'0045", "Bằng tốt nghiệp trung học cơ sở", "Phạm Văn Dũng",
                    None, "'036123456792", None, "10/02/2005", "Nam",
                    "Kinh", "Việt Nam", "Hòa Bình", "THCS Lê Quý Đôn", "2016-2020",
                    None, None, "TP Hòa Bình", "15/06/2020", "Phòng GD&ĐT TP Hòa Bình",
                    "Trưởng phòng", "Lê Văn Ký", "Cấp mới", None, "Đã số hóa đầy đủ, chính xác",
                    None, "002", "THCS Lê Quý Đôn"
                ],
                "school_code": "1462",
                "commune_code": "04894",
                "commune": "Phường Kỳ Sơn",
                "district_code": "H34",
                "district_name": "TP Hòa Bình",
                "province_name": "Hòa Bình",
                "source_file": "2020/THCS_LeQuyDon.xlsx",
                "source_sheet": "Lop9",
                "row_idx": 10
            }
        ]

        self.db.insert_students_batch(self.sample_students)

    def tearDown(self):
        shutil.rmtree(self.temp_dir, ignore_errors=True)

    def test_cohort_analysis(self):
        """Kiểm tra logic phân tích tính liên tục của niên khóa."""
        # 1 năm
        a1 = analyze_cohorts(["2016-2020"])
        self.assertTrue(a1["is_single"])
        self.assertTrue(a1["is_consecutive"])
        self.assertEqual(a1["min_year"], 2020)
        self.assertEqual(a1["max_year"], 2020)

        # 3 năm liên tiếp
        a2 = analyze_cohorts(["2016-2020", "2017-2021", "2018-2022"])
        self.assertFalse(a2["is_single"])
        self.assertTrue(a2["is_consecutive"])
        self.assertEqual(a2["min_year"], 2020)
        self.assertEqual(a2["max_year"], 2022)

        # 2 năm rời rạc (2020 và 2023)
        a3 = analyze_cohorts(["2016-2020", "2019-2023"])
        self.assertFalse(a3["is_single"])
        self.assertFalse(a3["is_consecutive"])
        self.assertEqual(a3["min_year"], 2020)
        self.assertEqual(a3["max_year"], 2023)

    def test_case_1_single_year_export(self):
        """
        Kiểm tra TH 1A: Xuất 1 năm duy nhất:
        - Xuất 1 file dữ liệu học sinh {Ten}_{Nam}.xlsx
        - Xuất 1 file DS_TL_So_Hoa với ĐÚNG 1 DÒNG DUY NHẤT
        """
        res = self.export_mgr.execute_export(
            db_manager=self.db,
            output_dir=self.out_dir,
            branch="branch_history",
            huyen_cu="TP Hòa Bình",
            ten_truong="THCS Độc Lập",
            selected_cohorts=["2016-2020"]
        )

        self.assertEqual(len(res["student_files"]), 1)
        self.assertEqual(len(res["doc_files"]), 1)

        st_file = res["student_files"][0]
        doc_file = res["doc_files"][0]

        self.assertTrue(st_file.name.endswith("_2020.xlsx"))
        self.assertTrue(doc_file.name.endswith("_2020_DS_TL_So_Hoa.xlsx"))

        # Đọc file học sinh: kiểm tra Text @ bảo toàn số 0 ở đầu (0042)
        wb_st = openpyxl.load_workbook(st_file)
        ws_st = wb_st.active
        self.assertEqual(ws_st.cell(3, 3).value, "0042")
        self.assertEqual(ws_st.cell(3, 3).data_type, 's')
        wb_st.close()

        # Đọc file tài liệu số hóa: phải có đúng 1 dòng dữ liệu (Row 3)
        wb_doc = openpyxl.load_workbook(doc_file)
        ws_doc = wb_doc.active
        self.assertEqual(ws_doc.cell(3, 1).value, 1)  # STT 1
        self.assertEqual(ws_doc.cell(3, 4).value, "1460_2020")  # Mã sổ gốc trường
        self.assertEqual(ws_doc.cell(3, 5).value, "Sổ gốc cấp bằng THCS trường THCS Độc Lập năm 2020")
        self.assertEqual(ws_doc.cell(3, 6).value, "2020")
        self.assertEqual(ws_doc.cell(3, 7).value, "2020")
        self.assertIsNone(ws_doc.cell(4, 1).value)  # Không có dòng thứ 2
        wb_doc.close()

    def test_case_2_consecutive_years_export(self):
        """
        Kiểm tra TH 1B: Xuất nhiều năm LIÊN TIẾP:
        - Xuất 1 file dữ liệu gộp {Ten}_{NamNho}_{NamLon}.xlsx
        - Xuất 1 file DS_TL_So_Hoa với ĐÚNG 1 DÒNG DUY NHẤT
        """
        res = self.export_mgr.execute_export(
            db_manager=self.db,
            output_dir=self.out_dir,
            branch="branch_history",
            huyen_cu="TP Hòa Bình",
            ten_truong="THCS Độc Lập",
            selected_cohorts=["2016-2020", "2017-2021"]
        )

        self.assertEqual(len(res["student_files"]), 1)
        self.assertEqual(len(res["doc_files"]), 1)

        st_file = res["student_files"][0]
        doc_file = res["doc_files"][0]

        self.assertTrue(st_file.name.endswith("_2020_2021.xlsx"))
        self.assertTrue(doc_file.name.endswith("_2020_2021_DS_TL_So_Hoa.xlsx"))

        # Kiểm tra file tài liệu số hóa: 1 dòng duy nhất có từ năm nhỏ đến năm lớn
        wb_doc = openpyxl.load_workbook(doc_file)
        ws_doc = wb_doc.active
        self.assertEqual(ws_doc.cell(3, 1).value, 1)
        self.assertEqual(ws_doc.cell(3, 4).value, "1460_2020-2021")
        self.assertEqual(ws_doc.cell(3, 5).value, "Sổ gốc cấp bằng THCS trường THCS Độc Lập từ 2020 đến 2021")
        self.assertEqual(ws_doc.cell(3, 6).value, "2020")
        self.assertEqual(ws_doc.cell(3, 7).value, "2021")
        self.assertIsNone(ws_doc.cell(4, 1).value)
        wb_doc.close()

    def test_case_3_discrete_years_export(self):
        """
        Kiểm tra TH 2: Xuất nhiều năm RỜI RẠC (2020 và 2023):
        - Xuất K file dữ liệu học sinh (mỗi năm 1 file)
        - VẪN CHỈ TẠO 1 FILE DS_TL_So_Hoa DUY NHẤT có K dòng!
        """
        res = self.export_mgr.execute_export(
            db_manager=self.db,
            output_dir=self.out_dir,
            branch="branch_history",
            huyen_cu="TP Hòa Bình",
            ten_truong="THCS Độc Lập",
            selected_cohorts=["2016-2020", "2019-2023"]
        )

        # 2 file học sinh riêng biệt
        self.assertEqual(len(res["student_files"]), 2)
        st_names = [f.name for f in res["student_files"]]
        self.assertTrue(any(n.endswith("_2020.xlsx") for n in st_names))
        self.assertTrue(any(n.endswith("_2023.xlsx") for n in st_names))

        # Chỉ có đúng 1 file DS_TL_So_Hoa duy nhất
        self.assertEqual(len(res["doc_files"]), 1)
        doc_file = res["doc_files"][0]
        self.assertEqual(doc_file.name, "THCS Độc Lập_DS_TL_So_Hoa.xlsx")

        # Đọc file DS_TL_So_Hoa: phải có 2 dòng tương ứng
        wb_doc = openpyxl.load_workbook(doc_file)
        ws_doc = wb_doc.active

        # Dòng 1 ứng với năm 2020
        self.assertEqual(ws_doc.cell(3, 1).value, 1)
        self.assertEqual(ws_doc.cell(3, 4).value, "1460_2020")
        self.assertEqual(ws_doc.cell(3, 5).value, "Sổ gốc cấp bằng THCS trường THCS Độc Lập năm 2020")
        self.assertEqual(ws_doc.cell(3, 6).value, "2020")
        self.assertEqual(ws_doc.cell(3, 7).value, "2020")

        # Dòng 2 ứng với năm 2023
        self.assertEqual(ws_doc.cell(4, 1).value, 2)
        self.assertEqual(ws_doc.cell(4, 4).value, "1460_2023")
        self.assertEqual(ws_doc.cell(4, 5).value, "Sổ gốc cấp bằng THCS trường THCS Độc Lập năm 2023")
        self.assertEqual(ws_doc.cell(4, 6).value, "2023")
        self.assertEqual(ws_doc.cell(4, 7).value, "2023")

        self.assertIsNone(ws_doc.cell(5, 1).value)
        wb_doc.close()

    def test_level_district_override_code_and_name(self):
        """
        Kiểm tra Quy tắc Cấp Huyện cũ:
        - Tên huyện = cụ thể, Tên trường = 'Tất cả'
        - Cột 27 được gán bằng Mã huyện ('H34')
        - Cột 28 được gán bằng Tên huyện ('TP Hòa Bình')
        """
        res = self.export_mgr.execute_export(
            db_manager=self.db,
            output_dir=self.out_dir,
            branch="branch_history",
            huyen_cu="TP Hòa Bình",
            ten_truong="Tất cả",
            selected_cohorts=["2016-2020"]
        )

        st_file = res["student_files"][0]
        wb = openpyxl.load_workbook(st_file)
        ws = wb.active

        # Dòng 3 (HS1) và Dòng 4 (HS4) thuộc 2 trường khác nhau nhưng đều gán Cột 27 và 28 theo Huyện
        self.assertEqual(ws.cell(3, 27).value, "H34")
        self.assertEqual(ws.cell(3, 28).value, "TP Hòa Bình")
        self.assertEqual(ws.cell(4, 27).value, "H34")
        self.assertEqual(ws.cell(4, 28).value, "TP Hòa Bình")
        wb.close()

        # Kiểm tra file tài liệu số hóa cấp Huyện
        doc_file = res["doc_files"][0]
        wb_doc = openpyxl.load_workbook(doc_file)
        ws_doc = wb_doc.active
        self.assertEqual(ws_doc.cell(3, 4).value, "H34_2020")
        self.assertEqual(ws_doc.cell(3, 5).value, "Sổ gốc cấp bằng THCS TP Hòa Bình năm 2020")
        wb_doc.close()

    def test_error_file_export(self):
        """Kiểm tra xuất file lỗi có kèm 3 cột vết nguồn gốc (AC..AE)."""
        res = self.export_mgr.execute_export(
            db_manager=self.db,
            output_dir=self.out_dir,
            branch="branch_history",
            huyen_cu="TP Hòa Bình",
            ten_truong="THCS Độc Lập",
            selected_cohorts=["2019-2023"],
            export_error_files=True
        )

        self.assertEqual(len(res["error_files"]), 1)
        err_file = res["error_files"][0]
        self.assertTrue(err_file.name.endswith("_DS_Loi_2023.xlsx"))

        wb_err = openpyxl.load_workbook(err_file)
        ws_err = wb_err.active
        self.assertEqual(ws_err.cell(2, 29).value, "Tệp nguồn gốc")
        self.assertEqual(ws_err.cell(2, 30).value, "Sheet nguồn")
        self.assertEqual(ws_err.cell(2, 31).value, "Dòng Excel gốc")

        # Dòng dữ liệu học sinh bị lỗi (Học sinh 3)
        self.assertEqual(ws_err.cell(3, 5).value, "Lê Văn Cường")
        self.assertEqual(ws_err.cell(3, 26).value, "Chỉ có năm sinh")
        self.assertEqual(ws_err.cell(3, 29).value, "2023/THCS_DocLap.xlsx")
        self.assertEqual(ws_err.cell(3, 30).value, "Lop9B")
        self.assertEqual(str(ws_err.cell(3, 31).value), "6")
        wb_err.close()

    def test_level_commune_override(self):
        """
        Kiểm tra Cấp Xã/phường mới:
        - Xã = cụ thể, Đơn vị = 'Tất cả'
        - Cột 27 = '04894' (Mã xã), Cột 28 = 'Phường Kỳ Sơn' (Tên xã)
        - Tên tài liệu: 'Sổ gốc cấp bằng THCS Phường Kỳ Sơn...'
        """
        res = self.export_mgr.execute_export(
            db_manager=self.db,
            output_dir=self.out_dir,
            branch="branch_new_admin",
            xa_phuong="Phường Kỳ Sơn",
            ten_don_vi="Tất cả",
            selected_cohorts=["2016-2020"]
        )

        st_file = res["student_files"][0]
        wb = openpyxl.load_workbook(st_file)
        ws = wb.active
        self.assertEqual(ws.cell(3, 27).value, "04894")
        self.assertEqual(ws.cell(3, 28).value, "Phường Kỳ Sơn")
        wb.close()

        doc_file = res["doc_files"][0]
        wb_doc = openpyxl.load_workbook(doc_file)
        ws_doc = wb_doc.active
        self.assertEqual(ws_doc.cell(3, 2).value, "04894")
        self.assertEqual(ws_doc.cell(3, 3).value, "Phường Kỳ Sơn")
        self.assertEqual(ws_doc.cell(3, 4).value, "04894_2020")
        self.assertEqual(ws_doc.cell(3, 5).value, "Sổ gốc cấp bằng THCS Phường Kỳ Sơn năm 2020")
        wb_doc.close()

    def test_level_unit_preserve(self):
        """
        Kiểm tra Cấp Đơn vị quản lý:
        - Xã = cụ thể, Đơn vị = cụ thể
        - Cột 27 và Cột 28 GIỮ NGUYÊN
        - Tên tài liệu: 'Sổ gốc cấp bằng THCS {Tên Đơn vị}...'
        """
        res = self.export_mgr.execute_export(
            db_manager=self.db,
            output_dir=self.out_dir,
            branch="branch_new_admin",
            xa_phuong="Phường Kỳ Sơn",
            ten_don_vi="PT DTBT Độc Lập",
            selected_cohorts=["2016-2020"]
        )

        st_file = res["student_files"][0]
        wb = openpyxl.load_workbook(st_file)
        ws = wb.active
        self.assertEqual(ws.cell(3, 27).value, "001")
        self.assertEqual(ws.cell(3, 28).value, "PT DTBT Độc Lập")
        wb.close()

        doc_file = res["doc_files"][0]
        wb_doc = openpyxl.load_workbook(doc_file)
        ws_doc = wb_doc.active
        self.assertEqual(ws_doc.cell(3, 4).value, "001_2020")
        self.assertEqual(ws_doc.cell(3, 5).value, "Sổ gốc cấp bằng THCS PT DTBT Độc Lập năm 2020")
        wb_doc.close()

    def test_sorting_stt_school_cohort(self):
        """Kiểm tra sắp xếp theo STT -> Trường -> Niên khóa."""
        students = self.db.query_students_for_export(
            sort_mode="stt_school_cohort"
        )
        self.assertEqual(len(students), 4)
        # Các bản ghi được sắp xếp theo stt tăng dần
        stt_list = [s["stt"] for s in students]
        self.assertEqual(stt_list, sorted(stt_list))

if __name__ == "__main__":
    unittest.main()
